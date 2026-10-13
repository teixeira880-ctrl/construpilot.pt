"""Étape 1 — Analyse documentaire vérifiable (XLSX + PDF).

Objectifs (déterministes, sans invention) :
- Extraire les données avec leur PROVENANCE exacte (feuille!cellule pour Excel,
  page + zone pour PDF).
- Lire séparément la FORMULE et son RÉSULTAT en cache : si le résultat est absent,
  le SIGNALER (jamais l'inventer).
- Distinguer clairement zéro / cellule vide / illisible / absente.
- Ne pas ignorer en silence les feuilles ou lignes masquées (les inclure + signaler).
- Éviter le double comptage détail/sous-total.
- Contrôler l'arithmétique (quantité × PU, TVA, arrondis) via Decimal.

Le cœur est 100 % déterministe. L'IA n'est pas requise ici (l'interprétation des
colonnes ambiguës = tranche suivante). Réutilise la détection d'en-têtes du DPGF.
"""
from __future__ import annotations

import io
import re
import time
from decimal import Decimal, ROUND_HALF_UP, InvalidOperation

import openpyxl
from openpyxl.utils import get_column_letter

try:
    import fitz  # PyMuPDF
except Exception:  # pragma: no cover
    fitz = None

from engines.dpgf_engine import _find_header, _is_lot_title, fold, _SKIP

ENGINE_VERSION = "doc_intel/1.0.0"

# ---------------------------------------------------------------------------
# Nombres FR / PT (virgule décimale, séparateurs de milliers espace/point)
# ---------------------------------------------------------------------------

def parse_number(v) -> dict:
    """Renvoie {value, status, raw}. status ∈ ok|empty|absent|illegible.

    - absent   : la cellule n'existe pas / None
    - empty    : chaîne vide
    - illegible: contenu non convertible en nombre
    - ok       : nombre lu (0 reste une valeur « ok », jamais confondu avec vide)
    """
    if v is None:
        return {"value": None, "status": "absent", "raw": None}
    if isinstance(v, bool):
        return {"value": None, "status": "illegible", "raw": v}
    if isinstance(v, (int, float)):
        return {"value": float(v), "status": "ok", "raw": v}
    raw = str(v)
    s = raw.replace("\u00a0", " ").replace("\u202f", " ").strip()
    if s == "":
        return {"value": None, "status": "empty", "raw": raw}
    cleaned = re.sub(r"[^\d,.\-]", "", s)
    if cleaned in ("", "-", ".", ",", "-.", "-,"):
        return {"value": None, "status": "illegible", "raw": raw}
    has_c, has_d = "," in cleaned, "." in cleaned
    if has_c and has_d:
        if cleaned.rfind(",") > cleaned.rfind("."):          # 1.234,56 → FR/PT
            cleaned = cleaned.replace(".", "").replace(",", ".")
        else:                                                 # 1,234.56 → EN
            cleaned = cleaned.replace(",", "")
    elif has_c:
        cleaned = cleaned.replace(",", "") if cleaned.count(",") > 1 else cleaned.replace(",", ".")
    elif has_d and cleaned.count(".") > 1:
        cleaned = cleaned.replace(".", "")                    # 1.234.567 → milliers
    try:
        return {"value": float(cleaned), "status": "ok", "raw": raw}
    except ValueError:
        return {"value": None, "status": "illegible", "raw": raw}


def guess_language(sample: str) -> str | None:
    s = fold(sample or "")
    if not s:
        return None
    pt = sum(s.count(k) for k in ("betao", "area", "empreitada", "orcamento", "revestimento", "cao "))
    fr = sum(s.count(k) for k in ("maconnerie", "cloison", "platre", "peinture", "plancher", "devis", "lot "))
    if pt > fr and pt > 0:
        return "pt"
    if fr > 0:
        return "fr"
    return None


def _looks_subtotal(text: str) -> bool:
    t = fold(text)
    return any(k in t for k in ("sous total", "sous-total", "s/total", "s total",
                                "total lot", "total ht", "total general", "sous total ht",
                                "subtotal", "total do lote", "total geral"))


# ---------------------------------------------------------------------------
# EXCEL
# ---------------------------------------------------------------------------

def extract_xlsx(raw: bytes) -> dict:
    t0 = time.time()
    wb_v = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)   # résultats en cache
    wb_f = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)  # formules

    sheets: list[dict] = []
    lines: list[dict] = []
    warnings: list[dict] = []
    unread: list[dict] = []
    sample_text: list[str] = []

    for name in wb_v.sheetnames:
        ws, wsf = wb_v[name], wb_f[name]
        state = getattr(ws, "sheet_state", "visible")
        if state != "visible":
            warnings.append({"type": "hidden_sheet", "sheet": name, "state": state,
                             "message": f"Feuille masquée « {name} » analysée (non ignorée)"})
        hrow, cols = _find_header(ws)
        if not hrow:
            sheets.append({"sheet": name, "state": state, "recognised": False, "lines": 0})
            unread.append({"type": "sheet_unrecognised", "sheet": name,
                           "message": "En-têtes non reconnus (interprétation IA prévue en tranche suivante)"})
            continue

        merged: dict = {}
        for rng in ws.merged_cells.ranges:
            tl = ws.cell(row=rng.min_row, column=rng.min_col).value
            for rr in range(rng.min_row, rng.max_row + 1):
                for cc in range(rng.min_col, rng.max_col + 1):
                    merged[(rr, cc)] = tl

        c_des, c_unit = cols.get("design"), cols.get("unit")
        c_qty, c_pu, c_total = cols.get("qty"), cols.get("pu"), cols.get("total")

        def cval(r, c):
            return merged.get((r, c), ws.cell(row=r, column=c).value) if c else None

        def coord(r, c):
            return f"{name}!{get_column_letter(c)}{r}" if c else None

        def formula_flag(r, c, key, flags):
            if not c:
                return
            fc = wsf.cell(row=r, column=c)
            if fc.data_type == "f":
                cached = ws.cell(row=r, column=c).value
                if cached is None:
                    flags.append(f"formula_result_missing:{key}")

        count, lot, pending = 0, "", []
        for r in range(hrow + 1, ws.max_row + 1):
            hidden = bool(ws.row_dimensions[r].hidden) if r in ws.row_dimensions else False
            text = str(cval(r, c_des) or "").strip()
            qty = parse_number(cval(r, c_qty))
            unit = str(cval(r, c_unit) or "").strip()
            pu = parse_number(cval(r, c_pu))
            amt = parse_number(cval(r, c_total))
            has_values = any(x["value"] is not None for x in (qty, pu, amt)) or bool(unit)
            if not text and not has_values:
                continue
            if text:
                sample_text.append(text)
            low = fold(text)
            if text and any(k in low for k in (fold(s) for s in _SKIP)) and not has_values:
                continue

            flags: list[str] = []
            if hidden:
                flags.append("hidden_row")
            formula_flag(r, c_pu, "pu", flags)
            formula_flag(r, c_total, "amount", flags)
            formula_flag(r, c_qty, "qty", flags)

            prov = {"header_row": hrow, "designation": coord(r, c_des), "qty": coord(r, c_qty),
                    "unit": coord(r, c_unit), "pu": coord(r, c_pu), "amount": coord(r, c_total)}

            if text and _is_lot_title(text, qty["value"], pu["value"]):
                lot, pending = text, []
                lines.append({"line_id": f"l{len(lines)+1}", "sheet": name, "row": r,
                              "kind": "lot_title", "hidden": hidden, "lot": lot,
                              "designation": text, "description": "", "unit": None,
                              "qty": qty, "pu": pu, "amount": amt, "provenance": prov, "flags": flags})
                continue

            if _looks_subtotal(text) and (qty["value"] in (None, 0)) and (pu["value"] in (None, 0)):
                flags.append("subtotal_excluded_from_sum")
                lines.append({"line_id": f"l{len(lines)+1}", "sheet": name, "row": r,
                              "kind": "subtotal", "hidden": hidden, "lot": lot or "Sans lot",
                              "designation": text, "description": "", "unit": unit or None,
                              "qty": qty, "pu": pu, "amount": amt, "provenance": prov, "flags": flags})
                continue

            if qty["value"] is None and not unit and pu["value"] is None and amt["value"] is None:
                pending.append(text)
                pending = pending[-6:]
                continue

            lines.append({"line_id": f"l{len(lines)+1}", "sheet": name, "row": r,
                          "kind": "detail", "hidden": hidden, "lot": lot or "Sans lot",
                          "designation": text, "description": " ".join(pending)[:1500],
                          "unit": unit or None, "qty": qty, "pu": pu, "amount": amt,
                          "provenance": prov, "flags": flags})
            pending = []
            count += 1

        sheets.append({"sheet": name, "state": state, "recognised": True, "lines": count,
                       "header_row": hrow, "columns": {k: get_column_letter(v) for k, v in cols.items()}})

    return {"format": "xlsx", "language": guess_language(" ".join(sample_text[:200])),
            "sheets": sheets, "lines": lines, "warnings": warnings, "unread": unread,
            "duration_ms": int((time.time() - t0) * 1000)}


# ---------------------------------------------------------------------------
# PDF (texte natif + coordonnées ; scans signalés pour OCR ultérieur)
# ---------------------------------------------------------------------------

def extract_pdf(raw: bytes) -> dict:
    if fitz is None:
        raise RuntimeError("PyMuPDF (fitz) indisponible")
    t0 = time.time()
    doc = fitz.open(stream=raw, filetype="pdf")
    pages: list[dict] = []
    lines: list[dict] = []
    tables: list[dict] = []
    warnings: list[dict] = []
    unread: list[dict] = []
    sample_text: list[str] = []

    for pno in range(doc.page_count):
        page = doc[pno]
        d = page.get_text("dict")
        n = 0
        text_present = False
        for block in d.get("blocks", []):
            for ln in block.get("lines", []):
                txt = "".join(sp.get("text", "") for sp in ln.get("spans", [])).strip()
                if not txt:
                    continue
                text_present = True
                bb = ln.get("bbox", [0, 0, 0, 0])
                n += 1
                if len(sample_text) < 200:
                    sample_text.append(txt)
                lines.append({"line_id": f"p{pno+1}_{n}", "page": pno + 1,
                              "bbox": [round(float(x), 1) for x in bb], "text": txt[:1000]})
        img_count = len(page.get_images(full=True))
        scanned = (not text_present) and img_count > 0
        if scanned:
            warnings.append({"type": "scanned_page", "page": pno + 1,
                             "message": "Page probablement scannée — OCR à activer (IA vision)"})
            unread.append({"type": "scanned_page", "page": pno + 1})

        n_tab = 0
        try:
            finder = page.find_tables()
            for ti, tab in enumerate(getattr(finder, "tables", []) or []):
                cells = tab.extract()
                tables.append({"table_id": f"p{pno+1}_t{ti+1}", "page": pno + 1,
                               "bbox": [round(float(x), 1) for x in tab.bbox], "rows": cells})
                n_tab += 1
        except Exception:
            pass
        pages.append({"page": pno + 1, "lines": n, "scanned": scanned,
                      "images": img_count, "tables": n_tab})
    doc.close()
    return {"format": "pdf", "language": guess_language(" ".join(sample_text)),
            "pages": pages, "lines": lines, "tables": tables,
            "warnings": warnings, "unread": unread,
            "duration_ms": int((time.time() - t0) * 1000)}


# ---------------------------------------------------------------------------
# Contrôles arithmétiques déterministes (Decimal)
# ---------------------------------------------------------------------------

def _D(x) -> Decimal:
    try:
        return Decimal(str(x))
    except (InvalidOperation, ValueError):
        return Decimal(0)


def arithmetic_check(lines: list[dict], vat_rate: float | None = None) -> dict:
    checks: list[dict] = []
    sum_computed = Decimal(0)
    sum_file = Decimal(0)
    for ln in lines:
        if ln.get("kind") != "detail":
            continue
        q = ln["qty"]["value"]
        pu = ln["pu"]["value"]
        amt = ln["amount"]["value"]
        e = {"line_id": ln["line_id"], "sheet": ln.get("sheet"), "row": ln.get("row"),
             "provenance": ln["provenance"].get("amount") or ln["provenance"].get("pu")}
        if q is None or pu is None:
            e["status"] = "insufficient"
            e["reason"] = "quantité ou PU manquant"
            if amt is not None:
                sum_file += _D(amt)
            checks.append(e)
            continue
        computed = (_D(q) * _D(pu)).quantize(Decimal("0.01"), ROUND_HALF_UP)
        sum_computed += computed
        e["computed"] = float(computed)
        if amt is None:
            e["status"] = "computed_only"
        else:
            f = _D(amt)
            sum_file += f
            diff = abs(f - computed)
            tol = max(Decimal("0.01"), computed * Decimal("0.005"))
            e["file_amount"] = float(f)
            e["diff"] = float(diff)
            e["status"] = "ok" if diff <= tol else "mismatch"
        checks.append(e)

    subtotals = [{"line_id": l["line_id"], "row": l["row"], "sheet": l.get("sheet"),
                  "designation": l["designation"], "amount": l["amount"]["value"]}
                 for l in lines if l.get("kind") == "subtotal"]

    total_ht = sum_computed
    vat_amount = (total_ht * _D(vat_rate) / Decimal(100)).quantize(Decimal("0.01"), ROUND_HALF_UP) if vat_rate else None
    total_ttc = (total_ht + vat_amount) if vat_amount is not None else None

    summary = {
        "detail_lines": sum(1 for l in lines if l.get("kind") == "detail"),
        "sum_computed_ht": float(sum_computed),
        "sum_file_amounts": float(sum_file),
        "mismatches": sum(1 for c in checks if c.get("status") == "mismatch"),
        "insufficient": sum(1 for c in checks if c.get("status") == "insufficient"),
        "subtotals_excluded": len(subtotals),
        "vat_rate": vat_rate,
        "vat_amount": float(vat_amount) if vat_amount is not None else None,
        "total_ttc": float(total_ttc) if total_ttc is not None else None,
    }
    return {"checks": checks, "subtotals": subtotals, "summary": summary}


# ---------------------------------------------------------------------------
# Orchestration
# ---------------------------------------------------------------------------

def analyse(raw: bytes, filename: str, vat_rate: float | None = None) -> dict:
    ext = (filename or "").lower().rsplit(".", 1)[-1] if "." in (filename or "") else ""
    if ext in ("xlsx", "xlsm"):
        res = extract_xlsx(raw)
    elif ext == "pdf":
        res = extract_pdf(raw)
    else:
        raise ValueError(f"Format non pris en charge à cette étape : .{ext or '?'} (XLSX ou PDF attendu)")
    res["arithmetic"] = arithmetic_check(res.get("lines", []), vat_rate) if res["format"] == "xlsx" else None
    res["engine_version"] = ENGINE_VERSION
    return res
