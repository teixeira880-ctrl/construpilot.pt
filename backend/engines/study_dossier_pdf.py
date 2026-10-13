"""Reusable beige/copper estimate dossier approved on the real plan exercises."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape

from reportlab.lib.colors import HexColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import SimpleDocTemplate, Paragraph, Table, TableStyle, Spacer, PageBreak, Image, KeepTogether

COPPER, CREAM, PALE = map(HexColor, ("#925535", "#F7F1E9", "#F0E7DC"))
INK, MUTED, LINE = map(HexColor, ("#302B26", "#6E6359", "#DFD1C2"))
WHITE = HexColor("#FFFFFF")
WIDTH = A4[0] - 96


def fonts():
    from reportlab import __file__ as rlfile
    candidates = [Path(__file__).parent.parent / "assets/fonts", Path("/usr/share/fonts/truetype/dejavu")]
    folder = next((p for p in candidates if (p/"DejaVuSans.ttf").exists()), None)
    for name, filename in (("Dossier", "DejaVuSans.ttf"), ("DossierBold", "DejaVuSans-Bold.ttf")):
        if name not in pdfmetrics.getRegisteredFontNames():
            path = folder/filename if folder else Path(rlfile).parent/"fonts"/("VeraBd.ttf" if name.endswith("Bold") else "Vera.ttf")
            pdfmetrics.registerFont(TTFont(name, str(path)))
    pdfmetrics.registerFontFamily("Dossier", normal="Dossier", bold="DossierBold")


def fmt(value, currency=False):
    if value is None:
        return "N/C"
    return f"{value:,.2f}".replace(",", " ").replace(".", ",") + (" €" if currency else "")


def text(value):
    return escape(str(value or "")).replace("\n", "<br/>")


class NumberedCanvas(Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.saved_pages = []

    def showPage(self):
        self.saved_pages.append(dict(self.__dict__))
        self._startPage()

    def save(self):
        count = len(self.saved_pages)
        for state in self.saved_pages:
            self.__dict__.update(state)
            self.setFillColor(MUTED)
            self.setFont("Dossier", 7.2)
            self.drawRightString(A4[0]-48, 29, f"{self._pageNumber} / {count}")
            super().showPage()
        super().save()


def build_dossier_pdf(plan, study, branding=None, source_images=None):
    fonts()
    branding = branding or {}
    title = (study.get("project") or {}).get("title") or plan.get("filename") or "Étude de construction"
    address = (study.get("project") or {}).get("address") or ""
    ref = str(plan.get("plan_id") or "DOSSIER")[-32:]
    date = str(study.get("generated_at") or "")[:10]
    s = {"body": ParagraphStyle("b", fontName="Dossier", fontSize=9.1, leading=13, textColor=INK, spaceAfter=8),
         "small": ParagraphStyle("s", fontName="Dossier", fontSize=7.9, leading=10.8, textColor=MUTED, spaceAfter=7),
         "cell": ParagraphStyle("c", fontName="Dossier", fontSize=8.1, leading=10.8, textColor=INK),
         "head": ParagraphStyle("h", fontName="DossierBold", fontSize=8, leading=10.8, textColor=WHITE),
         "title": ParagraphStyle("t", fontName="DossierBold", fontSize=21, leading=26, textColor=INK, spaceAfter=14),
         "h2": ParagraphStyle("h2", fontName="DossierBold", fontSize=11, leading=15, textColor=COPPER, spaceBefore=10, spaceAfter=9, keepWithNext=True),
         "money": ParagraphStyle("m", fontName="DossierBold", fontSize=24, leading=32, textColor=WHITE),
         "white": ParagraphStyle("w", fontName="Dossier", fontSize=9, leading=13, textColor=WHITE)}
    def para(value, style="body", raw=False):
        return Paragraph(str(value) if raw else text(value), s[style])
    def table(headers, data, widths, amounts=()):
        rows = [[para(v, "head") for v in headers]]
        rows.extend([[v if isinstance(v, Paragraph) else para(v, "cell") for v in row] for row in data])
        t = Table(rows, colWidths=widths, repeatRows=1, hAlign="LEFT")
        styles = [("BACKGROUND", (0, 0), (-1, 0), COPPER), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                  ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, PALE]),
                  ("LINEBELOW", (0, 1), (-1, -1), .3, LINE),
                  ("LEFTPADDING", (0, 0), (-1, -1), 7), ("RIGHTPADDING", (0, 0), (-1, -1), 7),
                  ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5)]
        for col in amounts:
            for row in rows[1:]:
                row[col].style = ParagraphStyle("amount", parent=s["cell"], alignment=2, fontName="DossierBold")
        t.setStyle(TableStyle(styles))
        return t
    def heading(value):
        return para(value, "title")
    buf = BytesIO()
    pdf = SimpleDocTemplate(buf, pagesize=A4, leftMargin=48, rightMargin=48, topMargin=69, bottomMargin=60,
                            title=f"Devis estimatif gros œuvre - {title}", author=branding.get("company_name") or "ConstruPilot")
    totals = study["totals"]
    story = [heading("Devis estimatif gros œuvre"), para(title + (" | " + address if address else ""), "small"),
             para(f"Fourniture et pose - {sum(len(l['lines']) for l in study['lots'])} postes. {study.get('scope') or 'Gros œuvre'}."),
             para("Estimation à partir des pièces du chantier. Les quantités, leurs formules et les provisions sont détaillées. Prix et choix constructifs restent à confirmer.")]
    budget = Table([[para("BUDGET PARTIEL HT" if totals.get("partial") else "BUDGET ESTIMATIF HT", "white")],
                    [para(fmt(totals["ht"], True), "money")],
                    [para(f"{fmt(totals['ttc'], True)} TTC | TVA {fmt(totals.get('tva_rate')*100) if totals.get('tva_rate') is not None else 'N/C'} %", "white")]], colWidths=[WIDTH])
    budget.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, -1), COPPER), ("LEFTPADDING", (0, 0), (-1, -1), 17),
                                ("TOPPADDING", (0, 0), (-1, -1), 7), ("BOTTOMPADDING", (0, 0), (-1, -1), 7)]))
    story += [budget, Spacer(1, 18), table(["Poste", "Montant HT"], [[l["lot"]+(" (partiel)" if l.get("partial") else ""),fmt(l["total_ht"], True)] for l in study["lots"]], [WIDTH-129,129], (1,))]
    story += [para("Lecture du budget", "h2"), para("M : mesure contrôlée par l’utilisateur ; H : hypothèse ou mesure extraite à contrôler ; P : provision ; N/C : non chiffré.", "small")]
    for variant in study.get("variants", []):
        story += [para(f"{variant['label']} : déduction de {fmt(variant['deduction_ht'], True)} HT ; budget résultant {fmt(variant['total_ht'], True)} HT. {variant.get('note') or ''}", "small")]
    story += [para("Document budgétaire non contractuel. Les réserves et exclusions font partie intégrante de cette estimation.", "small"), PageBreak(), heading("Devis / ouvrages détaillés")]
    data = []
    for count, r in enumerate([r for lot in study['lots'] for r in lot['lines']], 1):
        data.append([str(count).zfill(2), r['designation'], para(fmt(r['qty'])+"<br/><font color='#925535'>"+text(r['quantity_status'])+"</font>", "cell", True), r['unit'], fmt(r['unit_price']), fmt(r['total'])])
    story.append(table(["N°", "Désignation", "Qté", "U", "PU HT €", "Total HT €"], data, [30,WIDTH-285,54,35,76,90], (4,5)))
    story += [Spacer(1, 14), table(["Récapitulatif", "Montant"], [["Total HT"+(" partiel" if totals.get("partial") else ""),fmt(totals['ht'], True)],["TVA",fmt(totals['tva'], True)],["Total TTC estimatif",fmt(totals['ttc'], True)]],[WIDTH-150,150],(1,)),
              PageBreak(), heading("Métrés / calculs et sources")]
    geometry = [r for r in study["items"] if r["included"] and r["unit"] not in ("ft", "ens.", "u", "1") and (r["sources"] or r.get("error"))]
    story.append(para("Chaque formule utilise les repères de mesures ci-dessous. Les forfaits restent des provisions, distinctes des quantités géométriques.", "small"))
    story.append(table(["Ouvrage / repères de calcul", "Quantité", "Statut"],
        [[para(text(r['designation'])+"<br/><font size='7' color='#6E6359'>"+text(r['expression'])+"</font>"+("<br/>"+text(r['error']) if r.get('error') else ''), "cell", True),
          fmt(r['qty'])+" "+r['unit'],r['quantity_status']] for r in geometry], [WIDTH-133,87,46]))
    evidence = {p['id']:p for r in geometry for p in r['sources'] if not p.get('expression')}
    story += [para("Registre des mesures / sources", "h2"), table(["Repère / mesure", "Valeur", "Source / hypothèse"],
        [[p['id']+" / "+p['label'], fmt(p['value'])+" "+p['unit'],
          "; ".join(f"p. {x['page']} : {x['text']}" for x in p['sources']) or p.get('note') or "Valeur saisie / à confirmer"]
         for p in evidence.values()], [150,70,WIDTH-220])]
    derived = {p['id']:p for r in geometry for p in r['sources'] if p.get('expression')}
    if derived:
        story += [para("Mesures dérivées", "h2"), table(["Repère", "Formule", "Valeur"],
            [[p['id'],p['expression'],fmt(p['value'])+" "+p['unit']] for p in derived.values()], [115,WIDTH-195,80])]
    if source_images:
        story += [PageBreak(), heading("Pièces de référence")]
        for label, png in source_images:
            from PIL import Image as PILImage
            with PILImage.open(BytesIO(png)) as im:
                width, height = im.size
            ratio = min(WIDTH/width, 440/height)
            story.append(KeepTogether([para(label,"h2"),Image(BytesIO(png), width=width*ratio,height=height*ratio),Spacer(1,12)]))
    story += [PageBreak(), heading("Prestations et conventions")]
    for lot in study['lots']:
        story.append(para(lot['lot'], 'h2'))
        for r in lot['lines']:
            story.append(para("<b>"+text(r['id'])+" · "+text(r['designation'])+"</b> : "+text(r.get('detail') or 'Périmètre à confirmer.'), 'small', True))
    story += [PageBreak(), heading("Hypothèses et points à confirmer")]
    for name, values in (("Hypothèses de chantier", study.get("assumptions", [])),
                         ("Contrôles et informations manquantes", study.get("warnings", [])),
                         ("Contrôles de calcul", [f"{r['item_id']} : {r['message']}" for r in study.get("controls", [])])):
        if values:
            story.append(para(name, "h2"))
            story.extend(para("• "+str(v), "body") for v in values)
    story += [para("Structure et exécution", "h2"),para("Un PC ne dimensionne pas les fondations, les soutènements ni les armatures. Les hypothèses de matériau et de système constructif servent au budget ; elles doivent être réconciliées avec les études d’exécution. Une quantité manquante ne vaut jamais zéro.")]
    story += [PageBreak(), heading("Prix, périmètre et traçabilité"),para("Base des prix", "h2"),
              para("Prix de vente fourniture et pose. Le barème budgétaire n’est pas un tarif fournisseur ni une offre d’entreprise. Aucune marge n’est ajoutée automatiquement ; les coûts de revient et la marge ne sont pas déterminés par cette étude."),
              para("Origine des prix", "h2"),
              table(["Origine", "Postes"], [[label, ", ".join(str(i+1).zfill(2) for i,r in enumerate([x for x in study['items'] if x['included']]) if r['price_source']==key)]
                  for key,label in [("prix_utilisateur","Prix saisis"),("estimation_budgetaire","Barème budgétaire à confirmer"),("non_chiffre","Prix non fournis")]
                  if any(r['price_source']==key for r in study['items'] if r['included'])],[185,WIDTH-185]),
              para("Hors périmètre", "h2")]
    story.extend(para("• "+str(x)) for x in study.get("exclusions", []))
    coverage = study.get("coverage") or {}
    story += [para("Sources et vérification", "h2"),para(f"Source : {(study.get('source') or {}).get('filename') or plan.get('filename') or 'dossier'}. Pages lues : {', '.join(map(str,coverage.get('analyzed_pages', []))) or 'saisie documentée'}. Pages en échec : {', '.join(map(str,coverage.get('failed_pages', []))) or 'aucune déclarée'}. Pages hors sélection : {', '.join(map(str,coverage.get('omitted_pages', []))) or 'aucune' }.","small"),
              para(f"Calcul {study.get('schema_version')} ; barème {study.get('catalog_version')}. Révision {study.get('revision','')[:16]}. Les totaux sont recalculés à partir des lignes, avec arrondi monétaire. L’interprétation complète du plan n’est pas certifiée par le contrôle arithmétique.","small"),
              para("Avant engagement : confirmer les données manquantes, les systèmes constructifs, le sol, le voisinage et les prix avec les intervenants. Délais, assurances et conditions commerciales restent à définir.","small")]
    def on_page(c, doc):
        c.saveState();c.setFillColor(CREAM);c.rect(0,0,A4[0],A4[1],fill=1,stroke=0)
        c.setFillColor(COPPER);c.rect(0,A4[1]-8,A4[0],8,fill=1,stroke=0)
        c.setFont("DossierBold",10);c.drawString(48,A4[1]-42,(branding.get("company_name") or "CONSTRUPILOT")[:55])
        c.setStrokeColor(LINE);c.line(48,43,A4[0]-48,43);c.setFont("Dossier",7.2);c.setFillColor(MUTED)
        c.drawString(48,29,f"{ref} | Estimation à valider | {date}");c.restoreState()
    pdf.build(story,onFirstPage=on_page,onLaterPages=on_page,canvasmaker=NumberedCanvas)
    return buf.getvalue()


def build_dossier_docx(plan, study, branding=None):
    """Same data/unknowns as PDF; no legacy fixed assumptions in Word exports."""
    from docx import Document
    from docx.shared import Pt, RGBColor
    doc=Document()
    doc.styles['Normal'].font.name='DejaVu Sans'
    doc.styles['Normal'].font.size=Pt(10)
    doc.add_heading('Devis estimatif gros œuvre',0)
    doc.add_paragraph((study.get('project') or {}).get('title') or plan.get('filename') or '')
    for lot in study['lots']:
        doc.add_heading(lot['lot'],1)
        table=doc.add_table(rows=1, cols=6)
        for c,v in zip(table.rows[0].cells,['Désignation','Qté','U','Statut','PU HT','Total HT']): c.text=v
        for r in lot['lines']:
            for c,v in zip(table.add_row().cells,[r['designation'],fmt(r['qty']),r['unit'],r['quantity_status'],fmt(r['unit_price']),fmt(r['total'])]): c.text=v
    for key,label in [('ht','HT'),('tva','TVA'),('ttc','TTC')]:doc.add_paragraph(f"{label} : {fmt(study['totals'][key],True)}")
    doc.add_heading('Métrés et sources',1)
    for r in study['items']:
        doc.add_heading(r['designation'],2)
        doc.add_paragraph(f"{r['expression']} = {fmt(r['qty'])} {r['unit']} [{r['quantity_status']}]")
        doc.add_paragraph(r.get('detail') or '')
        for p in r['sources']:doc.add_paragraph(f"{p['label']} : {fmt(p['value'])} {p['unit']} ; "+'; '.join(f"p. {x['page']} {x['text']}" for x in p['sources'])+" "+p.get('note',''))
    for title,key in [('Hypothèses','assumptions'),('Points à confirmer','warnings'),('Exclusions','exclusions')]:
        doc.add_heading(title,1)
        for x in study.get(key,[]):doc.add_paragraph(str(x))
    doc.add_paragraph('Prix budgétaires ou saisis par l’utilisateur. Estimation non contractuelle ; études d’exécution et prix entreprise à confirmer.')
    for name in ['Title','Heading 1','Heading 2']:doc.styles[name].font.color.rgb=RGBColor.from_string('925535')
    out=BytesIO();doc.save(out);return out.getvalue()
