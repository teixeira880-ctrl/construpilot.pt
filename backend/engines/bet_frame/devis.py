"""Métré et devis estimatif de l'ossature et des réseaux collectifs (Excel à formules : quantité × PU)."""
from __future__ import annotations

from .calc import area
from .lang import _
from .model import Frame

STEEL = 7850.0
# clé → (PU FR € HT, PU PT € s/IVA)
PRICES = {
    "terr": (18.0, 9.0), "fouille": (32.0, 16.0), "propre": (120.0, 85.0), "beton_fond": (165.0, 105.0), "beton_elev": (185.0, 115.0),
    "beton_dalle": (175.0, 110.0), "acier": (1.85, 1.25), "coff_pot": (42.0, 22.0), "coff_poutre": (48.0, 26.0), "coff_dalle": (34.0, 18.0),
    "coff_voile": (38.0, 20.0), "dallage": (55.0, 32.0), "escalier": (2600.0, 1500.0), "drain": (65.0, 35.0), "etanch": (38.0, 22.0),
    "col_eau": (55.0, 32.0), "chute": (70.0, 38.0), "col_elec": (85.0, 48.0), "branch_eau": (1450.0, 850.0), "tableau": (1350.0, 780.0),
    "compteur": (420.0, 250.0),
}


def build_lines(m: Frame, r: dict, e: dict) -> list[dict]:
    pt = m.country == "PT"
    P = {k: v[1 if pt else 0] for k, v in PRICES.items()}
    L = []

    def add(lot, des, unit, qty, pu):
        if qty and qty > 0:
            L.append({"lot": lot, "des": des, "unit": unit, "qty": round(qty, 2), "pu": pu})
    f = r["foundations"]
    bas = r.get("basement")
    lot1 = _("1. Terrassements et fondations", "1. Movimento de terras e fundações")
    if bas:
        add(lot1, _("Terrassement général en pleine masse du sous-sol", "Escavação geral da cave"), "m³", m.footprint * (bas["H"] + 0.3) * 1.08, P["terr"])
    add(lot1, _("Fouilles des semelles", "Escavação de sapatas"), "m³", sum((p["B"] + 0.4) ** 2 * (p["h"] + 0.3) for p in f["pads"]), P["fouille"])
    add(lot1, _("Béton de propreté e=10 cm", "Betão de limpeza e=10 cm"), "m³", sum((p["B"] + 0.1) ** 2 * 0.10 for p in f["pads"]), P["propre"])
    vf = sum(p["B"] ** 2 * p["h"] for p in f["pads"])
    add(lot1, _(f"Béton {r['mat']['concrete']} des semelles", f"Betão {r['mat']['concrete']} em sapatas"), "m³", vf, P["beton_fond"])
    kg = sum(2 * p["mesh"]["As"] * p["B"] * (p["B"] + 0.30) * STEEL for p in f["pads"])
    lt = 0.0
    pads = {p["node"]: p for p in f["pads"]}
    for (i, j), p in pads.items():
        for di, dj in ((1, 0), (0, 1)):
            q = pads.get((i + di, j + dj))
            if q:
                lt += abs(q["x"] - p["x"]) + abs(q["y"] - p["y"]) - (p["B"] + q["B"]) / 2
    tb, th = f["ties"]
    add(lot1, _("Béton des longrines", "Betão das vigas de fundação"), "m³", lt * tb * th, P["beton_fond"])
    kg += lt * (4 * area(16) * 1.1 + 2 * (tb + th) * area(8) / 0.20) * STEEL
    for c in f["core_rafts"]:
        v = (c["x1"] - c["x0"]) * (c["y1"] - c["y0"]) * c["h"]
        add(lot1, _(f"Radier du noyau {c['id']} e={c['h']:.2f}", f"Ensoleiramento do núcleo {c['id']} e={c['h']:.2f}"), "m³", v, P["beton_fond"])
        kg += v * 90
    if bas:
        per = 2 * (m.L + m.W)
        add(lot1, _("Semelle filante sous murs de sous-sol", "Sapata contínua dos muros de cave"), "m³", per * bas["footing"][0] * bas["footing"][1], P["beton_fond"])
        kg += per * 45
    add(lot1, _(f"Aciers {r['mat']['steel']} façonnés — fondations", f"Aço {r['mat']['steel']} moldado — fundações"), "kg", kg * 1.08, P["acier"])
    # ---------------------------------------------------------------- structure
    lot2 = _("2. Gros œuvre — ossature béton armé", "2. Estrutura de betão armado")
    vcol = kgcol = fcol = 0.0
    for c in r["columns"]:
        for sg in c["segs"]:
            H = next(l.height for l in m.levels if l.id == sg["level"])
            vcol += c["b"] * c["h"] * H
            fcol += 2 * (c["b"] + c["h"]) * H
            kgcol += (sg.get("As", 0) * (H + 0.6) + 2 * (c["b"] + c["h"]) * area(8 if sg.get("phi", 12) >= 20 else 6) * H / 0.15) * STEEL
    add(lot2, _("Béton des poteaux", "Betão em pilares"), "m³", vcol, P["beton_elev"])
    add(lot2, _("Coffrage des poteaux", "Cofragem de pilares"), "m²", fcol, P["coff_pot"])
    vb = fb = kgb = 0.0
    vs = fs = kgs = 0.0
    for fl in r["floors"]:
        s = fl["slab"]
        for bm in fl["beams"]:
            hw = max(bm["h"] - s.h, 0)
            for sp in bm["spans"]:
                vb += bm["b"] * hw * sp["L"]
                fb += (bm["b"] + 2 * hw) * sp["L"]
                kgb += (sp["bot"]["As"] * (sp["L"] + 0.5) + (sp["top_l"]["As"] + sp["top_r"]["As"]) * sp["L"] * 0.35
                        + 2 * (bm["b"] + bm["h"]) * area(8) * sp["L"] / sp["s"]) * STEEL
        core_a = sum((c["x1"] - c["x0"]) * (c["y1"] - c["y0"]) for c in m.cores)
        a_ = m.footprint - core_a
        vs += a_ * s.h
        fs += a_
        for p in fl["panels"]:
            pa = p["dx"] * p["dy"]
            top = sum(t["As"] for t in p["top"].values()) / 4
            kgs += (p["mesh_short"]["As"] + p["mesh_long"]["As"] + top * 0.6) * pa * 1.10 * STEEL
    add(lot2, _("Béton des poutres (retombées)", "Betão em vigas (abaixo da laje)"), "m³", vb, P["beton_elev"])
    add(lot2, _("Coffrage des poutres", "Cofragem de vigas"), "m²", fb, P["coff_poutre"])
    add(lot2, _("Béton des dalles pleines", "Betão em lajes maciças"), "m³", vs, P["beton_dalle"])
    add(lot2, _("Coffrage et étaiement des dalles", "Cofragem e escoramento de lajes"), "m²", fs, P["coff_dalle"])
    vbal = sum((bc["a1"] - bc["a0"]) * bc["depth"] * bc["h"] * max(len(bc["levels"]), 1) for bc in r.get("balconies") or [])
    add(lot2, _("Béton des balcons en console", "Betão em varandas em consola"), "m³", vbal, P["beton_dalle"])
    hcore = r["floors"][-1]["slab"].z - m.levels[0].z
    vcore = sum(2 * ((c["x1"] - c["x0"]) + (c["y1"] - c["y0"])) * c["wall"] * hcore for c in m.cores)
    add(lot2, _("Béton des voiles des noyaux", "Betão em paredes dos núcleos"), "m³", vcore, P["beton_elev"])
    add(lot2, _("Coffrage des voiles (2 faces)", "Cofragem de paredes (2 faces)"), "m²", sum(4 * ((c["x1"] - c["x0"]) + (c["y1"] - c["y0"])) * hcore for c in m.cores), P["coff_voile"])
    if m.cores:
        add(lot2, _("Escaliers BA (volée + palier) par étage", "Escadas BA (lanço + patamar) por piso"), "u", len(m.levels), P["escalier"])
    if bas:
        per = 2 * (m.L + m.W)
        vw = per * bas["H"] * bas["t"]
        add(lot2, _(f"Béton des murs de sous-sol e={bas['t']:.2f}", f"Betão em muros de cave e={bas['t']:.2f}"), "m³", vw, P["beton_elev"])
        add(lot2, _("Coffrage des murs de sous-sol", "Cofragem dos muros de cave"), "m²", 2 * per * bas["H"], P["coff_voile"])
        kgw = (bas["vert"]["As"] + bas["outer"]["As"] + 2 * bas["hor"]["As"]) * per * bas["H"] * 1.1 * STEEL
        add(lot2, _("Étanchéité + drainage des murs enterrés", "Impermeabilização + drenagem dos muros enterrados"), "m²", per * bas["H"], P["etanch"] + P["drain"] / max(bas["H"], 1))
    else:
        kgw = 0.0
    lv0 = m.levels[0]
    if lv0.on_grade:
        add(lot2, _(f"Dallage / dalle sur terre-plein {lv0.name} e=0,15 (treillis soudé)", f"Laje térrea {lv0.name} e=0,15 (malha eletrossoldada)"), "m²", m.footprint, P["dallage"])
    kg2 = kgcol + kgb + kgs + vbal * 110 + vcore * 80 + kgw
    add(lot2, _(f"Aciers {r['mat']['steel']} façonnés — ossature", f"Aço {r['mat']['steel']} moldado — estrutura"), "kg", kg2 * 1.05, P["acier"])
    # ---------------------------------------------------------------- réseaux collectifs
    lot3 = _("3. Réseaux collectifs (eau, évacuations, électricité)", "3. Redes coletivas (águas, esgotos, eletricidade)")
    H = r["floors"][-1]["slab"].z - m.levels[0].z
    for ri in e["risers"]:
        if not ri["sections"]:
            continue
        d = ri["sections"][0].get("d") or ri["sections"][0].get("sec")
        if ri["kind"] in ("AF", "ECS"):
            add(lot3, _(f"Colonne montante {ri['id']} Ø{d} (PEX/PPR, calorifugée)", f"Coluna montante {ri['id']} Ø{d} (PEX/PPR, isolada)"), "m", H + 3, P["col_eau"])
        elif ri["kind"] == "EU":
            add(lot3, _(f"Chute {ri['id']} PVC Ø{d} + ventilation primaire", f"Tubo de queda {ri['id']} PVC DN{d} + ventilação primária"), "m", H + 2, P["chute"])
        elif ri["kind"] == "ELEC":
            add(lot3, _(f"Colonne électrique {ri['id']} {d} mm² Cu (4 conducteurs)", f"Coluna montante elétrica {ri['id']} {d} mm² Cu (4 condutores)"), "m", H + 5, P["col_elec"])
    n = len(e["units"])
    add(lot3, _("Branchement eau froide / eau chaude / évacuations par fraction (depuis colonne)", "Ramal de águas / esgotos por fração (desde a coluna)"), "u", n, P["branch_eau"])
    add(lot3, _("Compteurs d'eau individuels + vanne", "Contadores de água individuais + válvula"), "u", n, P["compteur"])
    add(lot3, _("Tableau électrique de fraction + dérivation individuelle", "Quadro elétrico de fração + entrada individual"), "u", n, P["tableau"])
    return L


def totals(lines: list[dict], country: str) -> dict:
    ht = sum(x["qty"] * x["pu"] for x in lines)
    vat = 0.23 if country == "PT" else 0.20
    return {"ht": round(ht, 2), "vat": vat, "ttc": round(ht * (1 + vat), 2)}


def write_xlsx(path, m: Frame, lines: list[dict]):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    wb = Workbook()
    ws = wb.active
    ws.title = _("Devis", "Orçamento")
    pt = m.country == "PT"
    hdr = [_("Lot", "Capítulo"), _("Désignation", "Designação"), _("Unité", "Un."), _("Quantité", "Quantidade"),
           _("PU € HT", "PU € s/IVA"), _("Total € HT", "Total € s/IVA")]
    ws.append([m.project.get("name", "")])
    ws.append([_("Estimation — pré-dimensionnement (prix unitaires modifiables)", "Estimativa — pré-dimensionamento (preços unitários editáveis)")])
    ws.append([])
    ws.append(hdr)
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
    ws["A1"].font = Font(bold=True, size=13)
    row = 5
    first = row
    for x in lines:
        ws.append([x["lot"], x["des"], x["unit"], x["qty"], x["pu"], f"=D{row}*E{row}"])
        ws[f"E{row}"].number_format = "#,##0.00"
        ws[f"F{row}"].number_format = "#,##0.00"
        row += 1
    last = row - 1
    vat = 0.23 if pt else 0.20
    ws.append([])
    row += 1
    ws.append(["", _("TOTAL HT", "TOTAL s/IVA"), "", "", "", f"=SUM(F{first}:F{last})"])
    ws.append(["", _("TVA", "IVA") + f" {int(vat * 100)} %", "", "", "", f"=F{row}*{vat}"])
    ws.append(["", _("TOTAL TTC", "TOTAL c/IVA"), "", "", "", f"=F{row}+F{row + 1}"])
    for rr in (row, row + 1, row + 2):
        ws[f"B{rr}"].font = Font(bold=True)
        ws[f"F{rr}"].font = Font(bold=True)
        ws[f"F{rr}"].number_format = "#,##0.00"
    for col, w in zip("ABCDEF", (34, 70, 7, 12, 12, 15)):
        ws.column_dimensions[col].width = w
    for r_ in ws.iter_rows(min_row=5, max_row=last):
        r_[1].alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A5"
    wb.save(str(path))
