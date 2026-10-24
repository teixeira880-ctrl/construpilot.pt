"""Mapa de quantidades e caderno de encargos por especialidade a partir de um rascunho de orçamento ConstruPilot (Portugal).

Opção « Mapa de quantidades e caderno de encargos » do programa IA por opções. Determinístico: as linhas, quantidades e
preços são os do orçamento; a especialidade é deduzida da designação / descrição do artigo.
O caderno de encargos não redige nenhuma prescrição técnica: por especialidade, cita apenas referências gerais e
verificáveis (Eurocódigos NP EN, RTIEBT, DR 23/95, REH, especificações LNEC aplicáveis…) sem números de artigo, e lista
os trabalhos com as suas quantidades; as prescrições ficam « a completar pelo autor do caderno de encargos ».
"""
from __future__ import annotations

import io
import re
import unicodedata

A_COMPLETAR = "a completar pelo autor do caderno de encargos"
REF_GERAIS = ["Marcação CE e declarações de desempenho dos produtos (Regulamento (UE) n.º 305/2011)",
              "Fichas técnicas e documentos de aplicação dos fabricantes dos produtos escolhidos"]

# (especialidade, palavras-chave da designação / descrição, referências gerais)
LOTES = [
    ("Demolições e trabalhos preparatórios", ("demolic", "desmontagem", "remocao", "entulho"), []),
    ("Estrutura e alvenarias", ("parede", "alvenaria", "tijolo", "bloco", "betao", "laje", "sapata", "fundac", "lintel", "viga",
                                "pilar", "cinta", "muro", "estrutura", "armadura", "cofragem"),
     ["Eurocódigos estruturais NP EN 1990 a NP EN 1998, com os Anexos Nacionais", "Especificações LNEC aplicáveis"]),
    ("Divisórias, forros e tetos em gesso cartonado", ("divisoria", "pladur", "gesso cartonado", "forro", "teto", "tecto"),
     ["RRAE — Regulamento dos Requisitos Acústicos dos Edifícios (DL 129/2002), quando aplicável"]),
    ("Isolamentos térmicos", ("isolamento", "isolante", "la mineral", "capoto", "etics", "xps", "eps"),
     ["REH — Regulamento de Desempenho Energético dos Edifícios de Habitação (DL 101-D/2020)", "Especificações LNEC aplicáveis"]),
    ("Betonilhas e piso radiante", ("betonilha", "piso radiante", "autonivelante", "chauffage_sol"),
     ["NP EN 1264 — Sistemas de aquecimento e arrefecimento por superfícies radiantes (piso radiante)", "Especificações LNEC aplicáveis"]),
    ("Impermeabilizações e revestimentos cerâmicos", ("ceramic", "azulejo", "mosaico", "cola para azulejo", "betume de juntas",
                                                      "impermeabiliza"),
     ["Especificações LNEC aplicáveis"]),
    ("Pavimentos de madeira e rodapés", ("soalho", "parquet", "flutuante", "subcamada", "rodape"), []),
    ("Estuques e pinturas", ("pintura", "tinta", "primario", "estuque", "barramento"), ["Especificações LNEC aplicáveis"]),
    ("Carpintarias, caixilharias e escadas", ("escada", "porta", "janela", "caixilharia", "carpintaria", "estore"),
     ["REH — Regulamento de Desempenho Energético dos Edifícios de Habitação (DL 101-D/2020), para os vãos envidraçados"]),
    ("Instalações elétricas", ("eletric", "electric", "tomada", "iluminacao", "ponto de luz", "quadro eletrico", "circuito"),
     ["RTIEBT — Regras Técnicas das Instalações Elétricas de Baixa Tensão"]),
    ("Abastecimento de água e drenagem de águas residuais", ("canaliza", "agua fria", "agua quente", "esgoto", "drenagem",
                                                             "sanitari", "loica", "autoclismo"),
     ["DR 23/95 — Regulamento Geral dos Sistemas Públicos e Prediais de Distribuição de Água e de Drenagem de Águas Residuais"]),
]
DIVERSOS = ("Diversos", [])
# Ordem de deteção: das palavras mais específicas às mais genéricas (« Pintura — paredes e tetos » é pintura e não teto
# falso nem estrutura; « Tubo PEX + placas isolantes » de piso radiante não é isolamento; « Azulejo (parede) » não é alvenaria).
DETECAO = ["Betonilhas e piso radiante", "Impermeabilizações e revestimentos cerâmicos", "Pavimentos de madeira e rodapés",
           "Estuques e pinturas", "Instalações elétricas", "Abastecimento de água e drenagem de águas residuais",
           "Divisórias, forros e tetos em gesso cartonado", "Isolamentos térmicos", "Carpintarias, caixilharias e escadas",
           "Demolições e trabalhos preparatórios", "Estrutura e alvenarias"]


def _fold(s: str) -> str:
    return unicodedata.normalize("NFKD", s or "").encode("ascii", "ignore").decode().lower()


def lote_de(line: dict) -> tuple[str, list[str]]:
    txt = _fold(" ".join(str(line.get(k) or "") for k in ("designation", "description")))
    for name, keys, refs in sorted(LOTES, key=lambda l: DETECAO.index(l[0])):
        if any(re.search(r"\b" + re.escape(k), txt) for k in keys):
            return name, refs
    return DIVERSOS


def por_lote(lines: list[dict]) -> list[dict]:
    """Agrupa as linhas orçamentáveis do orçamento por especialidade, pela ordem do quadro LOTES."""
    order = [n for n, *_ in LOTES] + [DIVERSOS[0]]
    lotes: dict[str, dict] = {}
    for ln in lines:
        if ln.get("kind", "line") != "line" or not ln.get("designation"):
            continue
        name, refs = lote_de(ln)
        lote = lotes.setdefault(name, {"lote": name, "refs": refs, "linhas": []})
        qty = float(ln.get("qty") or 0)
        pu = float(ln.get("unit_price_ht") or 0)
        falta = bool(ln.get("price_missing"))
        lote["linhas"].append({"designacao": ln["designation"], "detalhe": ln.get("description") or "", "unidade": ln.get("unit") or "",
                               "quantidade": round(qty, 2), "pu": None if falta else round(pu, 2),
                               "montante": None if falta else round(qty * pu, 2)})
    out = [lotes[n] for n in order if n in lotes]
    for i, lote in enumerate(out, 1):
        lote["numero"] = i
        lote["total"] = round(sum(l["montante"] or 0 for l in lote["linhas"]), 2)
        lote["por_orcamentar"] = sum(1 for l in lote["linhas"] if l["pu"] is None)
    return out


def _num(x: float) -> str:
    return f"{x:g}".replace(".", ",")


def _unidade(u: str) -> str:
    return {"m2": "m²", "m3": "m³"}.get(u, u)


def mapa_xlsx(lotes: list[dict], titulo: str) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Mapa de quantidades"
    ws.append([f"Mapa de quantidades — {titulo}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append(["Quantidades e preços retirados do rascunho de orçamento ConstruPilot. Linhas « por orçamentar »: preço a completar. "
               "Valores sem IVA."])
    ws.append([])
    ws.append(["N.º", "Designação", "Unid.", "Quantidade", "Preço unitário (€)", "Total (€)"])
    for c in ws[4]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E4A")
    total = 0.0
    for lote in lotes:
        ws.append([f"{lote['numero']}", f"{lote['numero']}. {lote['lote']}"])
        ws.cell(ws.max_row, 2).font = Font(bold=True)
        for j, l in enumerate(lote["linhas"], 1):
            ws.append([f"{lote['numero']}.{j}", l["designacao"], _unidade(l["unidade"]), l["quantidade"],
                       l["pu"] if l["pu"] is not None else "por orçamentar", l["montante"]])
        ws.append(["", f"Subtotal {lote['numero']}. {lote['lote']}", "", "", "", lote["total"]])
        ws.cell(ws.max_row, 2).font = Font(italic=True)
        total += lote["total"]
    ws.append([])
    ws.append(["", "TOTAL (sem IVA)", "", "", "", round(total, 2)])
    ws.cell(ws.max_row, 2).font = Font(bold=True)
    for row in ws.iter_rows(min_row=5, min_col=5, max_col=6):
        for c in row:
            if isinstance(c.value, (int, float)):
                c.number_format = "#,##0.00"
    for col, w in zip("ABCDEF", (8, 62, 8, 12, 18, 15)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def caderno_docx(lotes: list[dict], titulo: str) -> bytes:
    from docx import Document
    doc = Document()
    doc.core_properties.language = "pt-PT"
    doc.core_properties.title = f"Caderno de encargos — {titulo}"
    doc.add_heading(f"Caderno de encargos — {titulo}", 0)
    doc.add_paragraph("Estrutura de caderno de encargos (cláusulas técnicas) elaborada a partir do rascunho de orçamento ConstruPilot. "
                      "As referências normativas e regulamentares são citadas de forma geral por especialidade; as prescrições técnicas "
                      f"(materiais, execução, tolerâncias, ensaios e critérios de medição) estão por redigir — {A_COMPLETAR} — e devem ser "
                      "validadas antes de qualquer divulgação.")
    for lote in lotes:
        doc.add_heading(f"{lote['numero']}. {lote['lote']}", 1)
        doc.add_heading(f"{lote['numero']}.1 Documentos de referência", 2)
        for r in lote["refs"] + REF_GERAIS:
            doc.add_paragraph(r, style="List Bullet")
        doc.add_paragraph(f"Outras normas e regulamentos aplicáveis: {A_COMPLETAR}.")
        doc.add_heading(f"{lote['numero']}.2 Descrição dos trabalhos", 2)
        for j, l in enumerate(lote["linhas"], 1):
            doc.add_heading(f"{lote['numero']}.2.{j} {l['designacao']}", 3)
            doc.add_paragraph(f"Quantidade indicativa: {_num(l['quantidade'])} {_unidade(l['unidade'])} (medição sobre plantas).")
            doc.add_paragraph(f"Prescrições (materiais, execução, tolerâncias, critério de medição): {A_COMPLETAR}.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
