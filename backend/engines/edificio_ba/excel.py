"""Livro Excel com fórmulas (PT / FR) : resumo, medições, orçamento (P.U. editáveis), paredes, muros, sapatas, lajes, vãos.
Azul = valor medido / editável ; preto = fórmula ; verde = ligação a outra folha."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import estrutura as ES

HDR = PatternFill("solid", fgColor="1F4E46")
TOT = PatternFill("solid", fgColor="E8F3F0")
INP = Font(name="Arial", size=10, color="0000FF")
TXT = Font(name="Arial", size=10)
LNK = Font(name="Arial", size=10, color="008000")
BLD = Font(name="Arial", size=10, bold=True)
thin = Side(style="thin", color="BBBBBB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
N2 = "#,##0.00"
N0 = "#,##0"


def _sheet(wb, S, title, head, widths, note=None):
    ws = wb.create_sheet(title[:31])
    ws["A1"] = title
    ws["A1"].font = Font(name="Arial", size=13, bold=True, color="1F4E46")
    ws["A2"] = note or S.t(f"{S.name}. Azul = dado medido / editável ; preto = fórmula ; verde = ligação.",
                           f"{S.name}. Bleu = donnée mesurée / modifiable ; noir = formule ; vert = lien.")
    ws["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")
    for j, (h, w) in enumerate(zip(head, widths), 1):
        c = ws.cell(row=4, column=j, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = BOX
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 30
    ws.freeze_panes = "A5"
    return ws


def _put(ws, r, vals, fmt=None, inputs=()):
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=j, value=v)
        c.font = INP if (j in inputs and not (isinstance(v, str) and v.startswith("="))) else TXT
        c.border = BOX
        if fmt and j in fmt:
            c.number_format = fmt[j]


def _total(ws, r, label, cols, first, last, ncols):
    ws.cell(row=r, column=1, value=label).font = BLD
    for j in range(1, ncols + 1):
        ws.cell(row=r, column=j).fill = TOT
        ws.cell(row=r, column=j).border = BOX
    for j in cols:
        L = get_column_letter(j)
        c = ws.cell(row=r, column=j, value=f"=SUM({L}{first}:{L}{last})")
        c.font = BLD
        c.number_format = N2


def build(S, R, Q, O, path):
    T = S.t
    wb = Workbook()
    wb.remove(wb.active)
    res = wb.create_sheet(T("Resumo", "Synthese"))
    # ---------------------------------------------------------------- medições
    wm = _sheet(wb, S, T("Medicoes", "Metre"), ["Art.", T("Elemento", "Élément"), T("Dimensões", "Dimensions"), T("Quant.", "Quant."), T("Un.", "U."),
                                                T("Betão (m³)", "Béton (m³)"), T("Cofragem (m²)", "Coffrage (m²)"), T("Aço (kg)", "Acier (kg)")],
                [8, 50, 34, 12, 7, 13, 14, 12])
    r = 5
    for x in Q["lines"]:
        _put(wm, r, [x["art"], x["el"], x["dims"], x["q"], x["un"], x["vol"] or None, x["cof"] or None, x["aco"] or None], {4: N2, 6: N2, 7: N2, 8: N0},
             inputs=(4, 6, 7, 8))
        r += 1
    rmt = r
    _total(wm, r, "TOTAL", [6, 7, 8], 5, r - 1, 8)
    # ---------------------------------------------------------------- orçamento
    wo = _sheet(wb, S, T("Orcamento", "Estimation"), ["Art.", T("Designação", "Désignation"), T("Un.", "U."), T("Quant.", "Quant."), "P.U. (€)", "Total (€)"],
                [8, 70, 7, 13, 12, 15])
    r = 5
    first = r
    for a in O["arts"]:
        _put(wo, r, [a["art"], a["desc"], a["un"], a["q"], a["pu"], f"=D{r}*E{r}"], {4: N2, 5: N2, 6: N0}, inputs=(4, 5))
        r += 1
    rs = r
    _total(wo, r, T("TOTAL ESTRUTURA (cap. 1-6)", "TOTAL STRUCTURE (lots 1-6)"), [6], first, r - 1, 6)
    r += 2
    fp = r
    for p in O["par"]:
        _put(wo, r, [p["cap"], p["desc"], p["un"], p["q"], p["pu"], f"=D{r}*E{r}"], {4: N2, 5: N2, 6: N0}, inputs=(4, 5))
        r += 1
    rp = r
    _total(wo, r, T("TOTAL RESTANTES TRABALHOS (cap. 7-17)", "TOTAL AUTRES LOTS (7-17)"), [6], fp, r - 1, 6)
    r += 1
    _put(wo, r, ["", T("Estaleiro, segurança e ensaios (taxa)", "Installation de chantier, sécurité, essais (taux)"), "%", "", 0.03, f"=E{r}*(F{rs}+F{rp})"],
         {5: "0.0%", 6: N0}, inputs=(5,))
    re_ = r
    r += 1
    wo.cell(row=r, column=2, value=T("TOTAL SEM IVA", "TOTAL HT")).font = BLD
    wo.cell(row=r, column=6, value=f"=F{rs}+F{rp}+F{re_}").number_format = N0
    rt = r
    r += 1
    _put(wo, r, ["", S.prof["vat_note"], "%", "", S.prof["vat"], f"=E{r}*F{rt}"], {5: "0.0%", 6: N0}, inputs=(5,))
    rv = r
    r += 1
    wo.cell(row=r, column=2, value=T("TOTAL COM IVA", "TOTAL TTC")).font = BLD
    wo.cell(row=r, column=6, value=f"=F{rt}+F{rv}").number_format = N0
    rttc = r
    r += 1
    wo.cell(row=r, column=2, value=T("Área bruta (m²)", "Surface brute (m²)"))
    c = wo.cell(row=r, column=4, value=round(Q["areas"]["bruta"], 1))
    c.font = INP
    wo.cell(row=r, column=6, value=f"=F{rt}/D{r}").number_format = N0
    wo.cell(row=r, column=5, value="€/m²")
    rm2 = r
    # ---------------------------------------------------------------- paredes
    qd = Q["quant"]
    ch = qd["clear_h"]
    wp = _sheet(wb, S, T("Paredes", "Murs"), [T("Piso", "Niveau"), T("Perímetro fachada (ml)", "Périmètre façade (ml)"), T("Altura (m)", "Hauteur (m)"),
                                             T("Fachada bruta (m²)", "Façade brute (m²)"), T("Vãos (m²)", "Baies (m²)"), T("Betão empenas (ml)", "Béton pignons (ml)"),
                                             T("Alvenaria exterior (m²)", "Maçonnerie ext. (m²)"), T("Divisórias (ml)", "Cloisons (ml)"), T("Divisórias (m²)", "Cloisons (m²)"),
                                             T("Separação (ml)", "Séparatifs (ml)"), T("Separação (m²)", "Séparatifs (m²)"), T("Escada/hall (ml)", "Cage/hall (ml)"),
                                             T("Escada/hall (m²)", "Cage/hall (m²)")], [12, 14, 10, 14, 11, 13, 16, 12, 13, 13, 13, 13, 13])
    r = 5
    for p in qd["paredes"]:
        _put(wp, r, [p["piso"], p["L_fach"], p["h"], f"=B{r}*C{r}", round(p["a_vaos"], 2), p["L_bet"], f"=D{r}-E{r}-F{r}*C{r}", p["div"], f"=H{r}*{ch}",
                     p["sep"], f"=J{r}*{ch}", p["enu"], f"=L{r}*{ch}"], {k: N2 for k in range(2, 14)}, inputs=(2, 3, 5, 6, 8, 10, 12))
        r += 1
    rpw = r
    _total(wp, r, "TOTAL", [4, 5, 7, 9, 11, 13], 5, r - 1, 13)
    # ---------------------------------------------------------------- muros e betão
    wb_ = _sheet(wb, S, T("Muros e betao", "Voiles"), [T("Elemento", "Élément"), T("Comprimento (ml)", "Longueur (ml)"), T("Altura (m)", "Hauteur (m)"),
                                                     T("Espessura (m)", "Épaisseur (m)"), T("Área 1 face (m²)", "Surface 1 face (m²)"), T("Cofragem 2 faces (m²)", "Coffrage 2 faces (m²)"),
                                                     T("Volume (m³)", "Volume (m³)"), T("Terras (m)", "Terres (m)"), T("Impermeabilização (m²)", "Étanchéité (m²)")],
                 [34, 15, 11, 12, 15, 16, 12, 12, 17])
    r = 5
    for w in R["bwalls"]:
        _put(wb_, r, [w["name"], w["L"], S.H, w["t"], f"=B{r}*C{r}", f"=2*E{r}", f"=E{r}*D{r}", w["ht"], f"=B{r}*H{r}"], {k: N2 for k in range(2, 10)}, inputs=(2, 3, 4, 8))
        r += 1
    rmu = r
    _total(wb_, r, T("TOTAL muros de cave", "TOTAL voiles enterrés"), [2, 5, 6, 7, 9], 5, r - 1, 9)
    r += 2
    r0 = r
    ns = len(ES.storeys(S))
    for c_ in S.cores:
        cs = ES.core_section(c_)
        _put(wb_, r, [c_["id"], round(cs["A"] / c_["t"], 2), ns * S.H, c_["t"], f"=B{r}*C{r}", f"=2*E{r}", f"=E{r}*D{r}", 0, 0], {k: N2 for k in range(2, 10)}, inputs=(2, 3, 4))
        r += 1
    for w in S.walls:
        _put(wb_, r, [w["id"] + (" *" if w["proposed"] else ""), w["Lw"], ns * S.H, w["t"], f"=B{r}*C{r}", f"=2*E{r}", f"=E{r}*D{r}", 0, 0], {k: N2 for k in range(2, 10)},
             inputs=(2, 3, 4))
        r += 1
    rbe = r
    _total(wb_, r, T("TOTAL paredes de betão", "TOTAL voiles béton"), [5, 6, 7], r0, r - 1, 9)
    # ---------------------------------------------------------------- sapatas
    ws = _sheet(wb, S, T("Sapatas", "Semelles"), [T("Elemento", "Élément"), "N", T("Comprimento (ml)", "Longueur (ml)"), T("Largura (m)", "Largeur (m)"),
                                                  T("Altura (m)", "Hauteur (m)"), T("Área (m²)", "Surface (m²)"), T("Volume (m³)", "Volume (m³)")], [36, 7, 16, 12, 11, 12, 12])
    r = 5
    for w in R["bwalls"]:
        _put(ws, r, [T(f"Sapata contínua sob {w['name']}", f"Semelle filante sous {w['name']}"), 1, w["L"], 1.0, 0.40, f"=B{r}*C{r}*D{r}", f"=F{r}*E{r}"],
             {k: N2 for k in range(3, 8)}, inputs=(2, 3, 4, 5))
        r += 1
    rsc = r
    _total(ws, r, T("TOTAL sapatas contínuas", "TOTAL semelles filantes"), [3, 6, 7], 5, r - 1, 7)
    r += 2
    r1 = r
    for x in R["foots"]:
        _put(ws, r, [f"S{x['id'].lstrip('P')}", 1, round(x["y1"] - x["y0"], 2), round(x["x1"] - x["x0"], 2), x["h"], f"=B{r}*C{r}*D{r}", f"=F{r}*E{r}"],
             {k: N2 for k in range(3, 8)}, inputs=(2, 3, 4, 5))
        r += 1
    for x in R["fnd"]:
        _put(ws, r, [x["id"] + " (" + ", ".join(x["members"]) + ")", 1, x["L"], x["B"], x["h"], f"=B{r}*C{r}*D{r}", f"=F{r}*E{r}"], {k: N2 for k in range(3, 8)},
             inputs=(2, 3, 4, 5))
        r += 1
    rsi = r
    _total(ws, r, T("TOTAL sapatas isoladas + ensoleiramentos", "TOTAL semelles isolées + radiers"), [2, 6, 7], r1, r - 1, 7)
    # ---------------------------------------------------------------- lajes
    wl = _sheet(wb, S, T("Lajes", "Dalles"), [T("Elemento", "Élément"), T("Área (m²)", "Surface (m²)"), T("Espessura (m)", "Épaisseur (m)"), T("Volume (m³)", "Volume (m³)")],
                [56, 14, 14, 14])
    r = 5
    for x in [x for x in Q["lines"] if x["art"] in ("5.1", "5.2", "5.3", "5.6")]:
        h = round(x["vol"] / x["q"], 3) if x["q"] and x["vol"] else 0.15
        _put(wl, r, [x["el"], x["q"], h, f"=B{r}*C{r}"], {2: N2, 3: "0.000", 4: N2}, inputs=(2, 3))
        r += 1
    rl = r
    _total(wl, r, T("TOTAL lajes", "TOTAL dalles"), [2, 4], 5, r - 1, 4)
    # ---------------------------------------------------------------- vãos
    wj = _sheet(wb, S, T("Janelas", "Fenetres"), [T("Tipo", "Type"), T("Descrição", "Description"), T("Largura (m)", "Largeur (m)"), T("Altura (m)", "Hauteur (m)"),
                                                  T("Área unit. (m²)", "Surface unit. (m²)")] + [S.lname(k) for k in range(S.n)] + ["Total", T("Área total (m²)", "Surface totale (m²)")],
                [8, 40, 11, 10, 13] + [9] * S.n + [9, 14])
    r = 5
    ot = S.opening_types
    vaos = Q["quant"]["vaos"]
    nc = 5 + S.n
    for k, v in vaos.items():
        d = ot.get(k, {})
        vals = [k, d.get("label", ""), float(d.get("width") or 0), float(d.get("height") or 0), f"=C{r}*D{r}"] + [v["by_level"].get(l_, 0) for l_ in range(S.n)]
        cl, cr_ = get_column_letter(6), get_column_letter(nc)
        vals += [f"=SUM({cl}{r}:{cr_}{r})", f"=E{r}*{get_column_letter(nc + 1)}{r}"]
        _put(wj, r, vals, {3: N2, 4: N2, 5: N2, nc + 2: N2}, inputs=tuple([3, 4] + list(range(6, nc + 1))))
        r += 1
    rj = r
    _total(wj, r, "TOTAL", list(range(6, nc + 3)), 5, r - 1, nc + 2)
    # ---------------------------------------------------------------- resumo
    res["A1"] = T("Quadro de quantidades e orçamento — ", "Quantitatifs et estimation — ") + S.name
    res["A1"].font = Font(name="Arial", size=14, bold=True, color="1F4E46")
    res["A2"] = T("Valores ligados às folhas de detalhe (fórmulas).", "Valeurs liées aux feuilles de détail (formules).")
    res["A2"].font = Font(name="Arial", size=9, italic=True, color="666666")
    for j, (h, w) in enumerate(zip([T("Grandeza", "Grandeur"), T("Quantidade", "Quantité"), T("Unidade", "Unité")], [60, 18, 10]), 1):
        c = res.cell(row=4, column=j, value=h)
        c.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        c.fill = HDR
        res.column_dimensions[get_column_letter(j)].width = w
    sm, so, sp_, sb, ss, sl, sj = (f"'{x.title}'" for x in (wm, wo, wp, wb_, ws, wl, wj))
    rows = [
        (T("ESTRUTURA", "STRUCTURE"), None, None),
        (T("Betão estrutural", "Béton de structure"), f"={sm}!F{rmt}", "m³"), (T("Aço", "Acier"), f"={sm}!H{rmt}", "kg"),
        (T("Cofragem", "Coffrage"), f"={sm}!G{rmt}", "m²"),
        (T("PAREDES E MUROS", "MURS ET VOILES"), None, None),
        (T("Fachadas — alvenaria exterior (líquido de vãos)", "Façades — maçonnerie (déduction baies)"), f"={sp_}!G{rpw}", "m²"),
        (T("Paredes divisórias interiores", "Cloisons"), f"={sp_}!I{rpw}", "m²"), (T("Paredes de separação entre frações", "Murs séparatifs"), f"={sp_}!K{rpw}", "m²"),
        (T("Paredes da caixa de escada / hall", "Murs de cage / hall"), f"={sp_}!M{rpw}", "m²"), (T("Muros de cave (1 face)", "Voiles enterrés (1 face)"), f"={sb}!E{rmu}", "m²"),
        (T("Muros de cave — impermeabilização", "Voiles enterrés — étanchéité"), f"={sb}!I{rmu}", "m²"),
        (T("Paredes de betão (1 face)", "Voiles béton (1 face)"), f"={sb}!E{rbe}", "m²"),
        (T("SAPATAS", "SEMELLES"), None, None), (T("Sapatas contínuas sob muros", "Semelles filantes"), f"={ss}!C{rsc}", "ml"),
        (T("Sapatas isoladas + ensoleiramentos", "Semelles isolées + radiers"), f"={ss}!B{rsi}", "un"),
        (T("LAJES", "DALLES"), None, None), (T("Lajes, varandas e pavimento térreo", "Dalles, balcons et dallage"), f"={sl}!B{rl}", "m²"),
        (T("JANELAS E VÃOS", "FENÊTRES ET BAIES"), None, None), (T("Número de vãos exteriores", "Nombre de baies"), f"={sj}!{get_column_letter(nc + 1)}{rj}", "un"),
        (T("Área de vãos", "Surface de baies"), f"={sj}!{get_column_letter(nc + 2)}{rj}", "m²"),
        (T("ORÇAMENTO", "ESTIMATION"), None, None), (T("Estrutura", "Structure"), f"={so}!F{rs}", "€"), (T("Restantes trabalhos", "Autres lots"), f"={so}!F{rp}", "€"),
        (T("Total sem IVA", "Total HT"), f"={so}!F{rt}", "€"), (T("Total com IVA", "Total TTC"), f"={so}!F{rttc}", "€"), (T("Custo por m² bruto", "Coût au m² brut"), f"={so}!F{rm2}", "€/m²"),
    ]
    r = 5
    for lab, v, u in rows:
        a = res.cell(row=r, column=1, value=lab)
        if v is None:
            a.font = Font(name="Arial", size=10, bold=True, color="1F4E46")
            for j in (1, 2, 3):
                res.cell(row=r, column=j).fill = TOT
        else:
            a.font = TXT
            c = res.cell(row=r, column=2, value=v)
            c.font = LNK
            c.number_format = N0 if u in ("un", "€", "€/m²", "kg") else N2
            res.cell(row=r, column=3, value=u).font = TXT
        r += 1
    wb.save(path)
    return path
