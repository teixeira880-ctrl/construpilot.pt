"""Import tolérant d'une bibliothèque d'ouvrages (CSV ou Excel).

Reconnaît des en-têtes variés (FR, avec/sans accents) : catégorie, désignation,
unité, TVA, et prix par niveau (bas / médian / haut) — soit directement, soit via
matériaux + main d'œuvre par niveau. Retourne des lignes normalisées.
"""
from __future__ import annotations

import csv
import io
import re
import unicodedata


def _norm(h: str) -> str:
    s = (h or "").replace("œ", "oe").replace("æ", "ae").replace("’", " ").replace("'", " ")
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c))
    s = re.sub(r"[^a-z0-9]+", " ", s.lower()).strip()
    return re.sub(r"\s+", " ", s)


def _classify(header: str) -> str | None:
    """Associe un en-tête (quelconque) à un champ interne, par mots-clés."""
    n = _norm(header)
    if not n:
        return None
    # champs simples
    if any(k in n for k in ("designation", "ouvrage", "libelle", "intitule", "poste")) or n == "nom":
        return "designation"
    if "categorie" in n or "corps" in n or n in ("lot", "famille", "chapitre"):
        return "category"
    if n in ("unite", "u", "unit"):
        return "unit"
    if "tva" in n or n == "taux":
        return "vat"
    if "source" in n:
        return "source"
    if "date" in n:
        return "ref_date"
    if "idf" in n or "ile de france" in n:
        return "coef_idf"
    if "grandes villes" in n or ("coef" in n and "ville" in n):
        return "coef_ville"
    if "rural" in n:
        return "coef_rural"
    # niveau de prix
    level = None
    if any(k in n for k in ("bas", "mini", "min")):
        level = "low"
    elif any(k in n for k in ("haut", "maxi", "max")):
        level = "high"
    elif any(k in n for k in ("moyen", "median", "medium", "moy")):
        level = "mid"
    if level:
        if any(k in n for k in ("materiau", "materiel", "fourniture")) or n.startswith("mat"):
            return f"mat_{level}"
        if any(k in n for k in ("main", "oeuvre", "pose", "labor")) or n.startswith("mo"):
            return f"mo_{level}"
        return f"p_{level}"
    if "total ht" in n or n in ("prix", "prix ht", "pu ht", "prix unitaire", "prix vente"):
        return "p_mid"
    return None


def _map_headers(headers: list[str]) -> dict[str, int]:
    """field -> index de colonne (premier match)."""
    colmap: dict[str, int] = {}
    for i, h in enumerate(headers):
        field = _classify(h)
        if field and field not in colmap:
            colmap[field] = i
    return colmap


def _num(v) -> float | None:
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip().replace("\xa0", "").replace(" ", "").replace("€", "")
    if not s:
        return None
    s = s.replace(",", ".")
    # garder uniquement le dernier point comme séparateur décimal
    if s.count(".") > 1:
        parts = s.split(".")
        s = "".join(parts[:-1]) + "." + parts[-1]
    try:
        return float(re.sub(r"[^0-9.\-]", "", s))
    except ValueError:
        return None


def _level(direct, mat, mo) -> float | None:
    if direct is not None:
        return direct
    parts = [p for p in (mat, mo) if p is not None]
    return round(sum(parts), 2) if parts else None


def _build_row(cells: list, colmap: dict[str, int]) -> dict | None:
    def g(field):
        i = colmap.get(field)
        return cells[i] if (i is not None and i < len(cells)) else None

    designation = (str(g("designation")).strip() if g("designation") is not None else "")
    if not designation or designation.lower() in ("none", "nan"):
        return None

    low = _level(_num(g("p_low")), _num(g("mat_low")), _num(g("mo_low")))
    mid = _level(_num(g("p_mid")), _num(g("mat_mid")), _num(g("mo_mid")))
    high = _level(_num(g("p_high")), _num(g("mat_high")), _num(g("mo_high")))
    avail = [x for x in (low, mid, high) if x is not None]
    if not avail:
        return None
    if mid is None:
        mid = round(sum(avail) / len(avail), 2)
    if low is None:
        low = min(avail)
    if high is None:
        high = max(avail)

    vat = _num(g("vat"))
    if vat is not None and vat < 1:  # 0.1 -> 10
        vat = round(vat * 100, 1)
    unit = (str(g("unit")).strip() if g("unit") is not None else "") or "u"
    return {
        "category": (str(g("category")).strip() if g("category") is not None else "") or "Outro",
        "designation": designation[:220],
        "unit": unit[:16],
        "vat_rate": vat,
        "price_low": round(low, 2), "price_mid": round(mid, 2), "price_high": round(high, 2),
        "mat_low": _num(g("mat_low")), "mat_mid": _num(g("mat_mid")), "mat_high": _num(g("mat_high")),
        "labor_low": _num(g("mo_low")), "labor_mid": _num(g("mo_mid")), "labor_high": _num(g("mo_high")),
        "coef_idf": _num(g("coef_idf")), "coef_ville": _num(g("coef_ville")), "coef_rural": _num(g("coef_rural")),
        "source": (str(g("source")).strip() if g("source") is not None else "")[:120] or None,
        "ref_date": (str(g("ref_date")).strip() if g("ref_date") is not None else "")[:40] or None,
    }


def _parse_csv(content: bytes) -> tuple[list[list], list[str]]:
    try:
        text = content.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = content.decode("latin-1")
    sample = text[:4096]
    delim = ";" if sample.count(";") >= sample.count(",") else ","
    reader = csv.reader(io.StringIO(text), delimiter=delim)
    rows = [r for r in reader if any((c or "").strip() for c in r)]
    if not rows:
        return [], []
    return rows[1:], rows[0]


def _score_header(headers: list[str]) -> int:
    cm = _map_headers(headers)
    score = 0
    if "designation" in cm:
        score += 3
    if any(k in cm for k in ("p_low", "p_mid", "p_high", "mat_mid", "mo_mid", "mat_low")):
        score += 2
    return score


def _parse_xlsx(content: bytes) -> tuple[list[list], list[str]]:
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    best = None  # (score, data_rows, headers)
    skip = ("parametre", "tva", "taux", "source", "mapping", "note", "readme", "aide")
    for ws in wb.worksheets:
        name = _norm(ws.title)
        rows = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        rows = [r for r in rows if any(str(c).strip() for c in r if c is not None)]
        if len(rows) < 2:
            continue
        # cherche la ligne d'en-tête dans les 5 premières lignes
        for hi in range(min(5, len(rows))):
            headers = [str(c) if c is not None else "" for c in rows[hi]]
            sc = _score_header(headers)
            if name.split(" ")[0] in skip:
                sc -= 2
            if sc >= 3 and (best is None or sc > best[0]):
                best = (sc, rows[hi + 1:], headers)
    wb.close()
    if not best:
        return [], []
    return best[1], best[2]


def parse_shared_library(content: bytes, filename: str) -> dict:
    """Retourne {'rows': [...normalisées], 'columns': [en-têtes détectés], 'total_lines': n}."""
    fn = (filename or "").lower()
    if fn.endswith((".xlsx", ".xlsm", ".xls")):
        data_rows, headers = _parse_xlsx(content)
    else:
        data_rows, headers = _parse_csv(content)
    if not headers:
        return {"rows": [], "columns": [], "total_lines": 0}
    colmap = _map_headers(headers)
    if "designation" not in colmap:
        raise ValueError(
            "Coluna \'Designação\' não encontrada. Cabeçalhos detetados: "
            + ", ".join([h for h in headers if h][:30]))
    out = []
    for cells in data_rows:
        row = _build_row(list(cells), colmap)
        if row:
            out.append(row)
    detected = [headers[i] for f, i in sorted(colmap.items(), key=lambda x: x[1]) if i < len(headers)]
    return {"rows": out, "columns": detected, "total_lines": len(data_rows)}


CSV_TEMPLATE = (
    "Categorie;Designation;Unite;Prix_bas;Prix_moyen;Prix_haut;TVA;Source;Date_reference\nEstrutura;Parede em bloco de betão 20 cm;m2;45;55;70;23;Mercado 2026;2026-01\nCanalizacao;Duche de nivel;ens;750;950;1300;6;Mercado 2026;2026-01\nCaixilharia;Janela PVC vidro duplo;u;350;450;600;6;Mercado 2026;2026-01\n# Variante detalhada aceite: substitua Prix_* por Materiaux_bas;Materiaux_moyen;Materiaux_haut;MO_bas;MO_moyen;MO_haut\n"
)


# ========= Format « ConstruPilot FR – Bibliothèque de prix » (gammes Économique/Standard/Premium) =========

def _f(v):
    try:
        return round(float(str(v).replace(",", ".").replace("€", "").strip()), 2)
    except Exception:  # noqa: BLE001
        return None


def _hdr_idx(headers: list) -> dict:
    return {str(h).strip().lower(): i for i, h in enumerate(headers) if h is not None}


def parse_constructpilot_gammes(content: bytes) -> list[dict]:
    """Lit BIBLIOTHEQUE_PRIX (3 gammes → bas/moyen/haut) + PISCINES_CONSTRUCTION. Retourne des items shared_library."""
    import io
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    items: dict[tuple, dict] = {}
    if "BIBLIOTHEQUE_PRIX" in wb.sheetnames:
        rows = list(wb["BIBLIOTHEQUE_PRIX"].iter_rows(values_only=True))
        h = _hdr_idx(rows[0])
        gi = lambda r, k: r[h[k]] if k in h and h[k] < len(r) else None  # noqa: E731
        for r in rows[1:]:
            desig = str(gi(r, "designação do artigo") or "").strip()
            if not desig:
                continue
            cat = str(gi(r, "capítulo / família") or "Outro").strip()
            unit = str(gi(r, "unité") or "u").strip()
            key = (cat.lower(), desig.lower(), unit.lower())
            it = items.setdefault(key, {
                "category": cat[:80], "designation": desig[:220], "unit": unit[:16], "vat_rate": 6.0,
                "code": str(gi(r, "code") or "")[:24].rsplit("-", 1)[0], "lang": "pt", "country": "PT",
                "source": str(gi(r, "source") or "")[:120] or None, "ref_date": str(gi(r, "data do preço") or "")[:40] or None,
                "hours_mid": None})
            gamme = str(gi(r, "gamme") or "").lower()
            lvl = "low" if "conom" in gamme else "high" if "premium" in gamme else "mid"
            it[f"price_{lvl}"] = _f(gi(r, "preço mediano de referência")) or _f(gi(r, "preço calculado"))
            it[f"mat_{lvl}"] = _f(gi(r, "matériaux"))
            it[f"labor_{lvl}"] = _f(gi(r, "main-d'œuvre"))
            if lvl == "mid":
                it["hours_mid"] = _f(gi(r, "heures mo"))
    if "PISCINES_CONSTRUCTION" in wb.sheetnames:
        rows = list(wb["PISCINES_CONSTRUCTION"].iter_rows(values_only=True))
        h = _hdr_idx(rows[0])
        gi = lambda r, k: r[h[k]] if k in h and h[k] < len(r) else None  # noqa: E731
        for r in rows[1:]:
            desig = str(gi(r, "désignation") or "").strip()
            if not desig:
                continue
            cat = str(gi(r, "catégorie") or "Piscinas").strip()
            unit = str(gi(r, "unité") or "u").strip()
            key = (cat.lower(), desig.lower(), unit.lower())
            mid = _f(gi(r, "preço médio"))
            items[key] = {
                "category": cat[:80], "designation": desig[:220], "unit": unit[:16], "vat_rate": 6.0,
                "code": str(gi(r, "code") or "")[:24], "lang": "pt", "country": "PT",
                "price_low": _f(gi(r, "preço baixo")) or mid, "price_mid": mid, "price_high": _f(gi(r, "preço alto")) or mid,
                "mat_mid": _f(gi(r, "matériaux")), "labor_mid": _f(gi(r, "main-d'œuvre")),
                "source": "ConstruPilot PT 2026 – Piscinas", "ref_date": "2026-09"}
    out = []
    for it in items.values():
        mid = it.get("price_mid") or it.get("price_low") or it.get("price_high")
        if not mid:
            continue
        it["price_mid"] = mid
        it["price_low"] = it.get("price_low") or round(mid * 0.85, 2)
        it["price_high"] = it.get("price_high") or round(mid * 1.2, 2)
        out.append(it)
    return out
