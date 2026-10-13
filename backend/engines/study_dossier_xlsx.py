"""Excel projection of the same estimate, with formulas shown as inert evidence."""
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment


def build_dossier_xlsx(study):
    wb = Workbook()
    wb.remove(wb.active)

    def sheet(name, headers, widths, rows):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            ws.append(["N/C" if x is None else x for x in row])
        for row in ws:
            for cell in row:
                # Imported labels and formula evidence never become spreadsheet code.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill('solid', fgColor='925535')
                    cell.font = Font(bold=True, color='FFFFFF')
                elif cell.row % 2 == 0:
                    cell.fill = PatternFill('solid', fgColor='F7F1E9')
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'
        for col, width in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(1, col).column_letter].width = width
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        return ws

    sheet('Devis', ['Lot','Désignation','Détail','Unité','Quantité','Statut','PU HT','Total HT','Origine du prix','Inclus'],
          [24,44,60,9,14,10,15,17,26,10],
          [[r['group'],r['designation'],r['detail'],r['unit'],r['qty'],r['quantity_status'],r['unit_price'],r['total'],r['price_source'],r['included']] for r in study['items']])
    sheet('Calculs', ['Ouvrage','Formule','Résultat','Unité','Statut','Erreur'], [44,80,16,10,10,60],
          [[r['designation'],r['expression'],r['qty'],r['unit'],r['quantity_status'],r.get('error') or ''] for r in study['items']])
    sheet('Mesures et sources', ['Identifiant','Mesure','Valeur','Unité','Formule','Contrôlée','Source','Note'], [28,42,16,10,65,12,70,60],
          [[p['id'],p.get('label') or p['id'],p.get('value'),p.get('unit'),p.get('expression') or '',bool(p.get('reviewed')),
            '; '.join(f"p. {s['page']} : {s.get('text','')}" for s in p.get('sources',[])),p.get('note') or ''] for p in study['spec']['parameters']])
    t = study['totals']
    sheet('Synthèse', ['Indicateur','Valeur'], [55,70],
          [['HT partiel' if t['partial'] else 'HT',t['ht']],['TVA',t['tva']],['TTC',t['ttc']],
           ['TVA (%)',t['tva_rate']*100 if t['tva_rate'] is not None else None],['Révision',study['revision']],
           *[['Point à contrôler',w] for w in study['warnings']], *[['Hypothèse',x] for x in study['assumptions']],
           *[['Exclusion',x] for x in study['exclusions']]])
    out = BytesIO()
    wb.save(out)
    return out.getvalue()
