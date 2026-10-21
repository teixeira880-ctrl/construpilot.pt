"""Génère les PowerPoint brochure + catalogue ConstruPilot (python-pptx, 16:9)."""
import sys, os
sys.path.insert(0, os.path.dirname(__file__))
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR, MSO_AUTO_SIZE
from content import BRAND, INTRO, TIERS, TRANSVERSAL, COMPARE, HIGHLIGHTS, CONTACT


def rgb(h):
    return RGBColor.from_string(h.lstrip("#"))


DARK, LIME, AMBER, SOFT, GREY, MID, TEXT = (rgb(BRAND["dark"]), rgb(BRAND["lime"]), rgb(BRAND["amber"]),
                                            rgb(BRAND["soft"]), rgb(BRAND["grey"]), rgb(BRAND["mid"]), rgb("#1f2a24"))
SW, SH = Inches(13.333), Inches(7.5)
FONT = "Calibri"


def _rect(slide, x, y, w, h, color, shape=MSO_SHAPE.RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid(); s.fill.fore_color.rgb = color; s.line.fill.background()
    s.shadow.inherit = False
    return s


def _text(slide, x, y, w, h, text, size=14, bold=False, color=TEXT, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.05); tf.margin_top = tf.margin_bottom = Inches(0.02)
    p = tf.paragraphs[0]; p.alignment = align
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic; r.font.color.rgb = color; r.font.name = FONT
    return tb


def _bullets(slide, x, y, w, h, groups, size=11, title_size=12.5, gap=4):
    """groups: list of (title|None, [items])"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame; tf.word_wrap = True; tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    tf.margin_left = tf.margin_right = Inches(0.05)
    first = True
    for title, items in groups:
        if title:
            p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
            r = p.add_run(); r.text = title
            r.font.size = Pt(title_size); r.font.bold = True; r.font.color.rgb = MID; r.font.name = FONT
            p.space_before = Pt(gap + 3); p.space_after = Pt(1)
        for it in items:
            p = tf.paragraphs[0] if first else tf.add_paragraph(); first = False
            r = p.add_run(); r.text = "•  " + it
            r.font.size = Pt(size); r.font.color.rgb = TEXT; r.font.name = FONT
            p.space_after = Pt(gap)
    return tb


def _chrome(slide, title, kicker=None):
    _rect(slide, 0, 0, SW, Inches(0.12), DARK)
    _rect(slide, 0, 0, Inches(2.4), Inches(0.12), LIME)
    if kicker:
        _text(slide, Inches(0.5), Inches(0.25), Inches(8), Inches(0.35), kicker.upper(), 10, True, GREY)
    _text(slide, Inches(0.5), Inches(0.5), Inches(11), Inches(0.8), title, 30, True, DARK)
    _text(slide, Inches(0.5), SH - Inches(0.4), Inches(8), Inches(0.3), f"{BRAND['name']} — {BRAND['tagline']}", 9, False, GREY)
    _text(slide, SW - Inches(4.5), SH - Inches(0.4), Inches(4), Inches(0.3), BRAND["site"], 9, False, GREY, PP_ALIGN.RIGHT)


def _cover(prs, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, SW, SH, DARK)
    _rect(s, SW - Inches(4.6), -Inches(1.5), Inches(6), Inches(6), MID, MSO_SHAPE.OVAL)
    _rect(s, 0, Inches(4.05), SW, Inches(0.05), LIME)
    if os.path.exists(BRAND["logo"]):
        s.shapes.add_picture(BRAND["logo"], Inches(0.6), Inches(0.5), height=Inches(1.4))
    _text(s, Inches(0.6), Inches(2.3), Inches(10), Inches(1), BRAND["name"], 48, True, rgb("#ffffff"))
    _text(s, Inches(0.6), Inches(3.2), Inches(10), Inches(0.6), BRAND["tagline"], 20, False, LIME)
    _text(s, Inches(0.6), Inches(4.3), Inches(11), Inches(0.7), subtitle, 26, True, rgb("#ffffff"))
    _text(s, Inches(0.6), Inches(5.1), Inches(11), Inches(1.2),
          "Particuliers · Project Pass · Premium · Pro · Pro Business\nBlog & SEO · Radar Chantiers & prospection · API · Administration\nÉdition juin 2026 — document en français",
          13, False, SOFT)
    _text(s, Inches(0.6), SH - Inches(0.6), Inches(8), Inches(0.4), BRAND["site"], 11, False, GREY)


def _intro(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "Pourquoi ConstruPilot ?", "Présentation")
    _text(s, Inches(0.5), Inches(1.4), Inches(12.3), Inches(1.5), INTRO, 14, False, TEXT)
    x = Inches(0.5); w = Inches(2.38)
    for tr in TIERS:
        _rect(s, x, Inches(3.3), w, Inches(3.2), SOFT, MSO_SHAPE.ROUNDED_RECTANGLE)
        _rect(s, x, Inches(3.3), w, Inches(0.08), LIME)
        _text(s, x + Inches(0.15), Inches(3.5), w - Inches(0.3), Inches(0.5), tr["name"], 16, True, DARK)
        _text(s, x + Inches(0.15), Inches(4.0), w - Inches(0.3), Inches(0.5), tr["price"], 14, True, AMBER)
        _text(s, x + Inches(0.15), Inches(4.55), w - Inches(0.3), Inches(1.9), tr["audience"] + "\n\n" + tr["pitch"], 10.5, False, TEXT)
        x += w + Inches(0.1)


def _band_slide(prs, name, price, sub, pitch, kicker):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "", kicker)
    _rect(s, Inches(0.5), Inches(0.55), SW - Inches(1.0), Inches(1.35), DARK, MSO_SHAPE.ROUNDED_RECTANGLE)
    _text(s, Inches(0.75), Inches(0.62), Inches(8), Inches(0.6), name, 26, True, rgb("#ffffff"))
    _text(s, Inches(8.5), Inches(0.62), Inches(4.1), Inches(0.6), price, 20, True, LIME, PP_ALIGN.RIGHT)
    _text(s, Inches(0.75), Inches(1.15), Inches(12), Inches(0.35), sub, 11, False, SOFT)
    _text(s, Inches(0.75), Inches(1.45), Inches(12), Inches(0.4), pitch, 12, False, rgb("#ffffff"), italic=True)
    return s


def _split(sections):
    a, b, la, lb = [], [], 0, 0
    for title, items in sections:
        n = sum(len(i) for i in items) + 80
        if la <= lb:
            a.append((title, items)); la += n
        else:
            b.append((title, items)); lb += n
    return a, b


def _split_n(sections, n):
    cols = [[] for _ in range(n)]; load = [0] * n
    for title, items in sections:
        i = load.index(min(load))
        cols[i].append((title, items)); load[i] += sum(len(x) for x in items) + 80
    return cols


def _tier_brochure(prs, tr):
    s = _band_slide(prs, tr["name"], tr["price"], tr["sub"], tr["pitch"], f"Offre · {tr['audience']}")
    _rect(s, Inches(0.5), Inches(2.1), Inches(7.4), Inches(4.8), SOFT, MSO_SHAPE.ROUNDED_RECTANGLE)
    _text(s, Inches(0.7), Inches(2.2), Inches(5), Inches(0.4), "L'essentiel", 16, True, DARK)
    _bullets(s, Inches(0.7), Inches(2.7), Inches(7.0), Inches(4.1), [(None, HIGHLIGHTS[tr["key"]])], 13, gap=7)
    _rect(s, Inches(8.2), Inches(2.1), Inches(4.6), Inches(4.8), DARK, MSO_SHAPE.ROUNDED_RECTANGLE)
    _text(s, Inches(8.45), Inches(2.3), Inches(4.1), Inches(0.5), "Pour qui ?", 14, True, LIME)
    _text(s, Inches(8.45), Inches(2.8), Inches(4.1), Inches(1.2), tr["audience"], 13, False, rgb("#ffffff"))
    _text(s, Inches(8.45), Inches(4.1), Inches(4.1), Inches(0.5), "Tarif", 14, True, LIME)
    _text(s, Inches(8.45), Inches(4.6), Inches(4.1), Inches(1.8), tr["price"] + "\n" + tr["sub"], 12, False, SOFT)
    total = sum(len(x) for _, items in tr["sections"] for x in items)
    n = 3 if total > 1800 else 2
    size = 9 if n == 3 else 11
    s2 = _band_slide(prs, f"{tr['name']} — dans le détail", tr["price"], "", tr["pitch"], f"Offre · {tr['audience']}")
    cw = (SW - Inches(1.0) - Inches(0.2) * (n - 1)) / n
    for i, col in enumerate(_split_n(tr["sections"], n)):
        _bullets(s2, Inches(0.5) + i * (cw + Inches(0.2)), Inches(2.1), cw, Inches(4.9), col, size, size + 2, 3)


def _module_brochure(prs, tv, kicker):
    s = _band_slide(prs, tv["name"], "", "", tv["pitch"], kicker)
    a, b = _split(tv["sections"])
    _bullets(s, Inches(0.5), Inches(2.1), Inches(6.1), Inches(4.9), a, 11.5, 13, 5)
    _bullets(s, Inches(6.8), Inches(2.1), Inches(6.1), Inches(4.9), b, 11.5, 13, 5)


def _chunks(sections, budget=1500):
    out, cur, load = [], [], 0
    for sec in sections:
        n = sum(len(x) for x in sec[1]) + 80
        if cur and (load + n > budget or len(cur) == 2):
            out.append(cur); cur, load = [], 0
        cur.append(sec); load += n
    if cur:
        out.append(cur)
    return out


def _section_slides(prs, title, price_sub, pitch, sections, kicker):
    for i, chunk in enumerate(_chunks(sections)):
        s = _band_slide(prs, title, price_sub if i == 0 else "", "", pitch, kicker)
        if len(chunk) == 1:
            _bullets(s, Inches(0.5), Inches(2.1), Inches(12.3), Inches(4.9), chunk, 12.5, 15, 6)
        else:
            _bullets(s, Inches(0.5), Inches(2.1), Inches(6.1), Inches(4.9), chunk[:1], 11.5, 14, 5)
            _bullets(s, Inches(6.8), Inches(2.1), Inches(6.1), Inches(4.9), chunk[1:], 11.5, 14, 5)


def _compare(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "Comparatif des offres", "Synthèse")
    rows, cols = len(COMPARE) + 1, 5
    tbl = s.shapes.add_table(rows, cols, Inches(0.5), Inches(1.35), Inches(12.3), Inches(5.5)).table
    widths = [Inches(5.1), Inches(1.8), Inches(1.8), Inches(1.8), Inches(1.8)]
    for i, w in enumerate(widths):
        tbl.columns[i].width = w
    heads = ["Fonction", "Particuliers", "Premium", "Pro", "Pro Business"]
    for c, h in enumerate(heads):
        cell = tbl.cell(0, c); cell.text = h
        cell.fill.solid(); cell.fill.fore_color.rgb = DARK
        p = cell.text_frame.paragraphs[0]; p.font.size = Pt(10); p.font.bold = True; p.font.color.rgb = rgb("#ffffff"); p.font.name = FONT
        p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
    for r, row in enumerate(COMPARE, 1):
        for c, val in enumerate(row):
            cell = tbl.cell(r, c); cell.text = val
            cell.fill.solid(); cell.fill.fore_color.rgb = SOFT if r % 2 == 0 else rgb("#ffffff")
            cell.margin_top = cell.margin_bottom = Emu(20000)
            p = cell.text_frame.paragraphs[0]; p.font.size = Pt(8.5); p.font.color.rgb = TEXT; p.font.name = FONT
            p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
        tbl.rows[r].height = Inches(0.27)
    tbl.rows[0].height = Inches(0.35)


def _closing(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _rect(s, 0, 0, SW, SH, DARK)
    _text(s, Inches(0.6), Inches(1.8), Inches(12), Inches(1), "Estimez. Comparez. Construisez.", 40, True, rgb("#ffffff"))
    _text(s, Inches(0.6), Inches(2.8), Inches(12), Inches(0.6), "Passez à ConstruPilot dès aujourd'hui", 20, False, LIME)
    _bullets(s, Inches(0.6), Inches(3.7), Inches(12), Inches(2.5), [(None, CONTACT)], 14, gap=8)
    for p in s.shapes[-1].text_frame.paragraphs:
        for r in p.runs:
            r.font.color.rgb = SOFT
    _text(s, Inches(0.6), SH - Inches(0.7), Inches(8), Inches(0.4), BRAND["site"], 12, False, GREY)


def build_brochure(path):
    prs = Presentation(); prs.slide_width, prs.slide_height = SW, SH
    _cover(prs, "Toutes les fonctionnalités — brochure")
    _intro(prs)
    for tr in TIERS:
        _tier_brochure(prs, tr)
    for tv in TRANSVERSAL:
        _module_brochure(prs, tv, "Module transversal")
    _compare(prs)
    _closing(prs)
    prs.save(path)


def build_catalogue(path):
    prs = Presentation(); prs.slide_width, prs.slide_height = SW, SH
    _cover(prs, "Catalogue exhaustif des fonctionnalités")
    _intro(prs)
    s = prs.slides.add_slide(prs.slide_layouts[6])
    _chrome(s, "Sommaire", "Catalogue")
    toc = [f"{i}. {t['name']} — {t['price']}" for i, t in enumerate(TIERS, 1)]
    toc += [f"{i}. {t['name']}" for i, t in enumerate(TRANSVERSAL, len(TIERS) + 1)]
    toc.append(f"{len(TIERS) + len(TRANSVERSAL) + 1}. Comparatif des offres")
    _bullets(s, Inches(0.5), Inches(1.5), Inches(12), Inches(5.5), [(None, toc)], 15, gap=8)
    for i, tr in enumerate(TIERS, 1):
        _section_slides(prs, f"{i}. {tr['name']}", tr["price"], tr["pitch"], tr["sections"], f"Offre · {tr['sub']}")
    for i, tv in enumerate(TRANSVERSAL, len(TIERS) + 1):
        _section_slides(prs, f"{i}. {tv['name']}", "", tv["pitch"], tv["sections"], "Module transversal")
    _compare(prs)
    _closing(prs)
    prs.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "/app/frontend/public/docs"
    os.makedirs(out, exist_ok=True)
    build_brochure(f"{out}/ConstruPilot_Brochure_Fonctionnalites.pptx")
    build_catalogue(f"{out}/ConstruPilot_Catalogue_Complet.pptx")
    print("ok")
