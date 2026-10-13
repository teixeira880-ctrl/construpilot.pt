"""Routes /api/plans-pro/* — orchestration Autodesk + Vision cascade.

Endpoints :
  - ``POST /api/plans-pro/upload``           : DWG/RVT/IFC/PDF/PNG → job_id
  - ``GET  /api/plans-pro/job/{id}``         : progression
  - ``GET  /api/plans-pro/job/{id}/result``  : résultat complet
  - ``POST /api/plans-pro/analyze-image``    : analyse rapide image seule (base64)
  - ``GET  /api/plans-pro/history``          : historique du user

Le pipeline complet tourne en background ; le client interroge ``/job/{id}``
jusqu'à ``status == done``.
"""
import asyncio
import base64
import os
import logging
import uuid
from datetime import datetime, timezone

from typing import Optional

from fastapi import (APIRouter, BackgroundTasks, Depends, File, Form, HTTPException,
                     Request, Response, UploadFile)
from pydantic import BaseModel, Field

from core.db import db
from core.rate_limit import limiter
from core.security import current_user
from engines import storage
from plans_pro import autodesk_pipeline, vision_cascade

router = APIRouter(prefix="/plans-pro", tags=["plans-pro"])
logger = logging.getLogger(__name__)


# ext → (mime par défaut, needs_autodesk)
_ACCEPT_MIME: dict[str, tuple[str, bool]] = {
    ".dwg": ("model/vnd.dwg", True),
    ".dxf": ("image/vnd.dxf", True),
    ".rvt": ("application/octet-stream", True),
    ".ifc": ("application/x-step", True),
    ".nwc": ("application/octet-stream", True),
    ".nwd": ("application/octet-stream", True),
    ".pdf": ("application/pdf", True),
    ".png": ("image/png", False),
    ".jpg": ("image/jpeg", False),
    ".jpeg": ("image/jpeg", False),
    ".webp": ("image/webp", False),
}


FINAL_STATUSES = {"done", "partial", "needs_review", "failed", "cancelled"}
STATUS_HELP = {
    "queued": "En file d'attente.",
    "autodesk_translating": "Traduction Autodesk en cours (2 à 10 min selon le fichier).",
    "autodesk_done": "Modèle Autodesk lu, lecture IA de l'aperçu en cours.",
    "running_vision": "Lecture IA du plan en cours.",
    "done": "Analyse terminée : toutes les étapes ont produit un résultat exploitable.",
    "needs_review": "Analyse terminée mais des comptages divergent entre lectures : vérifiez les éléments signalés.",
    "partial": "Résultat partiel : une étape n'a pas abouti. Les données obtenues restent utilisables, les contrôles manquants sont indiqués.",
    "failed": "L'analyse n'a produit aucun résultat exploitable. Relancez-la ou changez de fichier.",
    "cancelled": "Analyse annulée.",
}


def _finalize(*, needs_autodesk: bool, autodesk_ok: bool, vision: Optional[dict]) -> tuple[str, list[str]]:
    """Contrat de résultat commun : `done` exige un résultat valide pour chaque étape attendue."""
    checks = []
    v_status = (vision or {}).get("status") if vision else None
    if needs_autodesk and not autodesk_ok:
        return "failed", ["Extraction Autodesk : échec"]
    if needs_autodesk:
        checks.append("Extraction Autodesk : réussie")
    if vision is None:
        checks.append("Lecture IA : non exécutée")
        return ("partial" if needs_autodesk else "failed"), checks
    if v_status == "failed":
        checks.append("Lecture IA : échec")
        return ("partial" if needs_autodesk and autodesk_ok else "failed"), checks
    if v_status == "needs_review":
        checks.append("Lecture IA : désaccords à vérifier")
        return "needs_review", checks
    checks.append("Lecture IA : réussie")
    return "done", checks


async def _pipeline(job_id: str, file_bytes: bytes, filename: str,
                    ext: str, needs_autodesk: bool, user_id: str):
    """Pipeline complet en background — ne termine jamais en `done` sans résultat exploitable."""
    async def _upd(**fields):
        await db.plans_pro_jobs.update_one({"job_id": job_id}, {"$set": fields})

    autodesk_ok = False
    vision: Optional[dict] = None
    vision_input: Optional[bytes] = None
    vision_mime = "image/png"
    vision_source = "upload"
    autodesk_error = None

    try:
        if needs_autodesk:
            if not autodesk_pipeline.is_configured():
                await _upd(status="failed", progress="0%", error="Autodesk (APS) n'est pas configuré : ce format ne peut pas être analysé.",
                           checks=["Extraction Autodesk : non configurée"], finished_at=datetime.now(timezone.utc))
                return
            await _upd(status="autodesk_translating", progress="10%")
            try:
                autodesk_result = await autodesk_pipeline.extract_full(file_bytes, filename, user_id=user_id)
                autodesk_ok = bool(autodesk_result.get("urn")) and bool(autodesk_result.get("views") or autodesk_result.get("quantities"))
                await _upd(
                    status="autodesk_done", progress="60%",
                    autodesk_urn=autodesk_result.get("urn"),
                    autodesk_views=[{"name": v.get("name"), "role": v.get("role"), "objects": v.get("objects")}
                                    for v in (autodesk_result.get("views") or [])],
                    autodesk_quantities=autodesk_result.get("quantities"),
                    autodesk_cached=autodesk_result.get("cached", False),
                    autodesk_status=autodesk_result.get("status"),
                )
                if autodesk_result.get("urn"):
                    thumb = await autodesk_pipeline.thumbnail(autodesk_result["urn"])
                    if thumb:
                        mime = "image/png" if thumb[:4] == b"\x89PNG" else "image/jpeg"
                        await _upd(autodesk_thumbnail=f"data:{mime};base64," + base64.b64encode(thumb).decode("ascii"))
                        vision_input, vision_mime, vision_source = thumb, mime, "autodesk_thumbnail"
            except Exception as e:
                logger.exception("Autodesk extract failed for %s", job_id)
                autodesk_error = str(e)[:300]
        else:
            vision_input = file_bytes
            vision_mime = _ACCEPT_MIME.get(ext, ("image/png", False))[0]

        if vision_input and (autodesk_ok or not needs_autodesk):
            await _upd(status="running_vision", progress="80%")
            vision = await vision_cascade.analyze_plan_image(image_bytes=vision_input, image_mime=vision_mime)
            vision["source"] = vision_source
            await _upd(vision=vision)

        status, checks = _finalize(needs_autodesk=needs_autodesk, autodesk_ok=autodesk_ok, vision=vision)
        error = None
        if status == "failed":
            error = autodesk_error and f"Autodesk n'a pas pu traiter ce fichier : {autodesk_error}" or (vision or {}).get("error") or "Aucun résultat exploitable."
        elif status == "partial":
            error = (vision or {}).get("error") or "Lecture IA indisponible sur l'aperçu Autodesk."
        await _upd(status=status, progress="100%", checks=checks, error=error, help=STATUS_HELP[status],
                   finished_at=datetime.now(timezone.utc))

    except Exception as e:
        logger.exception("Pipeline plans-pro failed for %s", job_id)
        try:
            await _upd(status="failed", error="Erreur interne pendant l'analyse : relancez-la.", technical_error=str(e)[:300],
                       checks=["Pipeline : interrompu"], finished_at=datetime.now(timezone.utc))
        except Exception:
            pass


def _job_path(user_id: str, job_id: str, ext: str) -> str:
    return f"private/plans-pro/{user_id}/{job_id}{ext}"


async def _start_job(background: BackgroundTasks, user: dict, filename: str, ext: str, data: bytes, *, retry_of: str | None = None) -> dict:
    needs_autodesk = _ACCEPT_MIME[ext][1]
    job_id = f"pplan_{uuid.uuid4().hex[:14]}"
    stored = False
    try:
        await asyncio.to_thread(storage.put_object, _job_path(user["user_id"], job_id, ext), data, _ACCEPT_MIME[ext][0])
        stored = True
    except Exception as exc:  # noqa: BLE001 — la relance ne sera simplement pas disponible
        logger.warning("plans-pro: stockage du fichier source impossible (%s)", exc)
    await db.plans_pro_jobs.insert_one({
        "job_id": job_id, "user_id": user["user_id"], "filename": filename, "size_bytes": len(data), "extension": ext,
        "needs_autodesk": needs_autodesk, "autodesk_configured": autodesk_pipeline.is_configured(),
        "status": "queued", "progress": "0%", "help": STATUS_HELP["queued"], "source_stored": stored, "retry_of": retry_of,
        "created_at": datetime.now(timezone.utc),
    })
    background.add_task(_pipeline, job_id, data, filename, ext, needs_autodesk, user["user_id"])
    return {"job_id": job_id, "status": "queued", "needs_autodesk": needs_autodesk, "autodesk_configured": autodesk_pipeline.is_configured(), "retry_of": retry_of}


@router.post("/upload")
@limiter.limit("20/hour")
async def upload(request: Request, response: Response, background: BackgroundTasks,
                 file: UploadFile = File(...), chantier_id: str | None = Form(None), user=Depends(current_user)):
    filename = file.filename or "plan"
    ext = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if ext not in _ACCEPT_MIME:
        raise HTTPException(400, f"Format non supporté ({ext}). Accepté : {', '.join(_ACCEPT_MIME)}")
    _, needs_autodesk = _ACCEPT_MIME[ext]
    if needs_autodesk and not autodesk_pipeline.is_configured():
        raise HTTPException(409, f"Le format {ext} nécessite Autodesk (APS), qui n'est pas configuré sur ce serveur. Importez une image PNG/JPG du plan ou contactez l'administrateur.")

    data = await file.read()
    if not data:
        raise HTTPException(400, "Fichier vide")
    if len(data) > 100 * 1024 * 1024:
        raise HTTPException(400, "Fichier trop volumineux (100 Mo max)")
    started = await _start_job(background, user, filename, ext, data)
    from routes.chantiers import owned_chantier_id
    owned = await owned_chantier_id(user["user_id"], chantier_id)
    if owned:
        await db.plans_pro_jobs.update_one({"job_id": started["job_id"]}, {"$set": {"chantier_id": owned}})
    return {**started, "chantier_id": owned}


@router.post("/job/{job_id}/retry")
@limiter.limit("20/hour")
async def retry_job(job_id: str, request: Request, response: Response, background: BackgroundTasks, user=Depends(current_user)):
    """Relance une analyse depuis le fichier source conservé (nouveau job lié)."""
    doc = await db.plans_pro_jobs.find_one({"job_id": job_id, "user_id": user["user_id"]}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Job introuvable")
    if doc.get("status") not in FINAL_STATUSES:
        raise HTTPException(409, "Analyse encore en cours")
    if not doc.get("source_stored"):
        raise HTTPException(409, "Fichier source non conservé : importez-le à nouveau.")
    ext = doc.get("extension") or ""
    if _ACCEPT_MIME.get(ext, ("", False))[1] and not autodesk_pipeline.is_configured():
        raise HTTPException(409, "Autodesk (APS) n'est pas configuré : relance impossible pour ce format.")
    try:
        data, _ = await asyncio.to_thread(storage.get_object, _job_path(user["user_id"], job_id, ext))
    except Exception:  # noqa: BLE001
        raise HTTPException(409, "Fichier source introuvable : importez-le à nouveau.") from None
    return await _start_job(background, user, doc.get("filename") or "plan", ext, data, retry_of=job_id)


@router.get("/job/{job_id}")
async def job_status(job_id: str, user=Depends(current_user)):
    doc = await db.plans_pro_jobs.find_one(
        {"job_id": job_id, "user_id": user["user_id"]},
        {"_id": 0, "user_id": 0},
    )
    if not doc:
        raise HTTPException(404, "Job introuvable")
    return doc


@router.get("/job/{job_id}/quantities.xlsx")
async def job_quantities_xlsx(job_id: str, user=Depends(current_user)):
    """Tableau des quantités Autodesk (catégories / calques) en Excel prêt pour le métré."""
    doc = await db.plans_pro_jobs.find_one({"job_id": job_id, "user_id": user["user_id"]},
                                           {"_id": 0, "filename": 1, "status": 1, "autodesk_quantities": 1, "autodesk_elements_count": 1, "autodesk_views": 1, "created_at": 1})
    if not doc:
        raise HTTPException(404, "Job introuvable")
    quantities = doc.get("autodesk_quantities") or {}
    if doc.get("status") not in FINAL_STATUSES or not quantities:
        raise HTTPException(409, "Aucune quantité Autodesk disponible pour ce job")
    data = await asyncio.to_thread(_quantities_workbook, doc, quantities)
    stem = os.path.splitext(doc.get("filename") or "plan")[0][:40].replace('"', "")
    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="quantites-{stem}.xlsx"', "Cache-Control": "private, no-store"})


def _quantities_workbook(doc: dict, quantities: dict) -> bytes:
    import io

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Quantités"
    head_fill, head_font = PatternFill("solid", fgColor="0F766E"), Font(bold=True, color="FFFFFF")
    ws.append(["ConstruPilot — Analyse de plan Pro", doc.get("filename")])
    ws.append(["Généré le", datetime.now(timezone.utc).strftime("%d/%m/%Y %H:%M UTC"), "Entités Autodesk", doc.get("autodesk_elements_count")])
    ws.append(["Source", "Comptage par calque (DWG) — quantités indicatives à valider par métré ; surfaces/volumes uniquement Revit/IFC"])
    ws.append([])
    headers = ["Ouvrage", "Nombre d'entités", "Surface (m²)", "Volume (m³)", "Longueur (m)", "Source", "Calques principaux"]
    ws.append(headers)
    for c in range(1, len(headers) + 1):
        cell = ws.cell(row=5, column=c)
        cell.fill, cell.font, cell.alignment = head_fill, head_font, Alignment(horizontal="center")
    for cat, q in quantities.items():
        ws.append([cat, q.get("count"), q.get("total_area_m2"), q.get("total_volume_m3"), q.get("total_length_m"),
                   "calques" if q.get("source") == "layers" else "catégories",
                   " ; ".join(f"{l['layer']} ({l['count']})" for l in (q.get("layers") or [])[:8])])
    for i, w in enumerate([22, 16, 14, 14, 14, 12, 70], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = "A6"

    ws2 = wb.create_sheet("Calques")
    ws2.append(["Ouvrage", "Calque", "Nombre d'entités"])
    for c in range(1, 4):
        ws2.cell(row=1, column=c).fill, ws2.cell(row=1, column=c).font = head_fill, head_font
    for cat, q in quantities.items():
        for l in q.get("layers") or []:
            ws2.append([cat, l["layer"], l["count"]])
    ws2.column_dimensions["A"].width, ws2.column_dimensions["B"].width, ws2.column_dimensions["C"].width = 22, 40, 18
    ws2.freeze_panes = "A2"

    ws3 = wb.create_sheet("Métré")
    ws3.append(["Ouvrage", "Désignation", "Unité", "Quantité", "PU HT (€)", "Total HT (€)"])
    for c in range(1, 7):
        ws3.cell(row=1, column=c).fill, ws3.cell(row=1, column=c).font = head_fill, head_font
    r = 2
    for cat, q in quantities.items():
        if cat in ("Annotations", "Trames / axes", "Détails", "Autres"):
            continue
        ws3.append([cat, "", "u", q.get("count"), None, f"=D{r}*E{r}"])
        r += 1
    ws3.append(["", "", "", "", "Total HT", f"=SUM(F2:F{max(r - 1, 2)})"])
    for i, w in enumerate([22, 40, 8, 12, 12, 14], start=1):
        ws3.column_dimensions[get_column_letter(i)].width = w

    if doc.get("autodesk_views"):
        ws4 = wb.create_sheet("Vues")
        ws4.append(["Vue", "Type", "Objets"])
        for v in doc["autodesk_views"]:
            ws4.append([v.get("name"), v.get("role"), v.get("objects")])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@router.get("/job/{job_id}/result")
async def job_result(job_id: str, user=Depends(current_user)):
    doc = await db.plans_pro_jobs.find_one(
        {"job_id": job_id, "user_id": user["user_id"]},
        {"_id": 0, "user_id": 0},
    )
    if not doc:
        raise HTTPException(404, "Job introuvable")
    if doc.get("status") not in FINAL_STATUSES:
        raise HTTPException(409, f"Job non terminé ({doc.get('status')})")
    return {
        "job_id": job_id,
        "filename": doc.get("filename"),
        "autodesk": {
            "urn": doc.get("autodesk_urn"),
            "status": doc.get("autodesk_status"),
            "thumbnail": doc.get("autodesk_thumbnail"),
            "views": doc.get("autodesk_views"),
            "quantities": doc.get("autodesk_quantities"),
            "cached": doc.get("autodesk_cached"),
            "elements_count": sum(
                (v.get("objects") or 0) for v in (doc.get("autodesk_views") or [])
            ) or None,
        } if doc.get("autodesk_urn") else None,
        "vision": doc.get("vision"),
        "status": doc.get("status"),
        "checks": doc.get("checks") or [],
        "help": doc.get("help"),
        "error": doc.get("error"),
        "can_retry": bool(doc.get("source_stored")),
        "progress": doc.get("progress"),
        "finished_at": doc.get("finished_at"),
    }


class ImageIn(BaseModel):
    image_b64: str = Field(..., description="Image encodée base64 (sans préfixe data:)")
    mime: str = "image/png"


@router.post("/analyze-image")
@limiter.limit("30/hour")
async def analyze_image(body: ImageIn, request: Request, response: Response, user=Depends(current_user)):
    """Analyse rapide sans conversion Autodesk (PNG/JPG directement)."""
    try:
        raw = base64.b64decode(body.image_b64)
    except Exception:
        raise HTTPException(400, "image_b64 invalide")
    if len(raw) > 20 * 1024 * 1024:
        raise HTTPException(400, "Image trop grande (20 Mo max)")
    return await vision_cascade.analyze_plan_image(image_bytes=raw, image_mime=body.mime)


@router.get("/history")
async def history(limit: int = 20, user=Depends(current_user)):
    limit = max(1, min(int(limit or 20), 100))
    items = await db.plans_pro_jobs.find(
        {"user_id": user["user_id"]},
        {"_id": 0, "user_id": 0, "vision.text_detected": 0, "autodesk_thumbnail": 0},
    ).sort("created_at", -1).to_list(limit)
    return {"items": items}
