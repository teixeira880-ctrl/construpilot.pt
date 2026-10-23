"""Testes do estudo de edifício multifamiliar (engines/edificio_ba) — caso de exemplo anonimizado e variantes."""
from __future__ import annotations

import asyncio
import importlib.util
import copy
import json
from pathlib import Path

import pytest

from engines.edificio_ba import assistant_tool, fem3d, grelha, leitura, medicoes, modelo, pipeline

CASE = Path(__file__).resolve().parents[2] / "engines" / "edificio_ba" / "casos" / "exemplo_pt.json"
CASE_FR = CASE.with_name("exemplo_fr.json")
HAS_PT = importlib.util.find_spec("engines.calc_lots.reh") is not None          # motores PT (construpilot.pt)
HAS_FR = importlib.util.find_spec("engines.calc_lots.re2020") is not None       # moteurs FR (construpilot.fr)


@pytest.fixture()
def raw():
    return json.loads(CASE.read_text())


@pytest.fixture()
def S(raw):
    return modelo.load(raw)


@pytest.fixture()
def R(S):
    return pipeline.calcular(S, use_fem=False)


def test_example_is_valid(S):
    assert modelo.validate(S) == []
    assert S.n == 5 and S.basement and len(S.units) == 20
    assert [lv[1] for lv in S.levels()] == ["Piso -1", "Piso 0", "Piso 1", "Piso 2", "Piso 3", "Piso 4", "Cobertura"]


def test_auto_beams_reach_cores(S):
    lines = grelha.beam_lines(S)
    ids = {(ln["a"], ln["b"]) for ln in lines}
    assert ((5.48, 0.0), (5.48, 5.91)) in ids                 # pilar de fachada → face do núcleo de escada
    assert ((10.82, 0.0), (10.82, 6.11)) in ids               # pilar de fachada → face do elevador
    assert sum(1 for ln in lines if ln["facade"]) == 4


def test_panels_cover_the_floor(S):
    segs = grelha.beam_segments(S, grelha.beam_lines(S))
    pans = grelha.panels(S, segs)
    cores = sum((c["box"][2] - c["box"][0]) * (c["box"][3] - c["box"][1]) for c in S.cores)
    area = sum(p["area"] for p in pans)
    assert 0.90 * (S.area_floor - cores) <= area <= 1.02 * (S.area_floor - cores)
    assert all(max(p["x1"] - p["x0"], p["y1"] - p["y0"]) <= 15 for p in pans)


def test_torsion_controlled_by_proposed_walls(R):
    assert 1 <= len(R["proposed_walls"]) <= 4
    assert R["seis"]["reg"]["ok"]
    assert R["seis"]["T1"] > 0.2


def test_columns_and_beams_verified(R):
    assert all(c["base"]["nu"] <= 0.65 + 1e-9 for c in R["cols"])
    assert all(b["ok"] for b in R["beams"])
    assert all(s["ok"] for s in R["slabs"])


def test_quantities_regression(S, R):
    Q = medicoes.medir(S, R)
    O = medicoes.orcamento(S, R, Q)
    assert 850 <= Q["totais"]["betao"] <= 1150             # estudo de referência : ≈ 980 m³
    assert 330_000 <= O["tot_s"] <= 460_000                 # estrutura ≈ 385 k€
    assert sum(v["n"] for v in Q["quant"]["vaos"].values()) == 65
    assert abs(Q["areas"]["bruta"] - 2122) < 5


def test_foundations(R):
    assert R["fnd"] and all(f["L"] and f["B"] for f in R["fnd"])
    assert all(f["sigma_Ed"] <= 450 + 1 for f in R["foots"])


def test_blocking_column_outside(raw):
    raw["columns"].append({"id": "PX", "x": 40.0, "y": 5.0})
    res = pipeline.run(raw, None, use_fem=False, especialidades=False)
    assert res["status"] == "MODELO_INVALIDO"
    assert any(a["code"] == "COL-OUT" for a in res["anomalies"])


def test_without_basement(raw):
    raw["levels"]["basement"] = False
    raw.pop("basement")
    S = modelo.load(raw)
    R = pipeline.calcular(S, use_fem=False)
    assert R["bwalls"] == [] and R["cols"]
    Q = medicoes.medir(S, R)
    assert not any(x["art"] == "1.1" for x in Q["lines"])


def test_french_profile_structure(raw):
    raw["project"]["country"] = "FR"
    raw["seismic"] = {"zone": "3", "q": 2.0, "importance": "II"}
    S = modelo.load(raw)
    assert S.lang == "fr" and S.lname(0) == "RDC" and S.lname(S.n) == "Toiture-terrasse"
    R = pipeline.calcular(S, use_fem=False)
    assert R["seis"]["ok"] and len(R["seis"]["acts"]) == 1 and R["seis"]["required"]


@pytest.mark.skipif(not HAS_PT, reason="motores de especialidades PT ausentes")
def test_full_run_files(raw, tmp_path):
    res = pipeline.run(raw, tmp_path, use_fem=False)
    assert res["status"] == "PRONTO"
    for k in ("memoria", "desenhos", "medicoes", "orcamento", "quantidades", "xlsx", "termico", "pichelaria", "eletricidade"):
        p = Path(res["files"][k])
        assert p.exists() and p.stat().st_size > 5000, k
    from openpyxl import load_workbook
    wb = load_workbook(res["files"]["xlsx"])
    assert len(wb.sheetnames) >= 8
    assert res["summary"]["especialidades"]["classes"]


@pytest.mark.skipif(not fem3d.available(), reason="OpenSeesPy não instalado")
def test_fem_model(S):
    R = pipeline.calcular(S, use_fem=True)
    fm = R["fem"]
    assert fm and 0.3 < fm["modes"][0]["T"] < 0.8
    assert sum(m["Mx"] for m in fm["modes"]) > 0.85 and sum(m["My"] for m in fm["modes"]) > 0.85


def test_tool_schema():
    t = assistant_tool.TOOLS[0]
    p = t["parameters"]
    assert t["strict"] and set(p["required"]) == set(p["properties"])
    assert assistant_tool.TOOL_NAMES == set(assistant_tool.TOOL_STATUS)


class _Col:
    def __init__(self):
        self.docs = {}

    async def update_one(self, flt, upd, upsert=False):
        key = flt.get("job_id") or flt.get("artifact_id")
        d = self.docs.setdefault(key, {})
        d.update(upd.get("$set", {}))
        for k, v in upd.get("$setOnInsert", {}).items():
            d.setdefault(k, v)

    async def find_one(self, flt, proj=None, sort=None):
        for d in reversed(list(self.docs.values())):
            if all(d.get(k) == v for k, v in flt.items()):
                return {k: v for k, v in d.items() if not proj or proj.get(k, 1)}
        return None


class _DB:
    def __init__(self):
        self.edificio_jobs = _Col()
        self.assistant_artifacts = _Col()


def test_assistant_demo_offline():
    db = _DB()
    user, conv = {"user_id": "u1"}, {"conversation_id": "c1"}
    out = asyncio.run(assistant_tool.execute("generate_multifamily_building_study", {"document_ids": [], "hypotheses": "", "country": "PT", "demo": True},
                                             user, conv, db, background=False))
    assert out["ok"]
    job = asyncio.run(assistant_tool.execute("read_multifamily_building_study", {"job_id": out["job_id"]}, user, conv, db))
    assert job["status"] == "done", job.get("error")
    assert len(job["artifacts"]) >= 6
    assert job["summary"]["betao_m3"] > 500


def test_reader_with_fake_llm(raw):
    pytest.importorskip("fitz")
    import fitz
    doc = fitz.open()
    pg = doc.new_page(width=842, height=595)
    pg.insert_text((50, 50), "PLANTA DO PISO 1")
    pdf = doc.tobytes()
    calls = []

    async def fake(system, text, images=None, model=None, max_tokens=4096, **kw):
        calls.append(1)
        return json.dumps(raw)
    rd = asyncio.run(leitura.read_building([("plantas.pdf", pdf)], "", "PT", llm=fake))
    assert rd["ok"] and len(calls) == 1 and rd["model"]["footprint"]["L"] == 21.25


def test_reader_correction_loop(raw):
    pytest.importorskip("fitz")
    import fitz
    doc = fitz.open()
    doc.new_page().insert_text((50, 50), "PLANTA DO PISO 1")
    bad = copy.deepcopy(raw)
    bad["columns"][0]["x"] = 60.0
    seq = [json.dumps(bad), json.dumps(raw)]

    async def fake(system, text, images=None, model=None, max_tokens=4096, **kw):
        return seq.pop(0)
    rd = asyncio.run(leitura.read_building([("p.pdf", doc.tobytes())], "", "PT", llm=fake))
    assert rd["ok"] and rd["attempts"] == 2


@pytest.mark.skipif(not HAS_FR, reason="moteurs FR absents")
def test_full_run_fr(tmp_path):
    raw = json.loads(CASE_FR.read_text())
    res = pipeline.run(raw, tmp_path, use_fem=False)
    assert res["status"] == "PRONTO"
    S = res["S"]
    assert S.lang == "fr" and not res["R"]["seis"]["required"]            # zone 1 : vent, pas de voile proposé
    assert res["R"]["proposed_walls"] == [] and res["R"]["seis"]["wind"]
    for k in ("memoria", "desenhos", "medicoes", "orcamento", "quantidades", "xlsx", "termico", "pichelaria", "eletricidade"):
        assert Path(res["files"][k]).exists(), k
    assert res["summary"]["especialidades"]["Bbio"] > 0


def test_fr_zone3_proposes_walls():
    raw = json.loads(CASE_FR.read_text())
    raw["seismic"] = {"zone": "3", "q": 2.0, "importance": "II"}
    S = modelo.load(raw)
    R = pipeline.calcular(S, use_fem=False)
    assert R["seis"]["required"] and R["proposed_walls"] and R["seis"]["reg"]["ok"]
