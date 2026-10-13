"""Iter 93 - Backend tests for:
- GET beton-xlsx export
- GET analysis includes acier_* + acier_ratios
- POST ai-refine (best-effort, IA can be slow/fail)
- GET /api/documents/{doc_id}/pdf?lang=fr
"""
import io
import os
import pytest
import requests
from openpyxl import load_workbook

def _load_backend_url():
    v = os.environ.get("REACT_APP_BACKEND_URL")
    if not v:
        # read /app/frontend/.env
        try:
            with open("/app/frontend/.env") as f:
                for line in f:
                    if line.startswith("REACT_APP_BACKEND_URL="):
                        v = line.split("=", 1)[1].strip()
                        break
        except Exception:
            pass
    assert v, "REACT_APP_BACKEND_URL not found"
    return v.rstrip("/")

BASE_URL = _load_backend_url()
DEMO_EMAIL = "demo@construpilot.fr"
DEMO_PASSWORD = "Demo2026!"


@pytest.fixture(scope="module")
def token():
    r = requests.post(f"{BASE_URL}/api/auth/login",
                      json={"email": DEMO_EMAIL, "password": DEMO_PASSWORD}, timeout=30)
    assert r.status_code == 200, r.text
    tok = r.json().get("token") or r.json().get("access_token")
    assert tok
    return tok


@pytest.fixture(scope="module")
def headers(token):
    return {"Authorization": f"Bearer {token}"}


@pytest.fixture(scope="module")
def plan_id(headers):
    r = requests.get(f"{BASE_URL}/api/takeoffs", headers=headers, timeout=30)
    assert r.status_code == 200, r.text
    data = r.json()
    items = data.get("plans", []) if isinstance(data, dict) else data
    assert items, "Aucun plan sur le compte démo"
    # Prefer plan_c56da42f0331 (coffrage) if present
    preferred = [it for it in items if it.get("plan_id") == "plan_c56da42f0331"]
    return (preferred or items)[0]["plan_id"]


# ---- Analysis: acier_* fields ----
class TestAnalysisAcier:
    def test_analysis_contains_acier_fields(self, headers, plan_id):
        r = requests.get(
            f"{BASE_URL}/api/takeoffs/{plan_id}/analysis/0?wall_height=2.5",
            headers=headers, timeout=60)
        assert r.status_code == 200, r.text
        data = r.json()
        d = data.get("descriptif") or {}
        for k in ["acier_total_kg", "acier_poteaux_kg", "acier_poutres_kg",
                  "acier_dalles_kg", "acier_voiles_kg", "acier_semelles_kg"]:
            assert k in d, f"Champ manquant: {k}"
            assert isinstance(d[k], (int, float)), f"{k} pas numérique: {d[k]}"
        ratios = d.get("acier_ratios") or {}
        assert ratios.get("semelles") == 40
        assert ratios.get("poteaux") == 130
        assert ratios.get("poutres") == 140
        assert ratios.get("dalles") == 85
        assert ratios.get("voiles") == 60

    def test_acier_total_matches_sum(self, headers, plan_id):
        r = requests.get(
            f"{BASE_URL}/api/takeoffs/{plan_id}/analysis/0?wall_height=2.5",
            headers=headers, timeout=60)
        d = r.json().get("descriptif") or {}
        total = d["acier_total_kg"]
        s = sum(d[k] for k in ["acier_poteaux_kg", "acier_poutres_kg",
                               "acier_dalles_kg", "acier_voiles_kg",
                               "acier_semelles_kg"])
        # rounded to 1 decimal, allow small delta
        assert abs(total - round(s, 1)) < 0.2, f"total={total} sum={s}"


# ---- Beton XLSX export ----
class TestBetonXlsx:
    def test_export_beton_xlsx_ok(self, headers, plan_id):
        r = requests.get(
            f"{BASE_URL}/api/takeoffs/{plan_id}/structure/0/beton-xlsx"
            f"?wall_height=2.5&dalle_ep=0.2&voile_ep=0.2",
            headers=headers, timeout=60)
        assert r.status_code == 200, r.text
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml.sheet" in ct or "openxmlformats" in ct, ct
        assert len(r.content) > 500
        wb = load_workbook(io.BytesIO(r.content))
        names = wb.sheetnames
        assert "Métré béton" in names, names
        assert "Poteaux" in names, names
        assert "Poutres" in names, names

    def test_export_beton_xlsx_404(self, headers):
        r = requests.get(
            f"{BASE_URL}/api/takeoffs/plan_does_not_exist/structure/0/beton-xlsx",
            headers=headers, timeout=30)
        assert r.status_code == 404


# ---- AI refine (best-effort — IA can time out) ----
class TestAiRefine:
    def test_ai_refine_empty_message_400(self, headers, plan_id):
        r = requests.post(
            f"{BASE_URL}/api/takeoffs/{plan_id}/structure/0/ai-refine",
            headers=headers, json={"message": "", "history": []}, timeout=30)
        assert r.status_code == 400

    def test_ai_refine_returns_structure(self, headers, plan_id):
        r = requests.post(
            f"{BASE_URL}/api/takeoffs/{plan_id}/structure/0/ai-refine",
            headers=headers,
            json={"message": "ignore les traits de cote", "history": []},
            timeout=120)
        assert r.status_code == 200, r.text
        data = r.json()
        assert "ok" in data
        assert "structure" in data
        # ok True or False, but structure must be present
        assert isinstance(data.get("structure"), dict)


# ---- Documents PDF with lang param ----
class TestDocumentPdfLang:
    @pytest.fixture(scope="class")
    def quote_id(self, headers):
        # Fetch any existing quote or create one
        r = requests.get(f"{BASE_URL}/api/documents?doc_type=quote",
                         headers=headers, timeout=30)
        if r.status_code == 200:
            data = r.json()
            items = data.get("documents", []) if isinstance(data, dict) else data
            if items:
                return items[0]["doc_id"]
        # Fallback: create a quote
        payload = {
            "doc_type": "quote",
            "client_name": "TEST Client Iter93",
            "lines": [{"designation": "TEST ligne", "quantity": 1, "unit": "u",
                       "unit_price": 100, "vat_rate": 20}]
        }
        c = requests.post(f"{BASE_URL}/api/documents", headers=headers,
                          json=payload, timeout=30)
        assert c.status_code in (200, 201), c.text
        return c.json()["doc_id"]

    def test_pdf_no_lang(self, headers, quote_id):
        r = requests.get(f"{BASE_URL}/api/documents/{quote_id}/pdf",
                         headers=headers, timeout=60)
        assert r.status_code == 200, r.text
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert r.content[:4] == b"%PDF"

    def test_pdf_lang_fr(self, headers, quote_id):
        r = requests.get(f"{BASE_URL}/api/documents/{quote_id}/pdf?lang=fr",
                         headers=headers, timeout=60)
        assert r.status_code == 200, r.text
        assert r.headers.get("content-type", "").startswith("application/pdf")
        assert r.content[:4] == b"%PDF"
