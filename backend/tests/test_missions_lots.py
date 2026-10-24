"""Opções « térmica », « estrutura », « mapa de quantidades / caderno de encargos » e « pré-certificado REH » do programa
IA de opções (.pt), mais a leitura da planta (depuis_plan PT). Base isolada, nenhuma chamada a modelo, dados fictícios."""
import base64
import copy
import importlib
import io

import pytest
from fastapi import HTTPException

from engines import missions as M
from engines import mapa_caderno as MC
from engines.calc_lots import depuis_plan as D
from routes import agents as A, missions as R
from test_missions import _run
from test_platform_supervisor import USER, Database

PLAN = {"plan_id": "p1", "user_id": "owner", "filename": "moradia.pdf", "calibrations": {"0": {"pixels_per_meter": 100}},
        "rooms": [{"room_id": "a", "page_index": 0, "name": "Sala", "points": [[0, 0], [600, 0], [600, 500], [0, 500]]},
                  {"room_id": "b", "page_index": 0, "name": "Quarto 1", "points": [[620, 0], [1020, 0], [1020, 500], [620, 500]]},
                  {"room_id": "g", "page_index": 0, "name": "Garagem", "points": [[1040, 0], [1540, 0], [1540, 500], [1040, 500]]}],
        "structure": {"0": {"openings": [{"kind": "fenetre", "x1": 100, "y1": -10, "x2": 300, "y2": -10},
                                         {"kind": "porte", "x1": 600, "y1": 200, "x2": 620, "y2": 290},
                                         {"kind": "porte_fenetre", "x1": 800, "y1": 510, "x2": 980, "y2": 510}]}}}


@pytest.fixture
def db(monkeypatch):
    database = Database()
    database.users.rows = [copy.deepcopy(USER)]
    for mod in (A, R, importlib.import_module("routes.chantiers"), importlib.import_module("routes.mep")):
        monkeypatch.setattr(mod, "db", database)
    access = importlib.import_module("engines.access")

    async def level(user):
        return "pro"
    monkeypatch.setattr(access, "compute_level", level)
    monkeypatch.setenv("MISSION_SYNTHESIS", "0")
    database.plans_takeoff.rows = [copy.deepcopy(PLAN)]
    return database


def _ctx(keys, chantier_id=None, results=None):
    run = {**_run(keys), "run_id": "r1", "chantier_id": chantier_id, "plan_status": [{"plan_id": "p1", "filename": "moradia.pdf"}]}
    return {"run": run, "user": USER, "cache": {}, "results": results or {}, "body": A.ExecIn()}


# ------------------------------------------------------------------ leitura da planta

def test_divisoes_from_plan_exterior_walls_and_openings():
    divisoes, hyp = D.divisoes_termicas(D.niveaux(PLAN))
    assert [d["nome"] for d in divisoes] == ["Sala", "Quarto 1"]                   # garagem não aquecida excluída
    assert [d["tipo"] for d in divisoes] == ["sala", "quarto"]
    from engines.calc_lots.aquecimento import T_INT
    assert all(d["tipo"] in T_INT for d in divisoes)
    sala = {w["tipo"]: w for w in divisoes[0]["paredes"]}
    assert sala["parede"]["area_m2"] == pytest.approx(17 * 2.5 - 2 * 1.2)            # 3 fachadas, parede meeira excluída
    assert sala["envidracado"]["area_m2"] == 2.4 and sala["pavimento"]["lado"] == "solo" and sala["cobertura"]["lado"] == "enu"
    assert all(w["U"] is None for w in divisoes[0]["paredes"])                        # U deixados ao motor (hipótese assinalada)
    vidros = [w for w in divisoes[1]["paredes"] if w["tipo"] == "envidracado"]
    assert len(vidros) == 1 and vidros[0]["area_m2"] == pytest.approx(1.8 * 2.1, abs=0.01)   # porta-janela contada uma vez
    assert (divisoes[0]["comprimento_m"], divisoes[0]["largura_m"]) == (6.0, 5.0)
    assert any("caixilharias" in h for h in hyp)


def test_structure_geometry_from_plan():
    geo, hyp = D.geometria_estrutura(D.niveaux(PLAN), "moradia")
    assert geo["outer_L_m"] == pytest.approx(15.8) and geo["outer_W_m"] == pytest.approx(5.4)
    assert {o["facade"] for o in geo["openings"]} == {"frente", "tras"}
    assert {o["label"] for o in geo["openings"]} == {"janela", "porta-janela"}
    assert geo["levels"][0]["storey_height_m"] == 2.7 and any("retângulo envolvente" in h for h in hyp)


def test_reh_inputs_from_plan():
    p, hyp = D.dados_reh(D.niveaux(PLAN))
    assert p["area_util_m2"] == 50 and p["tipologia"] == 1 and p["pisos"] == 1 and p["area_envidracada_m2"] == pytest.approx(6.2)
    assert any("T1" in h for h in hyp)


def test_plan_without_rooms_is_refused():
    from engines.mep.adapter import AdapterError
    with pytest.raises((ValueError, AdapterError), match="Nenhuma divisão"):
        D.niveaux({**PLAN, "rooms": []})


def test_concelho_from_city_or_address():
    assert R.concelho_do_processo({"city": "Braga"}) == "Braga"
    assert R.concelho_do_processo({"city": "", "address": "Rua Exemplo 10, 4700-000 Braga"}) == "Braga"
    assert R.concelho_do_processo({"city": "Cidade Inventada"}) is None


# ------------------------------------------------------------------ térmica e estrutura

@pytest.mark.asyncio
async def test_thermique_without_concelho_uses_default_climate_and_says_so(db):
    db.chantiers.rows = [{"chantier_id": "c1", "user_id": "owner", "city": ""}]
    res = await R._exec_thermique(_ctx(["thermique"], "c1"))
    n = res["notes"][0]
    assert n["concelho"] is None and n["theta_e"] == 0
    assert any("concelho" in h for h in n["hipoteses"])


@pytest.mark.asyncio
async def test_thermique_and_structure_store_pdf_notes(db):
    db.chantiers.rows = [{"chantier_id": "c1", "user_id": "owner", "city": "Porto"}]
    ctx = _ctx(["thermique", "structure"], "c1")
    db.agent_runs.rows = [copy.deepcopy(ctx["run"])]
    result = await R.run_mission(ctx["run"], A.ExecIn(), USER)
    assert [s["status"] for s in result["steps"]] == ["done", "done"], result["steps"]
    th = result["steps"][0]["result"]["notes"][0]
    assert th["concelho"] == "Porto" and th["perdas_w"] > 0 and th["bomba_calor_kw"] and th["circuitos"] >= 2
    assert th["pdf"].startswith("/missions/notes/")
    st = result["steps"][1]["result"]["notes"][0]
    assert st["implantacao_m"] == [15.8, 5.4] and st["linteis"] == 2
    assert any("sísmica" in w for w in st["avisos"])
    assert len(db.mission_notes.rows) == 2 and all(r["user_id"] == "owner" and r["run_id"] == "r1" for r in db.mission_notes.rows)
    resp = await R.note_file(th["note_id"], USER)
    assert resp.body[:4] == b"%PDF" and "Nota" in resp.headers["content-disposition"]
    with pytest.raises(HTTPException):
        await R.note_file(th["note_id"], {**USER, "user_id": "outro"})


# ------------------------------------------------------------------ mapa de quantidades / caderno de encargos

ORC = {"doc_id": "d1", "user_id": "owner", "title": "Orçamento parcial — Moradia Exemplo",
       "lines": [{"kind": "line", "designation": "Parede em alvenaria de tijolo 20 cm", "unit": "m2", "qty": 40, "unit_price_ht": 55},
                 {"kind": "line", "designation": "Azulejos", "description": "R/C · Azulejo (parede) — cozinha", "unit": "m2", "qty": 12.5,
                  "unit_price_ht": 0, "price_missing": True},
                 {"kind": "line", "designation": "Tomada 16 A", "unit": "u", "qty": 10, "unit_price_ht": 35},
                 {"kind": "line", "designation": "Mão de obra de aplicação", "description": "R/C · Pintura — paredes", "unit": "m2",
                  "qty": 80, "unit_price_ht": 6},
                 {"kind": "section", "designation": "R/C"}]}


def test_lotes_by_trade():
    lotes = {l["lote"]: l for l in MC.por_lote(ORC["lines"])}
    assert list(lotes) == ["Estrutura e alvenarias", "Impermeabilizações e revestimentos cerâmicos", "Estuques e pinturas",
                           "Instalações elétricas"]                                 # azulejo (parede) e pintura não vão para a estrutura
    assert [l["designacao"] for l in lotes["Estrutura e alvenarias"]["linhas"]] == ["Parede em alvenaria de tijolo 20 cm"]
    assert lotes["Impermeabilizações e revestimentos cerâmicos"]["por_orcamentar"] == 1
    assert any("RTIEBT" in r for r in lotes["Instalações elétricas"]["refs"])
    assert any("Eurocódigos" in r for r in lotes["Estrutura e alvenarias"]["refs"])
    assert MC.lote_de({"designation": "Sal de degelo"})[0] == "Diversos"
    assert MC.lote_de({"designation": "Tubo de queda PVC", "description": "drenagem de águas residuais"})[1][0].startswith("DR 23/95")


def test_mapa_xlsx_content():
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(MC.mapa_xlsx(MC.por_lote(ORC["lines"]), "Moradia Exemplo")))
    ws = wb["Mapa de quantidades"]
    rows = [[c for c in r] for r in ws.iter_rows(values_only=True)]
    assert rows[0][0] == "Mapa de quantidades — Moradia Exemplo"
    assert tuple(rows[3]) == ("N.º", "Designação", "Unid.", "Quantidade", "Preço unitário (€)", "Total (€)")
    flat = [r for r in rows if r and r[1]]
    assert any(r[1] == "Azulejos" and r[4] == "por orçamentar" and r[2] == "m²" for r in flat)
    assert rows[-1][1] == "TOTAL (sem IVA)" and rows[-1][5] == 40 * 55 + 10 * 35 + 80 * 6
    assert not any("HT" in str(v) or "DPGF" in str(v) for r in rows for v in r if v)


def test_caderno_docx_content():
    from docx import Document
    doc = Document(io.BytesIO(MC.caderno_docx(MC.por_lote(ORC["lines"]), "Moradia Exemplo")))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert doc.paragraphs[0].text == "Caderno de encargos — Moradia Exemplo"
    for s in ("2. Impermeabilizações e revestimentos cerâmicos", "Documentos de referência", "RTIEBT", "Eurocódigos estruturais",
              "Quantidade indicativa: 12,5 m²", "a completar pelo autor do caderno de encargos"):
        assert s in text, s
    assert "CCTP" not in text and "DTU" not in text and "artigo" not in text.lower().replace("artigos", "")


@pytest.mark.asyncio
async def test_dpgf_from_draft_quote(db):
    db.pro_documents.rows = [copy.deepcopy(ORC)]
    res = await R._exec_dpgf(_ctx(["devis", "dpgf"], results={"devis": {"doc_id": "d1"}}))
    assert res["kind"] == "files" and [f["ext"] for f in res["files"]] == ["xlsx", "docx"]
    assert res["total"] == 40 * 55 + 10 * 35 + 80 * 6 and res["files"][0]["title"] == "Mapa de quantidades — Moradia Exemplo"
    resp = await R.note_file(res["files"][0]["note_id"], USER)
    assert resp.body[:2] == b"PK" and "spreadsheet" in resp.media_type
    resp = await R.note_file(res["files"][1]["note_id"], USER)
    assert resp.body[:2] == b"PK" and "wordprocessing" in resp.media_type


@pytest.mark.asyncio
async def test_dpgf_without_quote_fails(db):
    with pytest.raises(HTTPException, match="rascunho de orçamento"):
        await R._exec_dpgf(_ctx(["dpgf"]))


# ------------------------------------------------------------------ pré-certificado REH

@pytest.mark.asyncio
async def test_reh_needs_concelho(db):
    db.chantiers.rows = [{"chantier_id": "c1", "user_id": "owner", "city": "Cidade Inventada"}]
    with pytest.raises(HTTPException, match="Concelho"):
        await R._exec_dpe(_ctx(["dpe"], "c1"))


@pytest.mark.asyncio
async def test_reh_note_from_plan(db):
    db.chantiers.rows = [{"chantier_id": "c1", "user_id": "owner", "city": "Évora"}]
    res = await R._exec_dpe(_ctx(["dpe"], "c1"))
    n = res["notes"][0]
    assert res["lot"] == "dpe" and n["area_util_m2"] == 50 and n["classe"] in ("A+", "A", "B", "B-", "C", "D", "E", "F")
    assert "ADENE" in n["aviso"] and n["pdf"].startswith("/missions/notes/")
    pdf = base64.b64decode(db.mission_notes.rows[0]["data_b64"])
    assert pdf[:4] == b"%PDF"


# ------------------------------------------------------------------ partilha das notas e ficheiros

@pytest.mark.asyncio
async def test_share_includes_notes_and_files(db, monkeypatch):
    monkeypatch.setenv("PUBLIC_BACKEND_URL", "https://api.exemplo.pt")
    db.mission_notes.rows = [{"note_id": "n1", "user_id": "owner", "run_id": "r1", "title": "Mapa de quantidades — X", "ext": "xlsx",
                              "data_b64": base64.b64encode(b"PK\x03\x04").decode(), "created_at": "x"},
                             {"note_id": "n2", "user_id": "owner", "run_id": "outro", "title": "Fora da partilha", "ext": "pdf",
                              "data_b64": "JVBERi0=", "created_at": "x"}]
    res = await R._exec_partage(_ctx(["partage"]))
    assert res["files"] == ["Mapa de quantidades — X"] and res["active"] is False
    token = db.mission_shares.rows[0]["token"]
    await R.activate_share(res["share_id"], R.ShareActivateIn(days=7), USER)
    page = await R.public_share(token)
    assert b"(XLSX)" in page.body and "Fora da partilha".encode() not in page.body
    f = await R.public_share_file(token, "n1")
    assert f.body[:2] == b"PK" and f.headers["x-robots-tag"] == "noindex"
    with pytest.raises(HTTPException):
        await R.public_share_file(token, "n2")


def test_new_options_are_in_catalog_order():
    keys = list(M.OPTIONS)
    assert keys.index("plan_plomberie") < keys.index("thermique") < keys.index("structure") < keys.index("devis") \
        < keys.index("dpgf") < keys.index("dpe") < keys.index("planning")
    assert M.OPTIONS["dpgf"]["depends_on"] == ["devis"] and M.OPTIONS["dpe"]["needs"] == ["plans", "chantier"]
    txt = " ".join(o["name"] + o["desc"] + o["deliverable"] + o["prompt"] for o in M.OPTIONS.values())
    for fr in ("NF C 15-100", "DTU", "DPE", "CCTP", "Devis", "chantier", "Métré"):
        assert fr not in txt


# ------------------------------------------------------------------ síntese facultativa (offline / demo)

@pytest.mark.asyncio
async def test_synthesis_offline_and_demo_quota(monkeypatch):
    monkeypatch.setenv("MISSION_SYNTHESIS", "1")
    ai = importlib.import_module("engines.ai")
    seen = {}

    async def down(*a, **k):
        raise OSError("network unreachable")
    monkeypatch.setattr(ai, "_complete", down)
    assert await M.synthesize("Resume.", {"x": 1}) is None
    guards = importlib.import_module("core.demo_guards")

    async def quota(user):
        raise HTTPException(429, "quota")
    monkeypatch.setattr(guards, "consume_demo_ai_quota", quota)
    assert await M.synthesize("Resume.", {"x": 1}, user={"user_id": "d", "is_demo": True}) is None

    async def complete(system, text, **k):
        seen.update(system=system, text=text)
        return "  Resultado: 2 kW.  "
    monkeypatch.setattr(ai, "_complete", complete)
    assert await M.synthesize("Resume.", {"perdas_w": 2000}) == "Resultado: 2 kW."
    assert '"perdas_w": 2000' in seen["text"] and "português europeu" in seen["system"]


# ------------------------------------------------------------------ medição e leitura das plantas

@pytest.mark.asyncio
async def test_owned_plans_load_rooms_so_room_only_plans_are_ready(db, monkeypatch):
    seen = []
    real = db.plans_takeoff.find
    monkeypatch.setattr(db.plans_takeoff, "find", lambda q, proj=None: seen.append(proj) or real(q, proj))
    out = await R.propose(R.MissionIn(options=["structure"], plan_ids=["p1"]), USER)
    assert seen[0].get("rooms") == 1
    assert out["run"]["plan_status"][0]["rooms"] == 3 and out["run"]["plan_status"][0]["ready"] is True
    assert not out["run"]["warnings"]


@pytest.mark.asyncio
async def test_metre_summary_adds_rooms_from_the_plan(db):
    db.plans_takeoff.rows = [{**copy.deepcopy(PLAN), "measurements": [{"measurement_id": "m", "poste": "peinture", "openings": [{}]}],
                              "rooms": [{**r, "surface_m2": 12} for r in PLAN["rooms"]]}]
    ctx = _ctx(["metre"])
    ctx["run"]["plan_status"] = [{"plan_id": "p1", "filename": "moradia.pdf", "level": "RDC", "measures": 0, "priced": 0, "openings": 0}]
    out = await R._exec_metre(ctx)
    assert out["kind"] == "summary" and out["total_measures"] == 1 and out["total_rooms"] == 3 and out["total_surface_m2"] == 36


@pytest.mark.parametrize("designation,description,lote", [
    ("Tubo PEX + placas isolantes", "R/C · Piso radiante (hidráulico)", "Betonilhas e piso radiante"),
    ("Tinta (2 demãos)", "Pintura — paredes e tetos", "Estuques e pinturas"),
    ("Isolante de lã mineral", "Piso 1 · Divisória (pladur)", "Divisórias, forros e tetos em gesso cartonado"),
    ("Isolante", "R/C · Isolamento — paredes exteriores", "Isolamentos térmicos"),
    ("Cola para azulejo", "R/C · Revestimento cerâmico (pavimento)", "Impermeabilizações e revestimentos cerâmicos"),
    ("Remoção do entulho", "Desmontagem / demolição", "Demolições e trabalhos preparatórios"),
    ("Janela em alumínio com corte térmico", "", "Carpintarias, caixilharias e escadas"),
    ("Laje aligeirada de vigotas", "", "Estrutura e alvenarias"),
])
def test_trade_detection_prefers_specific_words(designation, description, lote):
    assert MC.lote_de({"designation": designation, "description": description})[0] == lote


# ------------------------------------------------------------------ pré-certificado: estado da obra e sistemas

@pytest.mark.parametrize("chantier,obra", [
    ({"name": "Moradia nova T3", "notes": ""}, "nova"),
    ({"name": "Lote 4", "notes": "Construção nova, chave na mão"}, "nova"),
    ({"name": "Lote 4", "work_type": "new_build"}, "nova"),
    ({"name": "Remodelação apartamento", "notes": "obra em edifício de 1975"}, "existente"),
])
def test_estado_obra_from_chantier(chantier, obra):
    from datetime import date
    o, ano, origem = R.estado_obra(chantier)
    assert o == obra and (ano == date.today().year if obra == "nova" else ano is None) and origem


@pytest.mark.asyncio
async def test_reh_new_house_reuses_thermal_heat_pump(db):
    db.chantiers.rows = [{"chantier_id": "c1", "user_id": "owner", "city": "Coimbra", "name": "Moradia nova exemplo"}]
    ctx = _ctx(["thermique", "dpe"], "c1")
    db.agent_runs.rows = [copy.deepcopy(ctx["run"])]
    result = await R.run_mission(ctx["run"], A.ExecIn(), USER)
    assert [s["status"] for s in result["steps"]] == ["done", "done"], result["steps"]
    n = result["steps"][1]["result"]["notes"][0]
    assert n["obra"] == "nova" and n["aquecimento"] == "bomba_calor" and n["aquecimento_origem"] == "opção térmica deste programa"
    assert any("SCOP" in c for c in n["a_confirmar_perito"]) and any("AQS" in c for c in n["a_confirmar_perito"])
    assert any("PERITO QUALIFICADO" in h for h in n["hipoteses"]) and any("piso radiante" in h for h in n["hipoteses"])
    assert n["classe"] not in ("E", "F")


@pytest.mark.asyncio
async def test_reh_reuses_previous_thermal_note_of_the_process(db):
    db.chantiers.rows = [{"chantier_id": "c1", "user_id": "owner", "city": "Coimbra", "notes": "construção nova"}]
    db.mission_notes.rows = [{"note_id": "old", "user_id": "owner", "chantier_id": "c1", "kind": "thermique", "created_at": "2026-10-01",
                              "summary": {"plan_id": "p1", "bomba_calor_kw": 6, "cop_theta_e": 2.9, "emissores": "piso_radiante"}},
                             {"note_id": "alheia", "user_id": "outro", "chantier_id": "c1", "kind": "thermique", "created_at": "2026-10-20",
                              "summary": {"plan_id": "p1", "bomba_calor_kw": 16}}]
    res = await R._exec_dpe(_ctx(["dpe"], "c1"))
    n = res["notes"][0]
    assert n["aquecimento"] == "bomba_calor" and n["aquecimento_origem"].startswith("nota térmica anterior")
    assert any("bomba de calor ar/água 6 kW com piso radiante" in h for h in n["hipoteses"])
    assert any("COP indicativo a θe: 2,90" in h for h in n["hipoteses"])


@pytest.mark.asyncio
async def test_reh_existing_without_thermal_flags_defaults(db):
    db.chantiers.rows = [{"chantier_id": "c1", "user_id": "owner", "city": "Évora", "name": "Remodelação"}]
    n = (await R._exec_dpe(_ctx(["dpe"], "c1")))["notes"][0]
    assert n["obra"] == "existente" and n["aquecimento"] == "resistencia_eletrica" and n["aquecimento_origem"] is None
    assert "sistema de aquecimento e respetiva eficiência" in n["a_confirmar_perito"]
    assert any("1980 por defeito" in h for h in n["hipoteses"])
