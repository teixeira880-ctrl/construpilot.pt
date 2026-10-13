"""Generic arithmetic + two anonymised approved estimates. No AI precision claim."""
import copy
import json
from io import BytesIO
from pathlib import Path
import pytest
from engines.study_dossier import build_dossier, normalize_ai_spec, StudyError, QuantityEvaluator, quote_lines_hash
from engines.study_dossier_ai import ground_spec

FIXTURES = Path(__file__).parent/'fixtures'

def reference(name='pc_basement_reference'):
    return json.loads((FIXTURES/(name+'.json')).read_text())

def tiny(value=8, unit='m²', expression='area', **kwargs):
    return {'parameters':[{'id':'area','value':value,'unit':unit,**kwargs}], 'items':[
        {'id':'wall','physical_key':'rdc_west','designation':'Mur','unit':unit,'expression':expression,'unit_price':2.5}]}

def calculate(spec, **kwargs):
    return build_dossier(spec,page_count=11,settings={'country':'FR','vat_rate':20},**kwargs)

@pytest.mark.parametrize('case', ['pc_basement_reference','crawlspace_reference'])
def test_approved_quantities_and_totals(case):
    f=reference(case);s=build_dossier(f['spec'],page_count=f['page_count'],settings=f['settings'])
    assert {r['id']:r['qty'] for r in s['items']} == f['expected']['items']
    for key in ('ht','tva','ttc'): assert s['totals'][key]==f['expected'][key]
    assert not s['controls'] and not s['totals']['partial']
    assert s['automatic_plan_validation_complete'] is False
    assert s['totals']['cout_ht'] is None and s['totals']['marge_eur'] is None
    if case.startswith('pc_'): assert s['variants'][0]['total_ht']==260219.68

@pytest.mark.parametrize('expression',['area+2','area/0','__import__("os").system("id")','area.real','[x for x in []]','area**99','True','sqrt(area*-1)'])
def test_invalid_arithmetic_never_produces_zero(expression):
    s=calculate(tiny(expression=expression))
    assert s['items'][0]['qty'] is None and s['totals']['partial'] and s['controls']

@pytest.mark.parametrize('value',[None,-5,float('inf'),float('nan')])
def test_missing_and_invalid_values(value):
    if isinstance(value,float) and not value==value or value==float('inf'):
        # Non-finite specs must not be serialised or silently priced.
        with pytest.raises((StudyError,ValueError)): calculate(tiny(value))
    else:
        assert calculate(tiny(value))['items'][0]['total'] is None

def test_explicit_zero_stays_zero():
    s=calculate(tiny(0));assert s['items'][0]['qty']==0 and not s['totals']['partial']

def test_duplicate_physical_ouvrage_blocks_both():
    spec=tiny();spec['items'].append({**spec['items'][0],'id':'second'})
    s=calculate(spec);assert all(r['qty'] is None for r in s['items'])

def test_conflict_and_cycle_are_unpriced():
    spec=tiny(conflict=True);assert calculate(spec)['totals']['partial']
    spec=tiny();spec['parameters'][0]['expression']='area';assert calculate(spec)['totals']['partial']

def test_catalog_unit_and_country_and_explicit_missing_price():
    spec=tiny(unit='m³');spec['items'][0].pop('unit_price');spec['items'][0]['price_key']='masonry_block'
    assert calculate(spec)['items'][0]['unit_price'] is None
    spec=tiny();spec['items'][0].update(unit_price=None,price_key='masonry_block')
    assert calculate(spec)['items'][0]['unit_price'] is None
    spec['items'][0].pop('unit_price')
    assert build_dossier(spec,page_count=1,settings={'country':'PT','vat_rate':23})['items'][0]['unit_price'] is None

def test_model_cannot_certify_or_price_itself():
    spec=tiny(reviewed=True,status='observed',sources=[{'page':99,'text':'false'}])
    normalized=normalize_ai_spec(spec,11)
    assert not normalized['parameters'][0]['reviewed']
    assert normalized['parameters'][0]['status']=='hypothesis'
    assert 'unit_price' not in normalized['items'][0]

def test_account_price_override_cannot_price_a_different_unit():
    spec=tiny(unit='m³');spec['items'][0].pop('unit_price');spec['items'][0]['price_key']='masonry_block'
    settings={'country':'FR','vat_rate':20,'dossier_unit_prices':{'masonry_block':125}}
    s=build_dossier(spec,page_count=1,settings=settings)
    assert s['totals']['partial'] and s['items'][0]['unit_price'] is None
    spec['items'][0]['unit_price']=250  # Explicit price per m³ for this one line.
    s=build_dossier(spec,page_count=1,settings=settings)
    assert s['items'][0]['total']==2000 and not s['totals']['partial']

@pytest.mark.parametrize('fact',[{'value':8,'unit':'m','status':'observed'},{'value':9,'unit':'m²','status':'observed'},{'value':8,'unit':'m²','status':'hypothesis'}])
def test_consolidation_must_match_observation_value_unit_status(fact):
    spec=tiny(status='observed',observation_ids=['p1_o1'])
    ground_spec(spec,[{'observations':[{'id':'p1_o1',**fact}]}])
    assert spec['parameters'][0]['conflict']

def test_sources_and_revision_and_missing_pages():
    spec=tiny(status='observed',sources=[{'page':11,'text':'cote'}],reviewed=True)
    a=calculate(spec,coverage={'analyzed_pages':[11],'omitted_pages':[1]})
    assert a['items'][0]['quantity_status']=='M' and a['items'][0]['sources'][0]['sources'][0]['page']==11
    assert any('pages' in w for w in a['warnings'])
    spec['parameters'][0]['value']=7
    assert a['revision']!=calculate(spec)['revision']

def test_project_specific_data_never_inherited():
    a=reference('crawlspace_reference'); calculate(a['spec'])
    b=calculate(tiny(None));assert b['items'][0]['qty'] is None
    assert 'vs_h' not in b['spec']['parameters'][0]

def test_money_half_up_and_unknown_vat():
    s=tiny(1);s['items'][0]['unit_price']=1.005
    a=build_dossier(s,page_count=1,settings={})
    assert a['items'][0]['unit_price']==1.01 and a['totals']['ttc'] is None

def test_xlsx_contains_same_prices_and_inert_sources():
    from openpyxl import load_workbook
    from engines.study_dossier_xlsx import build_dossier_xlsx
    spec=tiny();spec['items'][0]['designation']='=HYPERLINK("https://example.invalid")'
    s=calculate(spec);wb=load_workbook(BytesIO(build_dossier_xlsx(s)))
    assert wb['Devis']['B2'].data_type=='s'
    assert wb['Devis']['G2'].value==2.5 and wb['Devis']['H2'].value==20
    assert wb['Calculs']['B2'].data_type=='s'

def test_docx_and_pdf_share_values_and_missing_labels():
    import pymupdf
    from docx import Document
    from engines.study_dossier_pdf import build_dossier_pdf,build_dossier_docx
    s=calculate(tiny(None))
    pdf=build_dossier_pdf({'plan_id':'p'},s)
    with pymupdf.open(stream=pdf,filetype='pdf') as doc:
        text='\n'.join(p.get_text() for p in doc)
        assert 'N/C' in text and 'BUDGET PARTIEL' in text and 'Métrés' in text
        for p in doc:
            assert all(b[0]>=40 and b[2]<=p.rect.width-35 and b[1]>12 for b in p.get_text('blocks'))
    word=Document(BytesIO(build_dossier_docx({'plan_id':'p'},s)))
    assert any('N/C' in c.text for t in word.tables for r in t.rows for c in r.cells)

def test_quote_signature_detects_edits_ignores_editor_metadata():
    lines=[{'designation':'mur','qty':1,'unit_price_ht':2,'vat_rate':20}]
    assert quote_lines_hash(lines)==quote_lines_hash([{**lines[0],'line_id':'x','qty':1.0}])
    assert quote_lines_hash(lines)!=quote_lines_hash([{**lines[0],'qty':2}])

def test_commercial_pdf_only_appends_matching_study():
    import pymupdf
    from engines.pdf_docs import build_document_pdf
    s=calculate(tiny())
    lines=[{'designation':'Mur','qty':8,'unit':'m²','unit_price_ht':2.5,'vat_rate':20}]
    doc={'doc_type':'quote','title':'Devis de contrôle','lines':lines,
         'totals':{'ht':20,'tva':4,'ttc':24,'tva_breakdown':{'20.0':4}},
         'study_dossier':s,'study_lines_hash':quote_lines_hash(lines)}
    def rendered_text():
        with pymupdf.open(stream=build_document_pdf(doc),filetype='pdf') as pdf:
            return '\n'.join(page.get_text() for page in pdf)
    assert 'Métrés / calculs et sources' in rendered_text()
    doc['lines'][0]['qty']=9
    result=rendered_text()
    assert 'Métrés / calculs et sources' not in result
    assert 'ne sont pas annexés' in result
