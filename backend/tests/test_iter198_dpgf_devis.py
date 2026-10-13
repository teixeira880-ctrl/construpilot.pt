"""DPGF : lecture multi-onglets/lots, consignes, CCTP, comparaison de prix, passage en devis.

Exécution : python backend/tests/test_iter198_dpgf_devis.py
"""
import io
import os
import re
import sys

import openpyxl
import requests

with open("/app/frontend/.env") as f:
    BASE = re.search(r"REACT_APP_BACKEND_URL=(\S+)", f.read()).group(1)
API = BASE.rstrip("/") + "/api"
EMAIL = os.environ.get("TEST_PRO_EMAIL", "pro@construpilot.fr")
PWD = os.environ.get("TEST_PRO_PASSWORD", "Pro2026!")

s = requests.Session()
ok = fail = 0


def check(label, cond, detail=""):
    global ok, fail
    if cond:
        ok += 1
        print(f"✅ {label}")
    else:
        fail += 1
        print(f"❌ {label} — {detail}")


def make_dpgf() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "DPGF"
    ws.append(["Bordereau — Villa TEST iter198"])
    ws.append(["Désignation", "Unité", "Quantité", "Prix unitaire HT", "Montant HT"])
    ws.append(["01 — MACONNERIE", None, None, None, None])
    ws.append(["Mur en blocs béton 20 cm hourdés au mortier", "m2", 45, None, None])
    ws.append(["Chape ciment dosée à 300 kg/m3", "m2", 60, None, None])
    ws.append(["03 — CARRELAGE", None, None, None, None])
    ws.append(["Carrelage grès cérame 60x60 collé", "m2", 60, None, None])
    ws.append(["Plinthes carrelage assorties", "ml", 40, None, None])
    ws.append(["05 — REVETEMENTS DE SOL", None, None, None, None])
    ws.append(["Parquet contrecollé chêne collé plein", "m2", 35, None, None])
    ws.append(["Sous-total lot", None, None, None, None])
    ws2 = wb.create_sheet("PEINTURE")
    ws2.append(["Désignation", "Unité", "Quantité", "Prix unitaire HT", "Montant HT"])
    ws2.append(["07 — PEINTURE", None, None, None, None])
    ws2.append(["Peinture acrylique mate 2 couches sur murs", "m2", 120, None, None])
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


CCTP = ("CAHIER DES CLAUSES TECHNIQUES PARTICULIERES — Villa TEST\n"
        "ARTICLE 1 — MACONNERIE : les blocs béton seront de classe B40, hourdés au mortier "
        "bâtard dosé à 350 kg/m3. Chape ciment armée d'un treillis soudé, épaisseur minimale 5 cm.\n"
        "ARTICLE 2 — PEINTURE : peinture acrylique de marque Dulux Valentine, application en "
        "trois couches sur enduit repassé, finition mate velours, degré de finition C soigné.\n"
        "ARTICLE 3 — SOLS : les carrelages seront de classement U4P4E3C2, pose collée double "
        "encollage sur chape désolidarisée. Les parquets contrecollés seront de classe 23/31.\n"
        "ARTICLE 4 — GENERALITES : l'entreprise prévoira le nettoyage quotidien du chantier, "
        "la protection des ouvrages livrés et le respect des normes NF DTU en vigueur. "
        "Les échafaudages et moyens de levage sont à la charge de l'entreprise.\n") * 3

r = s.post(f"{API}/auth/login", json={"email": EMAIL, "password": PWD}, timeout=30)
check("login pro", r.status_code == 200, r.text[:200])
if r.status_code != 200:
    sys.exit(1)

files = {"file": ("dpgf_test_iter198.xlsx", make_dpgf(),
                  "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"),
         "cctp": ("cctp_test_iter198.txt", CCTP.encode(), "text/plain")}
data = {"instructions": "ne chiffre pas le carrelage ni le parquet, marge 12 %, peinture Dulux",
        "margin_pct": "0", "analyse_cctp": "true", "market_compare": "true"}
r = s.post(f"{API}/pro/dpgf/import", files=files, data=data, timeout=240)
check("import DPGF + CCTP + consignes", r.status_code == 200, r.text[:400])
if r.status_code != 200:
    sys.exit(1)
res = r.json()
fid = res["file_id"]
lines = res["lines"]
print("   lots:", [x["lot"] for x in res["lots"]])

check("multi-onglets lus (DPGF + PEINTURE)",
      len({l["sheet"] for l in lines}) == 2, str({l["sheet"] for l in lines}))
check("lots détectés (4)", len(res["lots"]) == 4, str([x["lot"] for x in res["lots"]]))
check("lignes de titres et sous-totaux ignorés", res["line_count"] == 6, str(res["line_count"]))

carrelage = [l for l in lines if "arrelage" in l["designation"] or "linthes" in l["designation"]]
parquet = [l for l in lines if "arquet" in l["designation"]]
check("carrelage exclu par consigne (sans prix)",
      all(l["excluded"] and l["pu"] is None for l in carrelage), str(carrelage)[:300])
check("parquet exclu par consigne", all(l["excluded"] for l in parquet), str(parquet)[:300])
macon = [l for l in lines if "blocs" in l["designation"]][0]
check("maçonnerie chiffrée", (macon["pu"] or 0) > 0 and macon["total"] > 0, str(macon)[:300])
check("source de prix renseignée", macon["source"] in ("bibliotheque", "ia", "fichier"), str(macon["source"]))
check("marge 12 % détectée dans les consignes", res["directives"].get("margin_pct") == 12,
      str(res["directives"].get("margin_pct")))
check("total général = somme des lots chiffrés",
      abs(res["grand_total_ht"] - sum(x["total_ht"] for x in res["lots"])) < 0.05,
      str(res["grand_total_ht"]))
check("CCTP analysé (synthèse + prescriptions)",
      bool(res["cctp"].get("summary")) and len(res["cctp"].get("prescriptions") or []) > 0,
      str(res["cctp"])[:300])
check("prescriptions CCTP rattachées à au moins une ligne",
      any(l.get("cctp_note") for l in lines), "aucune note CCTP")
check("prix marché IA disponible pour comparaison",
      any(l.get("pu_ai") for l in lines), "aucun pu_ai")

# Excel rempli téléchargeable, exclusions mentionnées
x = s.get(f"{API}/pro/dpgf/file/{fid}")
check("Excel chiffré téléchargeable", x.status_code == 200 and x.content[:2] == b"PK", str(x.status_code))
wb = openpyxl.load_workbook(io.BytesIO(x.content), data_only=True)
vals = [str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row if c.value]
check("mention « Non chiffré — hors marché » écrite dans l'Excel",
      any("Non chiffré" in v for v in vals), "mention absente")

# Modification manuelle d'une ligne + réintégration d'une exclusion
lid = macon["line_id"]
res2 = s.patch(f"{API}/pro/dpgf/{fid}/line/{lid}", json={"pu": 99.0}).json()
m2 = [l for l in res2["lines"] if l["line_id"] == lid][0]
check("PU saisi manuellement retenu", m2["pu"] == 99.0 and m2["source"] == "manuel", str(m2)[:200])
cid = carrelage[0]["line_id"]
res3 = s.patch(f"{API}/pro/dpgf/{fid}/line/{cid}", json={"excluded": False}).json()
c3 = [l for l in res3["lines"] if l["line_id"] == cid][0]
check("ligne exclue réintégrable", not c3["excluded"], str(c3)[:200])
res4 = s.patch(f"{API}/pro/dpgf/{fid}/line/{cid}", json={"excluded": True}).json()
check("exclusion rétablie", [l for l in res4["lines"] if l["line_id"] == cid][0]["excluded"], "")

# Passage en devis
q = s.post(f"{API}/pro/dpgf/{fid}/quote", json={"vat_rate": 10, "full_descriptions": True})
check("passage en devis", q.status_code == 200, q.text[:300])
qd = q.json()
doc = s.get(f"{API}/documents/{qd['doc_id']}").json()
cats = {l.get("category") for l in doc["lines"]}
check("une catégorie par lot dans le devis", len(cats) >= 3, str(cats))
zero = [l for l in doc["lines"] if l["unit_price_ht"] == 0]
check("lignes non chiffrées conservées à 0 € avec mention",
      len(zero) == 3 and all("Non chiffré" in l["designation"] for l in zero), str(zero)[:300])
check("TVA 10 % appliquée", all(l["vat_rate"] == 10 for l in doc["lines"]), "TVA incorrecte")
check("total devis > 0 et hors lignes exclues", doc["totals"]["ht"] > 0, str(doc["totals"]["ht"]))
check("consignes et CCTP repris dans les notes du devis",
      "Consignes" in doc["notes"] and "hors marché" in doc["notes"], doc["notes"][:200])

# Nettoyage
s.delete(f"{API}/documents/{qd['doc_id']}")
s.delete(f"{API}/pro/dpgf/{fid}")
print(f"\n{ok} OK / {fail} KO")
sys.exit(1 if fail else 0)
