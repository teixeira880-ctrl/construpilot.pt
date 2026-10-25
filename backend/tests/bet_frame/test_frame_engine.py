"""Moteur ossature BA (bet_frame) : calculs de référence, contrôles, dossier complet FR / PT, détection du type, branchement Assistant."""
import asyncio
import copy
import io
import json
from pathlib import Path

import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from reportlab.lib.pagesizes import A3, landscape
from reportlab.pdfgen import canvas

from engines.bet_frame import calc, hook, mep, model, pipeline, reader
from engines.bet_frame.lang import use_lang

CASE = json.loads((Path(__file__).parents[2] / "engines/bet_frame/cases/immeuble_demo_pt.json").read_text(encoding="utf-8"))


def _mt(fck=25):
    return {"fck": fck, "fcd": fck / 1.5, "fctm": 0.30 * fck ** (2 / 3), "fyk": 500.0, "fyd": 500 / 1.15}


def test_trois_moments_deux_travees_egales():
    Ms = calc._three_moment([5.0, 5.0], [10.0, 10.0])
    assert Ms[0] == 0 and Ms[2] == 0 and abs(Ms[1] + 10 * 25 / 8) < 1e-9          # −wL²/8


def test_trois_moments_trois_travees():
    Ms = calc._three_moment([6.0, 6.0, 6.0], [12.0, 12.0, 12.0])
    assert abs(Ms[1] + 0.1 * 12 * 36) < 1e-6 and abs(Ms[2] - Ms[1]) < 1e-9        # −0,10 wL²


def test_flexion_simple_reference():
    r = calc.design_rect(100.0, 0.30, 0.45, _mt())
    assert r["ok"] and 5.3e-4 < r["As"] < 5.6e-4                                   # ≈ 5,4 cm²
    assert not calc.design_rect(800.0, 0.30, 0.45, _mt())["ok"]                    # μ > μlim


def test_poteau_compression_et_flexion():
    d = calc.design_column(0.40, 0.40, 1500.0, 40.0, 40.0, _mt())
    assert d["ok"] and d["MRdx"] > 40 and d["NRd"] > 1500
    assert not calc.design_column(0.30, 0.30, 4000.0, 0, 0, _mt())["ok"]           # ν > 0,85


def test_controles_bloquants():
    bad = copy.deepcopy(CASE)
    bad["grid"]["x"] = [0, 6, 5, 12]
    with use_lang("fr"):
        _, an = pipeline.validate(bad)
    assert any(a["code"] == "GRID-02" and a["level"] == model.BLOCK for a in an)
    res = pipeline.run(bad, "/tmp/_bf_bad")
    assert res["status"] == "MODELE_INVALIDE" and not res["files"]


def test_calculs_cas_immeuble():
    with use_lang("pt"):
        m = model.from_dict(copy.deepcopy(CASE))
        assert not model.check(m)
        r = calc.compute(m)
    assert len(r["columns"]) == 28 and len(r["floors"]) == 6
    base = max(c["segs"][0]["NEd"] for c in r["columns"])
    assert 2500 < base < 5000                                                       # 6 planchers × ~37 m² × ~13 kN/m²
    for c in r["columns"]:
        for s in c["segs"]:
            assert s["ok"] and s["nu"] <= 0.85 and s["rho"] <= 0.03
    for fl in r["floors"]:
        assert all(b["ok"] for b in fl["beams"])
        assert all(p["ok"] for p in fl["panels"] if not p["core"])
    assert r["basement"] and r["basement"]["t"] >= 0.25
    f = r["foundations"]
    assert all(p["sigma"] <= f["sigma_kpa"] * 1.001 for p in f["pads"])


def test_reseaux_colonnes():
    with use_lang("pt"):
        m = model.from_dict(copy.deepcopy(CASE))
        e = mep.compute(m)
    ris = {r["id"]: r for r in e["risers"]}
    assert ris["TQ1"]["sections"][0]["d"] >= 100                                     # chute avec WC
    af = ris["CAF1"]["sections"]
    assert af[0]["d"] >= af[-1]["d"] and af[0]["n"] == 10                            # diamètre décroissant vers le haut
    el = ris["CE"]["sections"][0]
    assert el["n"] == 20 and 0.3 < el["ks"] < 0.6


@pytest.mark.parametrize("country,word,status", [("PT", "Planta de estrutura", "PRONTO"), ("FR", "Plan de coffrage", "PRÊT")])
def test_dossier_complet(tmp_path, country, word, status):
    raw = copy.deepcopy(CASE)
    raw["project"]["country"] = country
    res = pipeline.run(raw, tmp_path, stem="T")
    assert res["status"].startswith(status), res
    assert len(res["sheets"]) >= 15 and any(word in s for s in res["sheets"])
    codes = [s.split()[0] for s in res["sheets"]]
    assert len(codes) == len(set(codes))                                             # repères de planches uniques
    pdf = PdfReader(res["files"]["pdf"])
    assert len(pdf.pages) > len(res["sheets"])
    txt = "".join(p.extract_text() or "" for p in PdfReader(res["files"]["plans_pdf"]).pages)
    if country == "PT":
        assert "Fundações" in txt or "fundações" in txt
        assert "Plan de coffrage" not in txt and "Semelles" not in txt                  # aucun reste de français dans les desenhos
    else:
        assert "Plan des fondations" in txt and "Planta de" not in txt
    ws = load_workbook(res["files"]["xlsx"]).active
    assert any(isinstance(c.value, str) and c.value.startswith("=D") for row in ws.iter_rows() for c in row)
    assert res["devis_ht"] > 100000


def _pdf(lines):
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=landscape(A3))
    for t in lines:
        c.drawString(50, 50, t)
        c.showPage()
    c.save()
    return buf.getvalue()


def test_detection_du_type():
    assert reader.detect_text([("a.pdf", _pdf(["PCMI 02 PLAN DU RDC — maison individuelle"]))]) == "maison"
    assert reader.detect_text([("b.pdf", _pdf(["PLANTA DO PISO TIPO — Fração A T3", "Fração B T2 — habitação coletiva"]))]) == "frame"
    assert reader.detect_text([("c.pdf", _pdf(["Ponte sobre a ribeira — viaduto", "tabuleiro do viaduto — ponte"]))]) == "civil"


class _Coll:
    def __init__(self):
        self.docs = []

    async def update_one(self, q, upd, upsert=False):
        for d in self.docs:
            if all(d.get(k) == v for k, v in q.items()):
                d.update(upd.get("$set", {}))
                return
        self.docs.append({**q, **upd.get("$set", {})})


class _DB:
    def __init__(self):
        self.bet_study_jobs = _Coll()


def test_hook_immeuble_de_bout_en_bout():
    db = _DB()
    steps = []

    async def step(label, **extra):
        steps.append((label, extra))
        await db.bet_study_jobs.update_one({"job_id": "j"}, {"$set": {"step": label, **extra}})
    stored = []

    async def store(db_, uid, cid, rid, kind, path, title):
        stored.append((kind, Path(path).stat().st_size))
        return {"kind": kind, "title": title}

    async def llm(system, text, images=None, model=None, max_tokens=0, json_schema=None):
        return json.dumps(CASE)
    docs = [("plantas.pdf", _pdf(["PLANTA DO PISO TIPO — Fração A, Fração B, Fração C, Fração D", "CORTE A-B", "ALÇADO SUL"]))]
    done = asyncio.run(hook.maybe_handle(db, "j", "u", "c", "r", docs, "", "PT", step=step, store=store, llm=llm))
    assert done
    final = steps[-1]
    assert final[0] == "Terminé" and final[1]["status"] == "done" and final[1]["building_kind"] == "frame"
    assert {k for k, _ in stored} == {"pdf", "plans_pdf", "xlsx", "zip"} and all(s > 1000 for _, s in stored)


def test_hook_maison_et_genie_civil():
    db = _DB()
    steps = []

    async def step(label, **extra):
        steps.append((label, extra))

    async def nope(*a, **k):
        raise AssertionError("pas d'appel IA attendu")
    maison = [("PC.pdf", _pdf(["PCMI 02 PLAN DU RDC", "PCMI 03 COUPE"]))]
    assert asyncio.run(hook.maybe_handle(db, "j", "u", "c", "r", maison, "", "FR", step=step, store=None, llm=nope)) is False
    civil = [("ouvrage.pdf", _pdf(["Viaduc — pont routier", "ouvrage d'art : pont, culées"]))]
    assert asyncio.run(hook.maybe_handle(db, "j", "u", "c", "r", civil, "", "FR", step=step, store=None, llm=nope)) is True
    assert steps[-1][1]["status"] == "failed" and "génie civil" in steps[-1][1]["error"]


def test_plans_de_lots_depuis_bet_core(tmp_path):
    from engines.bet_core import pipeline as core
    res = core.render_lots(copy.deepcopy(CASE), ["plomberie"], tmp_path / "lots.pdf", None, "pt", None)
    assert res["ok"] and (tmp_path / "lots.pdf").stat().st_size > 5000
    assert any("RED-" in s for s in res["sheets"])


def test_hooks_installes():
    from engines.bet_frame import install_hooks
    assert install_hooks.main(check_only=True) == 0
    src = (Path(install_hooks.BASE) / "assistant_tool.py").read_text(encoding="utf-8")
    assert "_frame_hook.maybe_handle" in src and "_FRAME_INSTRUCTIONS" in src
