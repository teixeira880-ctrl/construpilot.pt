"""Iteration 184 — Sprint 4:
  (A) « Devis depuis diagnostic » — POST /api/btp-solver/{problem_id}/to-quote
  (B) « Export quantités Excel » — GET /api/plans-pro/job/{job_id}/quantities.xlsx

Uses cached job pplan_4e216317b5b949 (DWG w/ Autodesk quantities) and existing
diagnostic prob_69d962d5f6. NO paid AI/Autodesk calls.
"""
import io
import os
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import pytest
import requests

BASE = os.environ.get("REACT_APP_BACKEND_URL", "").rstrip("/")
if not BASE:
    for line in open("/app/frontend/.env"):
        if line.startswith("REACT_APP_BACKEND_URL="):
            BASE = line.split("=", 1)[1].strip().rstrip("/")
            break
API = f"{BASE}/api"

PRO = ("pro@construpilot.fr", "Pro2026!")
CLIENT = ("client@construpilot.fr", "Client2026!")

KNOWN_PROBLEM = "prob_69d962d5f6"
KNOWN_DWG_JOB = "pplan_4e216317b5b949"
KNOWN_PNG_JOB = "pplan_0de14e2fddb346"


def _login(email, pwd):
    s = requests.Session()
    r = s.post(f"{API}/auth/login", json={"email": email, "password": pwd}, timeout=30)
    assert r.status_code == 200, f"login {email}: {r.status_code} {r.text[:200]}"
    assert s.cookies.get("session_token"), "no session_token"
    return s


@pytest.fixture(scope="module")
def pro_session():
    return _login(*PRO)


@pytest.fixture(scope="module")
def client_session():
    return _login(*CLIENT)


# --------------------- (A) BTP solver to-quote ---------------------

class TestBtpSolverToQuote:

    def _cleanup(self, sess, doc_id):
        try:
            sess.delete(f"{API}/documents/{doc_id}", timeout=20)
        except Exception:
            pass

    def test_unauthenticated_401(self):
        s = requests.Session()
        r = s.post(f"{API}/btp-solver/{KNOWN_PROBLEM}/to-quote", json={}, timeout=20)
        assert r.status_code == 401, f"expected 401, got {r.status_code} {r.text[:200]}"

    def test_client_403(self, client_session):
        r = client_session.post(f"{API}/btp-solver/{KNOWN_PROBLEM}/to-quote", json={}, timeout=20)
        assert r.status_code == 403, f"expected 403, got {r.status_code} {r.text[:200]}"

    def test_unknown_problem_404(self, pro_session):
        r = pro_session.post(f"{API}/btp-solver/prob_nonexistent99/to-quote", json={}, timeout=20)
        assert r.status_code == 404, f"expected 404, got {r.status_code} {r.text[:200]}"

    def test_to_quote_recommended(self, pro_session):
        # Fetch problem to determine recommended solution
        rd = pro_session.get(f"{API}/btp-solver/{KNOWN_PROBLEM}", timeout=30)
        assert rd.status_code == 200, rd.text[:200]
        prob = rd.json()
        sols = prob.get("solutions") or []
        assert sols, "no solutions in problem"
        recommended = next((s for s in sols if s.get("recommended")), sols[0])
        expected_mid = round(((recommended.get("cost_min_eur") or 0) + (recommended.get("cost_max_eur") or recommended.get("cost_min_eur") or 0)) / 2, 2)

        r = pro_session.post(f"{API}/btp-solver/{KNOWN_PROBLEM}/to-quote", json={}, timeout=30)
        assert r.status_code == 200, f"{r.status_code} {r.text[:300]}"
        j = r.json()
        assert j["doc_id"].startswith("doc_"), j
        assert j["lines"] >= 1
        assert isinstance(j["total_ht"], (int, float)) and j["total_ht"] > 0
        assert j["url"] == f"/devis-editeur/{j['doc_id']}"
        doc_id = j["doc_id"]

        try:
            # Verify GET /api/documents/<doc_id>
            g = pro_session.get(f"{API}/documents/{doc_id}", timeout=20)
            assert g.status_code == 200, g.text[:300]
            gd = g.json()
            assert gd.get("doc_type") == "quote", gd.get("doc_type")
            assert gd.get("status") == "draft", gd.get("status")
            lines = gd.get("lines") or []
            assert len(lines) == j["lines"], f"lines mismatch {len(lines)} vs {j['lines']}"
            total = 0.0
            for ln in lines:
                assert ln.get("unit") == "forfait", ln
                assert ln.get("qty") == 1, ln
                assert ln.get("price_source") == "btp_solver_estimate", ln
                total += (ln.get("qty") or 0) * (ln.get("unit_price_ht") or 0)
            assert abs(total - expected_mid) <= 1.0, f"total {total} vs midpoint {expected_mid}"
            notes = gd.get("notes") or ""
            assert KNOWN_PROBLEM in notes, f"problem_id not in notes: {notes[:200]}"
            # source/source_ref: may not be exposed in GET → skip if absent
            if "source" in gd:
                assert gd.get("source") == "btp_solver"
                assert gd.get("source_ref") == KNOWN_PROBLEM

            # Verify problem now lists doc_id
            rp = pro_session.get(f"{API}/btp-solver/{KNOWN_PROBLEM}", timeout=20)
            assert rp.status_code == 200
            qids = rp.json().get("quote_doc_ids") or []
            assert doc_id in qids, f"{doc_id} not in {qids}"
        finally:
            self._cleanup(pro_session, doc_id)

    def test_to_quote_specific_non_recommended(self, pro_session):
        rd = pro_session.get(f"{API}/btp-solver/{KNOWN_PROBLEM}", timeout=30)
        prob = rd.json()
        sols = prob.get("solutions") or []
        non_reco = next((s for s in sols if not s.get("recommended")), None)
        if not non_reco:
            pytest.skip("no non-recommended solution")
        sid = non_reco.get("id")
        label = non_reco["label"]

        r = pro_session.post(f"{API}/btp-solver/{KNOWN_PROBLEM}/to-quote", json={"solution_id": sid}, timeout=30)
        assert r.status_code == 200, r.text[:300]
        j = r.json()
        assert j["title"].startswith("Devis — "), j["title"]
        assert label[:40] in j["title"], f"title {j['title']} lacks label {label}"
        self._cleanup(pro_session, j["doc_id"])

    def test_to_quote_unknown_solution_falls_back(self, pro_session):
        r = pro_session.post(f"{API}/btp-solver/{KNOWN_PROBLEM}/to-quote", json={"solution_id": "sol_does_not_exist"}, timeout=30)
        assert r.status_code == 200, r.text[:300]
        j = r.json()
        assert j["doc_id"].startswith("doc_")
        self._cleanup(pro_session, j["doc_id"])


# --------------------- (B) plans-pro quantities.xlsx ---------------------

class TestPlansProXlsx:

    def test_client_gets_404(self, client_session):
        r = client_session.get(f"{API}/plans-pro/job/{KNOWN_DWG_JOB}/quantities.xlsx", timeout=30)
        assert r.status_code == 404, f"expected 404, got {r.status_code}"

    def test_unknown_job_404(self, pro_session):
        r = pro_session.get(f"{API}/plans-pro/job/pplan_zzz_does_not_exist/quantities.xlsx", timeout=30)
        assert r.status_code == 404, f"expected 404, got {r.status_code}"

    def test_png_job_conflict(self, pro_session):
        r = pro_session.get(f"{API}/plans-pro/job/{KNOWN_PNG_JOB}/quantities.xlsx", timeout=30)
        assert r.status_code == 409, f"expected 409, got {r.status_code} {r.text[:200]}"

    def test_dwg_xlsx_full(self, pro_session):
        r = pro_session.get(f"{API}/plans-pro/job/{KNOWN_DWG_JOB}/quantities.xlsx", timeout=60)
        assert r.status_code == 200, r.text[:200]
        ct = r.headers.get("content-type", "")
        assert "spreadsheetml" in ct, ct
        cd = r.headers.get("content-disposition", "")
        assert "attachment" in cd and "quantites-" in cd and ".xlsx" in cd, cd

        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(r.content))
        names = wb.sheetnames
        for req in ("Quantités", "Calques", "Métré"):
            assert req in names, f"missing sheet {req} in {names}"

        # Quantités row 5 headers
        ws = wb["Quantités"]
        headers = [ws.cell(row=5, column=c).value for c in range(1, 8)]
        assert headers[0] == "Ouvrage", headers
        assert headers[1] == "Nombre d'entités", headers

        # Find Murs row
        murs_row = None
        for row in ws.iter_rows(min_row=6, values_only=True):
            if row and row[0] == "Murs":
                murs_row = row
                break
        assert murs_row, "no 'Murs' row"
        assert murs_row[1] == 7607, f"Murs count {murs_row[1]}"
        assert murs_row[5] == "calques", f"Murs source {murs_row[5]}"

        # Calques sheet: >10 rows with layer names like A-WALL
        ws2 = wb["Calques"]
        rows = list(ws2.iter_rows(min_row=2, values_only=True))
        assert len(rows) > 10, f"only {len(rows)} rows in Calques"
        layer_names = [r[1] for r in rows if r and r[1]]
        assert any("A-WALL" in (n or "") or "WALL" in (n or "").upper() for n in layer_names), f"no wall-like layer in {layer_names[:20]}"

        # Métré: formulas in column F, excludes 'Annotations'
        ws3 = wb["Métré"]
        f_vals = []
        cats = []
        for row in ws3.iter_rows(min_row=2, values_only=False):
            cats.append(row[0].value)
            f_vals.append(row[5].value)
        formula_cells = [v for v in f_vals if isinstance(v, str) and v.startswith("=")]
        assert formula_cells, f"no formulas in col F: {f_vals}"
        assert "Annotations" not in cats, f"Annotations must be excluded: {cats}"
