"""Import Excel architecte (DPGF / BPU / DQE) + remplissage automatique des prix.

Espace Pro Business. Détecte les colonnes (Désignation / Unité / Quantité / PU / Total),
estime les prix unitaires manquants via l'IA, calcule les totaux, et renvoie le même
classeur Excel rempli (téléchargeable) + un aperçu des lignes.
"""
import io
import re
import uuid

import openpyxl
from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from fastapi.responses import StreamingResponse

from core.db import db
from core.security import current_user
from engines import storage
from engines.ai import price_dpgf_lines

router = APIRouter(prefix="/pro/dpgf", tags=["dpgf"])

_DESIGN = ("désignation", "designation", "libellé", "libelle", "prestation", "descriptif", "ouvrage")
_UNIT = ("unité", "unite", "un.", "u.", "unit")
_QTY = ("quantité", "quantite", "qté", "qte", "q.", "quant")
_PU = ("prix unitaire", "p.u", "pu ht", "pu", "prix u")
_TOTAL = ("montant", "total", "prix total", "p. total")


def _norm(v) -> str:
    return re.sub(r"\s+", " ", str(v or "").strip().lower())


def _find_header(ws):
    """Trouve la ligne d'en-tête + indices de colonnes (désignation/unité/qté/pu/total)."""
    for r in range(1, min(ws.max_row, 40) + 1):
        cells = {c: _norm(ws.cell(row=r, column=c).value) for c in range(1, min(ws.max_column, 30) + 1)}
        cols = {}
        for c, val in cells.items():
            if not val:
                continue
            if "design" not in cols and any(k in val for k in _DESIGN):
                cols["design"] = c
            elif "unit" not in cols and any(val == k or val.startswith(k) for k in _UNIT):
                cols["unit"] = c
            elif "qty" not in cols and any(k in val for k in _QTY):
                cols["qty"] = c
            elif "total" not in cols and any(k in val for k in _TOTAL):
                cols["total"] = c
            elif "pu" not in cols and any(k in val for k in _PU):
                cols["pu"] = c
        if "design" in cols and ("pu" in cols or "total" in cols):
            return r, cols
    return None, {}


def _num(v):
    if isinstance(v, (int, float)):
        return float(v)
    if v is None:
        return None
    s = re.sub(r"[^\d,.-]", "", str(v)).replace(",", ".")
    try:
        return float(s)
    except Exception:
        return None


@router.post("/import")
async def import_dpgf(file: UploadFile = File(...), user=Depends(current_user)) -> dict:
    raw = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(raw), data_only=False)
    except Exception:
        raise HTTPException(400, "Fichier Excel illisible (.xlsx attendu)")
    ws = wb.active
    hrow, cols = _find_header(ws)
    if not hrow:
        raise HTTPException(422, "Colonnes non reconnues. Attendu : Désignation / Unité / Quantité / Prix unitaire.")

    c_des, c_unit, c_qty = cols.get("design"), cols.get("unit"), cols.get("qty")
    c_pu, c_total = cols.get("pu"), cols.get("total")

    items, rows = [], []
    for r in range(hrow + 1, ws.max_row + 1):
        des = ws.cell(row=r, column=c_des).value if c_des else None
        if not _norm(des):
            continue
        qty = _num(ws.cell(row=r, column=c_qty).value) if c_qty else None
        unit = str(ws.cell(row=r, column=c_unit).value or "").strip() if c_unit else ""
        pu = _num(ws.cell(row=r, column=c_pu).value) if c_pu else None
        rows.append({"row": r, "designation": str(des).strip(), "unit": unit, "qty": qty, "pu": pu})
        items.append({"designation": str(des).strip(), "unit": unit, "qty": qty})

    if not rows:
        raise HTTPException(422, "Aucune ligne exploitable trouvée sous l'en-tête.")

    # Estime les PU manquants via l'IA (seulement les lignes sans prix)
    to_price = [i for i, x in enumerate(rows) if not x["pu"]]
    if to_price:
        est = await price_dpgf_lines([items[i] for i in to_price])
        for k, i in enumerate(to_price):
            if k < len(est) and est[k] > 0:
                rows[i]["pu"] = round(est[k], 2)
                rows[i]["price_source"] = "ia"

    # Remplit le classeur (PU + Total) et calcule le total général
    filled = 0
    grand_total = 0.0
    for x in rows:
        pu, qty = x.get("pu"), x.get("qty")
        if pu and c_pu:
            ws.cell(row=x["row"], column=c_pu).value = pu
            filled += 1
        if pu is not None and qty is not None:
            tot = round(pu * qty, 2)
            x["total"] = tot
            grand_total += tot
            if c_total:
                ws.cell(row=x["row"], column=c_total).value = tot

    out = io.BytesIO()
    wb.save(out)
    fid = "dpgf_" + uuid.uuid4().hex[:12]
    path = f"{storage.APP_NAME}/dpgf/{user['user_id']}/{fid}.xlsx"
    storage.put_object(path, out.getvalue(),
                       "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    await db.dpgf_imports.insert_one({
        "file_id": fid, "user_id": user["user_id"], "storage_path": path,
        "filename": file.filename, "lines": rows, "grand_total_ht": round(grand_total, 2),
    })
    return {
        "file_id": fid,
        "download_url": f"/api/pro/dpgf/file/{fid}",
        "columns": cols,
        "line_count": len(rows),
        "filled": filled,
        "grand_total_ht": round(grand_total, 2),
        "lines": rows[:500],
    }


@router.get("/file/{file_id}")
async def download_dpgf(file_id: str, user=Depends(current_user)):
    doc = await db.dpgf_imports.find_one({"file_id": file_id, "user_id": user["user_id"]},
                                         {"_id": 0, "storage_path": 1, "filename": 1})
    if not doc:
        raise HTTPException(404, "Fichier introuvable")
    data, ctype = storage.get_object(doc["storage_path"])
    name = (doc.get("filename") or "dpgf").rsplit(".", 1)[0] + "_rempli.xlsx"
    return StreamingResponse(
        io.BytesIO(data),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{name}"'})
