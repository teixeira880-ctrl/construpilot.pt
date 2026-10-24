"""Processos do Assistente: importação de um processo completo (dezenas de PDF/DWG/DXF/XLSX…) ou de um ZIP, ficheiro a ficheiro.

POST   /assistant/dossiers                      {name}          → cria o processo (dossier_id)
POST   /assistant/dossiers/{id}/files           ficheiro ou ZIP → acrescenta (ZIP descomprimido no servidor, classificado por especialidade/piso)
GET    /assistant/dossiers                                      → processos da conta
GET    /assistant/dossiers/{id}                                 → inventário
DELETE /assistant/dossiers/{id}                                 → elimina o processo e os seus ficheiros

Os ficheiros são guardados como anexos do Assistente (db.assistant_docs, em blocos de 1 MB acima de 8 MB) com dossier_id,
especialidade e piso: o Assistente recebe apenas o inventário e lê cada peça a pedido (read_document).
Um envio = um ficheiro (ou um ZIP): sem limite de número de ficheiros do lado da interface (teto do servidor: 400).
Porta de routes/assistant_dossiers.py do .fr.
"""
from __future__ import annotations

import asyncio
import base64
import hashlib
import re
import uuid
from datetime import datetime, timezone

from fastapi import APIRouter, Depends, File, HTTPException, UploadFile
from pydantic import BaseModel, Field

from core.db import db
from core.security import current_user
from engines import dossier as D
from engines.access import require_level
from engines.dossier_tool import list_dossier
from engines.office_text import OFFICE_MIMES, extract_office_text

router = APIRouter(prefix="/assistant/dossiers", tags=["assistant"], dependencies=[Depends(require_level("particulier_payant"))])
UPLOAD_MAX = 300 * 1024 * 1024          # por envio (um ficheiro ou um ZIP)
MIMES = {"pdf": "application/pdf", "dwg": "application/acad", "dxf": "image/vnd.dxf", "png": "image/png", "jpg": "image/jpeg",
         "jpeg": "image/jpeg", "webp": "image/webp", **OFFICE_MIMES}
CHUNK = 1024 * 1024


def _now():
    return datetime.now(timezone.utc).isoformat()


class DossierIn(BaseModel):
    name: str = Field(default="Processo", max_length=160)


def _pdf_info(data: bytes) -> tuple[str, int, int, list]:
    try:
        import pymupdf
        doc = pymupdf.open(stream=data, filetype="pdf")
        text = "".join(doc[i].get_text() for i in range(min(len(doc), 3)))
        reg = []
        if re.search(D._REG_TITLE, text[:3000], re.I):              # lista de desenhos: todas as páginas (30 no máximo)
            reg = D.register_from_text("\n".join(doc[i].get_text() for i in range(min(len(doc), 30))))
        return text[:20000], len(doc), len(doc.get_ocgs() or {}), reg
    except Exception:  # noqa: BLE001 — PDF ilegível: guardado na mesma, assinalado com 0 páginas
        return "", 0, 0, []


def _sheet_register(data: bytes, e: str) -> list:
    """Lista de desenhos XLS/XLSX → desenhos esperados; [] se não for uma lista ou se o livro for ilegível."""
    try:
        if e == "xls":
            import xlrd
            wb = xlrd.open_workbook(file_contents=data)
            rows = [sh.row_values(r) for sh in wb.sheets() for r in range(min(sh.nrows, 2000))]
        else:
            import io
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            rows = [list(r) for ws in wb.worksheets for r in ws.iter_rows(max_row=2000, values_only=True)]
        return D.register_from_rows(rows)[:1000]
    except Exception:  # noqa: BLE001
        return []


def _office_text(data: bytes, e: str) -> str:
    try:
        return extract_office_text(data, e)
    except Exception:  # noqa: BLE001 — folha protegida/danificada: guardada sem texto
        return ""


async def _store_file(uid: str, org, dossier_id: str, name: str, data: bytes) -> dict:
    e = D.ext_of(name)
    text, pages, layers, reg = (await asyncio.to_thread(_pdf_info, data)) if e == "pdf" else ("", 1, 0, [])
    if e in OFFICE_MIMES:
        text = await asyncio.to_thread(_office_text, data, e)
    if e in ("xls", "xlsx"):
        reg = await asyncio.to_thread(_sheet_register, data, e)
    cls = D.classify(name, text)
    if reg:
        cls["discipline"] = "lista_desenhos"
    doc = {"doc_id": f"doc_{uuid.uuid4().hex[:10]}", "user_id": uid, "organization_id": org, "dossier_id": dossier_id,
           "filename": name.rsplit("/", 1)[-1], "path": name, "mime": MIMES.get(e, "application/octet-stream"), "size": len(data),
           "pages": pages, "kind": "plan" if cls["cad"] or cls["discipline"] in D.PLAN_DISCIPLINES else ("photo" if e in D.IMAGES else "document"),
           "discipline": cls["discipline"], "level": cls["level"], "format": e, "sha256": hashlib.sha256(data).hexdigest(),
           "text": text, "has_text": len(text) >= 120, "cad_layers": layers if e == "pdf" else None, "created_at": _now(),
           **{k: cls[k] for k in ("zone", "nature", "number", "indice", "element", "phase", "plan_key") if k in cls}}
    if reg:
        doc["register"] = reg
    if len(data) > 8 * CHUNK:
        doc["chunk_count"] = (len(data) + CHUNK - 1) // CHUNK
        try:
            for i, s in enumerate(range(0, len(data), CHUNK)):
                await db.assistant_doc_chunks.insert_one({"doc_id": doc["doc_id"], "user_id": uid, "index": i,
                                                          "data_b64": base64.b64encode(data[s:s + CHUNK]).decode()})
        except Exception:
            await db.assistant_doc_chunks.delete_many({"doc_id": doc["doc_id"], "user_id": uid})
            raise
    else:
        doc["data_b64"] = base64.b64encode(data).decode()
    await db.assistant_docs.insert_one(dict(doc))
    return {**{k: doc[k] for k in ("doc_id", "filename", "path", "discipline", "level", "format", "pages", "size")},
            "rotulo": D.LABELS.get(doc["discipline"], doc["discipline"]),
            **{k: doc[k] for k in ("zone", "indice") if doc.get(k)}, **({"desenhos_na_lista": len(reg)} if reg else {})}


async def _owned(dossier_id: str, uid: str) -> dict:
    d = await db.assistant_dossiers.find_one({"dossier_id": dossier_id, "user_id": uid}, {"_id": 0})
    if not d:
        raise HTTPException(404, "Processo não encontrado")
    return d


@router.post("")
async def create_dossier(body: DossierIn, user=Depends(current_user)) -> dict:
    d = {"dossier_id": f"dos_{uuid.uuid4().hex[:10]}", "user_id": user["user_id"], "organization_id": user.get("organization_id"),
         "name": body.name.strip() or "Processo", "files_count": 0, "skipped": [], "created_at": _now(), "updated_at": _now()}
    await db.assistant_dossiers.insert_one(dict(d))
    return d


@router.post("/{dossier_id}/files")
async def add_file(dossier_id: str, file: UploadFile = File(...), user=Depends(current_user)) -> dict:
    uid = user["user_id"]
    dos = await _owned(dossier_id, uid)
    data = await file.read(UPLOAD_MAX + 1)
    if not data:
        raise HTTPException(422, "Ficheiro vazio")
    if len(data) > UPLOAD_MAX:
        raise HTTPException(413, "Envio demasiado volumoso (> 300 MB): divida o ZIP")
    try:
        files, skipped = await asyncio.to_thread(D.expand, file.filename or "ficheiro", data)
    except D.DossierError as exc:
        raise HTTPException(422, str(exc)) from None
    if dos.get("files_count", 0) + len(files) > D.MAX_FILES:
        raise HTTPException(413, f"Um processo contém no máximo {D.MAX_FILES} ficheiros")
    known = {d["sha256"] async for d in db.assistant_docs.find({"user_id": uid, "dossier_id": dossier_id}, {"_id": 0, "sha256": 1}) if d.get("sha256")}
    added = []
    for name, blob in files:
        h = hashlib.sha256(blob).hexdigest()
        if h in known:
            skipped.append(f"{name}: já está no processo")
            continue
        known.add(h)
        try:
            added.append(await _store_file(uid, user.get("organization_id"), dossier_id, name, blob))
        except Exception:  # noqa: BLE001 — um ficheiro com erro não anula os outros
            skipped.append(f"{name}: gravação impossível, tente novamente")
    await db.assistant_dossiers.update_one({"dossier_id": dossier_id, "user_id": uid},
                                           {"$inc": {"files_count": len(added)}, "$push": {"skipped": {"$each": skipped[-50:]}},
                                            "$set": {"updated_at": _now()}})
    return {"dossier_id": dossier_id, "added": added, "skipped": skipped}


@router.get("")
async def list_dossiers(user=Depends(current_user)) -> list:
    return [d async for d in db.assistant_dossiers.find({"user_id": user["user_id"]}, {"_id": 0, "skipped": 0}).sort("created_at", -1).limit(100)]


@router.get("/{dossier_id}")
async def get_dossier(dossier_id: str, user=Depends(current_user)) -> dict:
    await _owned(dossier_id, user["user_id"])
    return await list_dossier(db, user["user_id"], dossier_id)


@router.delete("/{dossier_id}")
async def delete_dossier(dossier_id: str, user=Depends(current_user)) -> dict:
    uid = user["user_id"]
    await _owned(dossier_id, uid)
    ids = [d["doc_id"] async for d in db.assistant_docs.find({"user_id": uid, "dossier_id": dossier_id}, {"_id": 0, "doc_id": 1})]
    if ids:
        await db.assistant_doc_chunks.delete_many({"user_id": uid, "doc_id": {"$in": ids}})
        await db.assistant_docs.delete_many({"user_id": uid, "dossier_id": dossier_id})
    await db.assistant_dossiers.delete_one({"dossier_id": dossier_id, "user_id": uid})
    return {"deleted": len(ids)}
