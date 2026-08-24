"""PDF exports + Quote sending + Templates + Multi-quote + Tracking +
Notifications + Signature client + Client portal + Reminders scheduler.

Refactor iter53 : extrait de `routes/plans.py`.
"""
from __future__ import annotations

import logging
import os
import re
import secrets
from datetime import datetime, timedelta, timezone
from io import BytesIO

from fastapi import Depends, HTTPException
from pydantic import BaseModel, EmailStr, Field

from core.db import db
from core.security import admin_user, current_user
from engines import geometry as G

from routes.plans_common import (
    QUOTE_SEND_PREFIX, SIG_TOKEN_PREFIX,
    _build_quote_and_pdf, _now, _resolve_branding,
    build_client_portal_token, quote_router, router,
)

logger = logging.getLogger(__name__)

REMINDER_DELAY_DAYS = 5
QUOTE_DEPOSIT_DEFAULT_PCT = 30.0


# ============ PDF EXPORT MÉTRÉ ============

def _aggregate_export_levels(rooms: list[dict], pages: list[dict]) -> list[dict]:
    """Reconstitue les niveaux à partir des pages+rooms (comme /takeoff)."""
    page_level = {pg.get("index"): {"index": pg.get("level_index"),
                                     "label": pg.get("level_label")}
                  for pg in pages}
    by_level: dict = {}
    for r in rooms:
        info = page_level.get(r.get("page_index"), {"index": None, "label": None})
        key = f"{info['index'] if info['index'] is not None else 'none'}"
        bucket = by_level.setdefault(key, {
            "level_index": info["index"],
            "level_label": info["label"] or "Sans niveau",
            "rooms": [], "totals": {"floor_m2": 0, "walls_net_m2": 0, "perimeter_m": 0}})
        bucket["rooms"].append(r)
        bucket["totals"]["floor_m2"] += r.get("surface_m2", 0)
        bucket["totals"]["walls_net_m2"] += r.get("wall_area_net_m2", 0)
        bucket["totals"]["perimeter_m"] += r.get("perimeter_m", 0)
    for b in by_level.values():
        for k in b["totals"]:
            b["totals"][k] = round(b["totals"][k], 2)
    return sorted(by_level.values(),
                  key=lambda x: (x["level_index"] is None,
                                 x["level_index"] if x["level_index"] is not None else 99))


def _takeoff_materials(rooms: list[dict], floor_material: str,
                       totals: dict) -> dict:
    """Compute the metric materials block used by the takeoff PDF."""
    doors_est = 2 * len(rooms)  # heuristique
    return {
        "floor": G.floor_quantity(totals["floor_m2"], floor_material),
        "paint": G.paint_liters(totals["walls_net_m2"] + totals["ceiling_m2"]),
        "placo": G.placo_quantity(totals["walls_net_m2"] + totals["ceiling_m2"]),
        "skirting": G.skirting_length_m(totals["perimeter_m"], doors_est * 0.9),
    }


@router.get("/{plan_id}/pdf")
async def export_takeoff_pdf(plan_id: str, include_materials: bool = True,
                              floor_material: str = "parquet",
                              user=Depends(current_user)):
    """Génère et renvoie le PDF de métré du plan (streaming binaire)."""
    from fastapi.responses import Response
    from engines.pdf_takeoff import build_takeoff_pdf

    p = await db.plans_takeoff.find_one(
        {"plan_id": plan_id, "user_id": user["user_id"]},
        {"_id": 0, "pages.image_png_b64": 0}
    )
    if not p:
        raise HTTPException(404, "Plan introuvable")

    rooms = p.get("rooms") or []
    pages = p.get("pages") or []
    totals = {
        "floor_m2": round(sum(r.get("surface_m2", 0) for r in rooms), 2),
        "ceiling_m2": round(sum(r.get("ceiling_area_m2", 0) for r in rooms), 2),
        "walls_gross_m2": round(sum(r.get("wall_area_gross_m2", 0) for r in rooms), 2),
        "walls_net_m2": round(sum(r.get("wall_area_net_m2", 0) for r in rooms), 2),
        "perimeter_m": round(sum(r.get("perimeter_m", 0) for r in rooms), 2),
    }

    takeoff = {
        "room_count": len(rooms), "rooms": rooms,
        "levels": _aggregate_export_levels(rooms, pages),
        "totals": totals,
    }
    materials = _takeoff_materials(rooms, floor_material, totals) if (include_materials and rooms) else None
    plan_light = {"plan_id": p.get("plan_id"), "filename": p.get("filename"),
                  "page_count": p.get("page_count", 1)}
    branding = await _resolve_branding(user["user_id"])

    # Vague B (iter58) — enrichir le PDF avec les modèles appliqués aux pièces
    preset_ids = list({r.get("preset_id") for r in rooms if r.get("preset_id")})
    presets_by_id: dict = {}
    if preset_ids:
        async for preset in db.tool_presets.find(
            {"preset_id": {"$in": preset_ids}, "user_id": user["user_id"]}, {"_id": 0}
        ):
            presets_by_id[preset["preset_id"]] = preset

    pdf_bytes = build_takeoff_pdf(plan_light, takeoff, materials,
                                   branding=branding, presets_by_id=presets_by_id,
                                   measurements=p.get("measurements") or [])
    return Response(
        content=pdf_bytes, media_type="application/pdf",
        headers={"Content-Disposition":
                 f'attachment; filename="metre-{plan_id}.pdf"'}
    )


# ============ PDF RAPPORT ÉCARTS (iter63) ============

async def _ecarts_pdf_for(plan_id: str, user_id: str, page_index: int):
    """Génère le rapport écarts (bytes) + méta. Lève HTTPException si prérequis manquants.

    Le matching pièces mesurées ↔ pièces IA est recalculé en direct
    (nom normalisé puis proximité de surface).
    """
    from engines import plan_texts as T
    from engines.pdf_takeoff import build_ecarts_pdf

    p = await db.plans_takeoff.find_one(
        {"plan_id": plan_id, "user_id": user_id},
        {"_id": 0, "pages.image_png_b64": 0})
    if not p:
        raise HTTPException(404, "Plan introuvable")
    ai = p.get("ai_analysis") or {}
    if not ai or ai.get("error"):
        raise HTTPException(400, "Lancez d'abord l'Analyse IA du plan")
    rooms = [r for r in (p.get("rooms") or []) if r.get("page_index") == page_index]
    if not rooms:
        raise HTTPException(400, "Aucune pièce mesurée sur cette page")

    calib = (p.get("calibrations") or {}).get(str(page_index))
    result = T.reconcile_written_surfaces(
        rooms, ai.get("rooms") or [],
        (calib or {}).get("pixels_per_meter") or 1.0)
    matches = result.get("matches", [])
    matched_ids = {m["room_id"] for m in matches}
    unmatched = [r for r in rooms if r.get("room_id") not in matched_ids]

    branding = await _resolve_branding(user_id)
    pdf_bytes = build_ecarts_pdf(
        {"plan_id": p.get("plan_id"), "filename": p.get("filename")},
        calib, matches, unmatched, branding=branding, page_index=page_index)
    ecarts = sum(1 for m in matches if abs(m.get("delta_pct") or 0) > 5.0)
    meta = {"filename": p.get("filename") or "plan",
            "matches": len(matches), "ecarts": ecarts,
            "unmatched": len(unmatched)}
    return pdf_bytes, meta


@router.get("/{plan_id}/ecarts-pdf")
async def export_ecarts_pdf(plan_id: str, page_index: int = 0,
                            user=Depends(current_user)):
    """Rapport PDF des écarts surfaces mesurées ↔ surfaces écrites sur le plan."""
    from fastapi.responses import Response
    pdf_bytes, _ = await _ecarts_pdf_for(plan_id, user["user_id"], page_index)
    return Response(
        content=pdf_bytes, media_type="application/pdf",
        headers={"Content-Disposition":
                 f'attachment; filename="rapport-ecarts-{plan_id}.pdf"'}
    )


class SendEcartsIn(BaseModel):
    page_index: int = Field(default=0, ge=0)
    to_email: EmailStr
    message: str = ""


@router.post("/{plan_id}/ecarts-pdf/send")
async def send_ecarts_pdf(plan_id: str, body: SendEcartsIn,
                          user=Depends(current_user)) -> dict:
    """Envoie le rapport d'écarts PDF au client par email avec un mot personnalisé (iter64)."""
    import html as _html
    from engines.emailer import send_email_with_pdf, wrap_html

    pdf_bytes, meta = await _ecarts_pdf_for(plan_id, user["user_id"], body.page_index)
    perso = _html.escape(body.message.strip()[:2000]).replace("\n", "<br/>")
    intro = f"<p>{perso}</p>" if perso else ""
    stats = (f"<p><b>{meta['matches']}</b> pièce(s) comparée(s) · "
             f"<b>{meta['ecarts']}</b> écart(s) &gt; 5 % · "
             f"{meta['unmatched']} sans surface écrite.</p>")
    html_body = (
        intro +
        "<p>Veuillez trouver ci-joint le <b>rapport de contrôle des surfaces</b> du plan "
        f"« {_html.escape(meta['filename'])} » : comparaison des surfaces mesurées "
        "numériquement avec les surfaces écrites sur le plan.</p>" + stats)
    html = wrap_html("Rapport de contrôle des surfaces", html_body)
    try:
        await send_email_with_pdf(
            body.to_email,
            f"Rapport de contrôle des surfaces — {meta['filename']}",
            html, pdf_bytes, filename=f"rapport-ecarts-{plan_id}.pdf")
    except RuntimeError as e:
        raise HTTPException(503, str(e))
    except Exception:
        raise HTTPException(502, "Envoi de l'email impossible — réessayez plus tard")
    # Historique des envois sur la fiche du plan (iter65)
    await db.plans_takeoff.update_one(
        {"plan_id": plan_id, "user_id": user["user_id"]},
        {"$push": {"sent_reports": {"$each": [{
            "kind": "ecarts", "to_email": body.to_email,
            "page_index": body.page_index, "sent_at": _now(),
            "ecarts": meta["ecarts"], "matches": meta["matches"],
        }], "$slice": -20}}})
    return {"ok": True, "sent_to": body.to_email, **meta}


# ============ EXPORT EXCEL (iter68) ============

@router.get("/{plan_id}/export-xlsx")
async def export_xlsx(plan_id: str, user=Depends(current_user)):
    """Export Excel des pièces et mesures — prêt pour les tableaux de chiffrage."""
    from fastapi.responses import Response
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    p = await db.plans_takeoff.find_one(
        {"plan_id": plan_id, "user_id": user["user_id"]},
        {"_id": 0, "filename": 1, "rooms": 1, "measurements": 1, "pc_study": 1})
    if not p:
        raise HTTPException(404, "Plan introuvable")

    header_font = Font(bold=True, color="FFFFFF", size=10)
    header_fill = PatternFill("solid", fgColor="1F2733")
    bold = Font(bold=True)

    def sheet_header(ws, headers: list, widths: list) -> None:
        ws.append(headers)
        for col, w in enumerate(widths, start=1):
            cell = ws.cell(row=1, column=col)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center")
            ws.column_dimensions[cell.column_letter].width = w

    wb = Workbook()
    ws = wb.active
    ws.title = "Pièces"
    sheet_header(ws, ["Pièce", "Page", "Surface nette (m²)", "Surface brute (m²)",
                      "Déductions (m²)", "Périmètre (m)", "Murs bruts (m²)",
                      "Murs nets (m²)", "Plafond (m²)", "HSP (m)", "Volume (m³)",
                      "Matériau sol", "Surface écrite (m²)", "Écart (%)"],
                 [22, 6, 16, 16, 14, 13, 13, 12, 12, 8, 12, 16, 16, 10])
    rooms = p.get("rooms") or []
    for r in rooms:
        ws.append([r.get("name"), (r.get("page_index") or 0) + 1,
                   r.get("surface_m2"),
                   r.get("surface_gross_m2") or r.get("surface_m2"),
                   r.get("deductions_area_m2") or 0,
                   r.get("perimeter_m"), r.get("wall_area_gross_m2"),
                   r.get("wall_area_net_m2"), r.get("ceiling_area_m2"),
                   r.get("ceiling_height_m"), r.get("volume_m3"),
                   r.get("floor_material") or "",
                   r.get("surface_written_m2") or "",
                   r.get("surface_delta_pct") if r.get("surface_delta_pct") is not None else ""])
    if rooms:
        last = ws.max_row
        totals = ["TOTAL", ""]
        for letter in ["C", "D", "E", "F", "G", "H", "I"]:
            totals.append(f"=SUM({letter}2:{letter}{last})")
        totals.append("")
        totals.append(f"=SUM(K2:K{last})")
        ws.append(totals)
        for cell in ws[ws.max_row]:
            cell.font = bold

    ws2 = wb.create_sheet("Mesures")
    sheet_header(ws2, ["Libellé", "Type", "Page", "Longueur (m)", "Surface (m²)",
                       "Périmètre (m)", "Créée le"],
                 [28, 12, 6, 13, 13, 13, 18])
    kind_fr = {"distance": "Distance", "polyline": "Polyligne", "polygon": "Surface"}
    for m in p.get("measurements") or []:
        is_poly = m.get("kind") == "polygon"
        ws2.append([m.get("label") or "",
                    kind_fr.get(m.get("kind"), m.get("kind") or ""),
                    (m.get("page_index") or 0) + 1,
                    "" if is_poly else (m.get("value_m") or m.get("length_m") or ""),
                    m.get("area_m2") if is_poly else "",
                    m.get("perimeter_m") if is_poly else "",
                    str(m.get("created_at") or "")[:16].replace("T", " ")])

    # Étude PC (gros œuvre) — feuille dédiée si l'analyse a été lancée
    study = p.get("pc_study")
    if study:
        ws3 = wb.create_sheet("Étude PC")
        sheet_header(ws3, ["Lot", "Désignation", "Détail", "Unité", "Qté",
                           "PU mat. (€)", "PU MO (€)", "Total HT (€)"],
                     [30, 42, 52, 8, 10, 12, 12, 14])
        for lot in study.get("lots") or []:
            for ln in lot.get("lines") or []:
                ws3.append([lot.get("lot"), ln.get("designation"), ln.get("detail"),
                            ln.get("unit"), ln.get("qty"), ln.get("pu_mat"),
                            ln.get("pu_mo"), ln.get("total")])
        t = study.get("totals") or {}
        ws3.append([])
        for label, val in [("Béton total (m³)", t.get("beton_m3")),
                           ("Aciers total (kg)", t.get("acier_kg")),
                           ("Briques isobrique (u)", t.get("briques_u")),
                           ("Total HT (€)", t.get("ht")),
                           ("TVA (€)", t.get("tva")),
                           ("TOTAL TTC (€)", t.get("ttc"))]:
            ws3.append(["", "", "", "", "", "", label, val])
            ws3.cell(row=ws3.max_row, column=7).font = bold
            ws3.cell(row=ws3.max_row, column=8).font = bold

    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition":
                 f'attachment; filename="metre-{plan_id}.xlsx"'}
    )


# ============ PDF DEVIS ============

@router.get("/{plan_id}/quote-pdf")
async def export_quote_pdf(plan_id: str, finish: str = "std",
                            floor_material: str = "carrelage",
                            include_paint: bool = True, include_placo: bool = True,
                            include_floor: bool = True, include_skirting: bool = True,
                            user=Depends(current_user)):
    """Génère un PDF Devis complet : main d'œuvre + matériaux pré-remplis depuis le takeoff."""
    from fastapi.responses import Response
    _q, pdf_bytes = await _build_quote_and_pdf(
        plan_id, user["user_id"], finish, floor_material,
        include_paint, include_placo, include_floor, include_skirting)
    return Response(
        content=pdf_bytes, media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="devis-{plan_id}.pdf"'}
    )


# ============ SEND QUOTE (email + tracking + optional TVA + optional video) ============

class SendQuoteIn(BaseModel):
    client_email: str
    client_name: str | None = None
    client_address: str | None = None
    work_address: str | None = None
    message: str | None = None
    subject: str | None = None
    intro: str | None = None
    finish: str = "std"
    floor_material: str = "carrelage"
    include_paint: bool = True
    include_placo: bool = True
    include_floor: bool = True
    include_skirting: bool = True
    cc_self: bool = True
    include_tva_attestation: bool = False
    attach_latest_video: bool = False


def _build_send_quote_html(*, hello: str, intro_html: str, custom_html: str,
                          ttc: float, portal_url: str, company: str, contact: str,
                          pixel_tag: str) -> str:
    """Assemble the HTML body of the devis email."""
    return (
        f"<p>{hello}</p>"
        f"{intro_html}"
        f"{custom_html}"
        f"<div style='background:#FF5A00;color:white;padding:16px 22px;border-radius:6px;"
        f"font-size:18px;font-weight:700;text-align:center;letter-spacing:1px'>"
        f"Total TTC : {ttc:,.2f} €</div>"
        f"<p style='margin-top:16px'>Ce document est un devis estimatif indicatif "
        f"(TVA 10 % rénovation logement &gt; 2 ans). Nous restons à votre disposition "
        f"pour toute question ou visite technique.</p>"
        f"<div style='margin-top:24px;padding:16px;background:#f8fafc;border-radius:6px;"
        f"border-left:3px solid #FF5A00'>"
        f"<p style='margin:0 0 6px;font-weight:700;font-size:13px'>📁 Votre espace devis</p>"
        f"<p style='margin:0 0 10px;font-size:12px;color:#64748B'>"
        f"Retrouvez tous vos devis reçus (celui-ci et les précédents) "
        f"dans votre espace personnel sécurisé :</p>"
        f"<a href='{portal_url}' style='color:#FF5A00;font-weight:700;text-decoration:none;font-size:13px'>"
        f"→ Ouvrir mon espace devis</a>"
        f"</div>"
        f"<p style='color:#666;font-size:12px;margin-top:24px'>"
        f"{company}<br/>{contact}</p>"
        f"{pixel_tag}"
    )


def _build_tva_attachment(quote: dict, body: SendQuoteIn, company: str,
                          contact: str, branding: dict) -> dict:
    """Build the TVA attestation PDF attachment payload."""
    from engines.pdf_takeoff import build_tva_attestation_pdf
    work_desc_parts = [li["name"] for li in quote.get("line_items", [])]
    work_desc = " ; ".join(work_desc_parts) if work_desc_parts else "Travaux de rénovation"
    client_addr = (body.client_address or "").strip()
    work_addr = (body.work_address or client_addr or "").strip()
    tva_pdf = build_tva_attestation_pdf(
        client_name=body.client_name or "",
        client_address=client_addr,
        work_address=work_addr,
        work_description=work_desc,
        company_name=company,
        company_address=contact or "",
        branding=branding, tva_rate=quote.get("tva_rate", 0.10),
    )
    return {
        "filename": f"attestation-tva-{quote['plan_id']}.pdf",
        "content": tva_pdf,
        "content_type": "application/pdf",
    }


async def _maybe_attach_latest_video(plan_id: str, user_id: str) -> tuple[dict | None, str | None]:
    """Return (attachment_dict, video_id) if a recent MP4 fits under 20 MB — else (None, None)."""
    latest = await db.plan_videos_3d.find_one(
        {"plan_id": plan_id, "user_id": user_id},
        sort=[("created_at", -1)],
    )
    if not latest:
        return None, None
    from engines.video_tour import get_video_path
    vpath = get_video_path(latest.get("video_id", ""))
    if not vpath:
        return None, None
    try:
        mp4_bytes = vpath.read_bytes()
    except Exception as e:
        logger.warning(f"failed to read video: {e}")
        return None, None
    if len(mp4_bytes) > 20 * 1024 * 1024:
        logger.warning(f"video too large to attach: {len(mp4_bytes)} bytes")
        return None, None
    return ({
        "filename": f"tour-3d-{plan_id}.mp4",
        "content": mp4_bytes,
        "content_type": "video/mp4",
    }, latest["video_id"])


@router.post("/{plan_id}/send-quote")
async def send_quote(plan_id: str, body: SendQuoteIn, user=Depends(current_user)) -> dict:
    """Envoie le PDF Devis directement au client par email + copie au sender + log."""
    if not re.match(r"^[^@\s]+@[^@\s]+\.[^@\s]+$", body.client_email or ""):
        raise HTTPException(400, "Email client invalide.")

    quote, pdf_bytes = await _build_quote_and_pdf(
        plan_id, user["user_id"], body.finish, body.floor_material,
        body.include_paint, body.include_placo, body.include_floor, body.include_skirting)

    if body.include_tva_attestation and not (body.client_address or "").strip():
        raise HTTPException(400,
            "Adresse client requise pour joindre l'attestation TVA.")

    branding = await _resolve_branding(user["user_id"]) or {}
    sender_u = await db.users.find_one({"user_id": user["user_id"]},
                                         {"_id": 0, "name": 1, "email": 1})
    company = branding.get("company_name") or (sender_u or {}).get("name") or "ConstruPilot"
    company_email = branding.get("company_email")
    contact = branding.get("contact_line") or ""

    ttc = quote.get("total_ttc", 0)
    finish_label = {"eco": "Éco", "std": "Standard", "premium": "Premium"}.get(
        body.finish, "Standard")

    from engines.emailer import wrap_company_html
    from routes.plans_common import deliver_email
    hello = f"Bonjour {body.client_name}," if body.client_name else "Bonjour,"
    intro_html = (f"<p>{body.intro}</p>" if body.intro
                   else "<p>Veuillez trouver ci-joint votre <b>devis estimatif</b>"
                        f" ({finish_label}) élaboré depuis votre plan.</p>")
    custom_html = f"<p>{body.message}</p>" if body.message else ""
    log_id = QUOTE_SEND_PREFIX + secrets.token_hex(6)

    backend_public = os.getenv("PUBLIC_BACKEND_URL") or os.getenv("FRONTEND_URL") or ""
    pixel_tag = (f'<img src="{backend_public}/api/quotes/track/open/{log_id}" '
                 f'width="1" height="1" alt="" style="display:block" />'
                 if backend_public else "")
    portal_token = build_client_portal_token(body.client_email or "")
    portal_url = (f"{os.getenv('FRONTEND_URL', 'https://construpilot.fr').rstrip('/')}"
                  f"/mes-devis?token={portal_token}")

    inner = _build_send_quote_html(
        hello=hello, intro_html=intro_html, custom_html=custom_html,
        ttc=ttc, portal_url=portal_url, company=company, contact=contact,
        pixel_tag=pixel_tag,
    )
    subject = (body.subject or "").strip()[:180] or f"Devis · {company}"
    html = wrap_company_html(subject, inner, company_name=company,
                             logo_b64=branding.get("logo_b64"), contact_line=contact,
                             siret=branding.get("siret"), vat=branding.get("vat"),
                             company_email=company_email)

    extra_attachments: list[dict] = []
    if body.include_tva_attestation:
        extra_attachments.append(_build_tva_attachment(quote, body, company, contact, branding))
        # Ajoute la mention dans le message
        inner_with_att = inner.replace(
            "Ce document est un devis estimatif indicatif",
            "Vous trouverez également jointe l'<b>attestation TVA à taux réduit</b> "
            "à nous retourner signée avant la facturation. Ce document est un "
            "devis estimatif indicatif"
        )
        html = wrap_company_html(subject, inner_with_att, company_name=company,
                                 logo_b64=branding.get("logo_b64"), contact_line=contact,
                                 siret=branding.get("siret"), vat=branding.get("vat"),
                                 company_email=company_email)

    attached_video_id: str | None = None
    if body.attach_latest_video:
        video_att, attached_video_id = await _maybe_attach_latest_video(
            plan_id, user["user_id"])
        if video_att:
            extra_attachments.append(video_att)
            html = html.replace(
                "Total TTC :",
                "<p style='margin:0 0 12px;font-size:13px;color:#FF5A00;"
                "font-weight:700'>🎬 Un tour vidéo 3D de votre projet est "
                "joint à cet email — regardez-le pour visualiser votre "
                "chantier avant travaux.</p>Total TTC :",
                1,
            )

    email_status = "sent"
    email_id = None
    error = None
    try:
        cc = (sender_u or {}).get("email") if body.cc_self else None
        res = await deliver_email(user["user_id"], body.client_email, subject,
                                  html, pdf_bytes=pdf_bytes,
                                  filename=f"devis-{plan_id}.pdf", cc=cc,
                                  extra_attachments=extra_attachments,
                                  reply_to=(company_email or cc),
                                  from_name=company)
        email_id = res.get("id")
    except Exception as e:
        logger.warning(f"Send quote email failed: {e}")
        email_status = "failed"
        error = str(e)

    log = {
        "log_id": log_id,
        "plan_id": plan_id, "user_id": user["user_id"],
        "client_email": body.client_email, "client_name": body.client_name,
        "client_address": body.client_address, "work_address": body.work_address,
        "cc": (sender_u or {}).get("email") if body.cc_self else None,
        "total_ttc": ttc, "finish": body.finish, "floor_material": body.floor_material,
        "subject": subject, "intro": body.intro, "message": body.message,
        "include_tva_attestation": body.include_tva_attestation,
        "attach_latest_video": body.attach_latest_video,
        "attached_video_id": attached_video_id,
        "email_status": email_status, "email_id": email_id, "error": error,
        "created_at": _now(),
        "opened": False, "open_count": 0,
        "signed": False,
        "reminder_sent": False,
    }
    await db.quote_sends.insert_one(log)
    return {"ok": email_status == "sent", "email_status": email_status,
            "email_id": email_id, "log_id": log["log_id"], "total_ttc": ttc}


@router.get("/{plan_id}/quote-sends")
async def list_quote_sends(plan_id: str, user=Depends(current_user)) -> dict:
    """Historique des envois de devis pour ce plan."""
    docs = await db.quote_sends.find(
        {"plan_id": plan_id, "user_id": user["user_id"]},
        {"_id": 0}
    ).sort("created_at", -1).limit(50).to_list(50)
    return {"plan_id": plan_id, "sends": docs}


# ============ WORKFLOW : Plan → Estimation → Devis ============

class TakeoffEstimateIn(BaseModel):
    project_type: str = "renovation_complete"
    finish_level: str = "standard"
    country: str = "FR"
    region: str | None = "idf"
    building_type: str = "maison"
    condition: str = "bon"
    construction_year: int | None = None


@router.post("/{plan_id}/estimate-from-takeoff")
async def estimate_from_takeoff(plan_id: str, body: TakeoffEstimateIn,
                                 user=Depends(current_user)) -> dict:
    """Envoie automatiquement la surface mesurée du plan à l'estimateur."""
    from engines.estimation import estimate as estim_run

    p = await db.plans_takeoff.find_one(
        {"plan_id": plan_id, "user_id": user["user_id"]},
        {"_id": 0, "rooms": 1, "filename": 1}
    )
    if not p:
        raise HTTPException(404, "Plan introuvable")
    rooms = p.get("rooms") or []
    if not rooms:
        raise HTTPException(400, "Aucune pièce mesurée sur ce plan.")

    surface = round(sum(r.get("surface_m2", 0) for r in rooms), 2)
    if surface <= 0:
        raise HTTPException(400, "Surface mesurée invalide.")

    est = estim_run(
        project_type=body.project_type, surface=surface,
        finish=body.finish_level, country=body.country, region=body.region,
        building_type=body.building_type, condition=body.condition,
        construction_year=body.construction_year,
    )

    proj = {
        "project_id": "proj_" + secrets.token_hex(6),
        "user_id": user["user_id"], "plan_id": plan_id,
        "name": p.get("filename") or "Projet depuis plan",
        "source": "plan_takeoff",
        "surface": surface, "room_count": len(rooms),
        "estimation": est,
        "created_at": _now(),
    }
    await db.projects.insert_one(proj)

    return {"ok": True, "project_id": proj["project_id"],
            "plan_id": plan_id, "surface_m2": surface,
            "room_count": len(rooms),
            "estimation": est,
            "is_renovation": "renovation" in body.project_type.lower()}


# ============ REMINDERS ============

def _build_reminder_html(*, hello: str, ttc: float, sign_url: str,
                        portal_url: str, company: str) -> str:
    return (
        f"<p>{hello}</p>"
        f"<p>Nous nous permettons un petit rappel : votre devis "
        f"(<b>{ttc:,.2f} € TTC</b>) est toujours disponible "
        f"pour votre signature.</p>"
        f"<p>Il vous suffit d'un clic pour valider :</p>"
        f"<p style='text-align:center;margin:24px 0'>"
        f"<a href='{sign_url}' style='background:#FF5A00;color:white;"
        f"padding:14px 28px;border-radius:6px;text-decoration:none;"
        f"font-weight:700;font-size:16px'>Signer le devis</a></p>"
        f"<p style='text-align:center;margin:8px 0 24px'>"
        f"<a href='{portal_url}' style='color:#64748B;font-size:12px;"
        f"text-decoration:underline'>Voir tous mes devis dans mon espace</a></p>"
        f"<p>Nous restons à votre disposition pour toute question ou ajustement.</p>"
        f"<p style='color:#666;font-size:12px;margin-top:24px'>"
        f"Cordialement,<br/>{company}</p>"
    )


async def _ensure_signature_token(log: dict) -> str:
    """Return an existing signature token or create+persist a new one."""
    token = log.get("signature_token")
    if token:
        return token
    token = SIG_TOKEN_PREFIX + secrets.token_urlsafe(18)
    await db.quote_sends.update_one({"log_id": log["log_id"]},
        {"$set": {"signature_token": token,
                  "signature_token_created_at": _now()}})
    return token


async def _send_single_reminder(c: dict) -> None:
    from engines.emailer import wrap_company_html
    from routes.plans_common import deliver_email
    branding = await _resolve_branding(c["user_id"]) or {}
    company = branding.get("company_name") or "ConstruPilot"
    company_email = branding.get("company_email")
    hello = f"Bonjour {c['client_name']}," if c.get("client_name") else "Bonjour,"
    token = await _ensure_signature_token(c)
    sign_url = f"{os.getenv('FRONTEND_URL', '')}/sign/{token}"
    portal_token = build_client_portal_token(c.get("client_email") or "")
    portal_url = (f"{os.getenv('FRONTEND_URL', 'https://construpilot.fr').rstrip('/')}"
                  f"/mes-devis?token={portal_token}")
    inner = _build_reminder_html(hello=hello, ttc=c.get("total_ttc", 0),
                                  sign_url=sign_url, portal_url=portal_url,
                                  company=company)
    await deliver_email(c["user_id"], c["client_email"], f"Rappel : votre devis · {company}",
                        wrap_company_html(f"Rappel devis · {company}", inner,
                                          company_name=company, logo_b64=branding.get("logo_b64"),
                                          contact_line=branding.get("contact_line"),
                                          siret=branding.get("siret"), vat=branding.get("vat"),
                                          company_email=company_email),
                        reply_to=company_email, from_name=company)
    await db.quote_sends.update_one({"log_id": c["log_id"]},
        {"$set": {"reminder_sent": True, "reminder_sent_at": _now()}})
    await db.notifications.insert_one({
        "notif_id": "n_" + secrets.token_hex(5),
        "user_id": c["user_id"], "type": "reminder_sent",
        "log_id": c["log_id"],
        "message": f"Rappel envoyé à {c['client_email']} (devis non signé "
                   f"depuis {REMINDER_DELAY_DAYS} j).",
        "created_at": _now(), "read": False,
    })


async def _run_reminders() -> dict:
    """Envoie un rappel aux devis ouverts mais non signés depuis > 5 jours."""
    cutoff = (datetime.now(timezone.utc) - timedelta(days=REMINDER_DELAY_DAYS)).isoformat()
    candidates = await db.quote_sends.find(
        {"opened": True, "signed": False, "reminder_sent": {"$ne": True},
         "opened_at": {"$lte": cutoff}, "email_status": "sent"},
        {"_id": 0, "log_id": 1, "user_id": 1, "client_email": 1, "client_name": 1,
         "total_ttc": 1, "signature_token": 1, "plan_id": 1, "subject": 1}
    ).to_list(200)

    sent, failed = 0, 0
    for c in candidates:
        try:
            await _send_single_reminder(c)
            sent += 1
        except Exception as e:
            logger.warning(f"reminder for {c.get('log_id')} failed: {e}")
            failed += 1
    return {"sent": sent, "failed": failed, "skipped": 0,
            "candidates": len(candidates)}


@quote_router.post("/reminders/run")
async def trigger_reminders(user=Depends(admin_user)):
    """Endpoint admin pour lancer manuellement le job de rappels."""
    return await _run_reminders()


# ============ QUOTE TEMPLATES ============

class QuoteTemplateIn(BaseModel):
    name: str = Field(min_length=2, max_length=80)
    finish: str = "std"
    floor_material: str = "carrelage"
    include_paint: bool = True
    include_placo: bool = True
    include_floor: bool = True
    include_skirting: bool = True
    is_shared: bool = False


@quote_router.post("/templates")
async def create_quote_template(body: QuoteTemplateIn, user=Depends(current_user)) -> dict:
    tpl_id = "qtpl_" + secrets.token_hex(5)
    u = await db.users.find_one({"user_id": user["user_id"]},
                                 {"_id": 0, "organization_id": 1}) or {}
    org_id = u.get("organization_id") if body.is_shared else None
    doc = {
        "template_id": tpl_id, "user_id": user["user_id"],
        "organization_id": org_id,
        "name": body.name.strip()[:80],
        "finish": body.finish, "floor_material": body.floor_material,
        "include_paint": body.include_paint, "include_placo": body.include_placo,
        "include_floor": body.include_floor, "include_skirting": body.include_skirting,
        "created_at": _now(),
    }
    await db.quote_templates.insert_one(doc)
    return {"ok": True, "template_id": tpl_id, **{k: v for k, v in doc.items() if k != "_id"}}


@quote_router.get("/templates")
async def list_quote_templates(user=Depends(current_user)) -> dict:
    u = await db.users.find_one({"user_id": user["user_id"]},
                                 {"_id": 0, "organization_id": 1}) or {}
    org_id = u.get("organization_id")
    query = {"$or": [{"user_id": user["user_id"]}]}
    if org_id:
        query["$or"].append({"organization_id": org_id})
    docs = await db.quote_templates.find(query, {"_id": 0}).sort("created_at", -1).to_list(100)
    return {"templates": docs}


@quote_router.delete("/templates/{template_id}")
async def delete_quote_template(template_id: str, user=Depends(current_user)) -> dict:
    tpl = await db.quote_templates.find_one({"template_id": template_id}, {"_id": 0})
    if not tpl:
        raise HTTPException(404, "Modèle introuvable")
    if tpl.get("user_id") != user["user_id"]:
        raise HTTPException(403, "Réservé au créateur du modèle")
    await db.quote_templates.delete_one({"template_id": template_id})
    return {"ok": True}


# ============ MULTI-PLAN QUOTE PDF ============

class MultiQuoteIn(BaseModel):
    plan_ids: list[str] = Field(min_length=1, max_length=20)
    project_name: str | None = None
    finish: str = "std"
    floor_material: str = "carrelage"
    include_paint: bool = True
    include_placo: bool = True
    include_floor: bool = True
    include_skirting: bool = True


@quote_router.post("/multi-pdf")
async def export_multi_quote_pdf(body: MultiQuoteIn, user=Depends(current_user)):
    """Génère un PDF Devis projet consolidé pour plusieurs plans."""
    from fastapi.responses import Response
    from engines.pdf_takeoff import build_project_quote_pdf

    sub_quotes = []
    for pid in body.plan_ids:
        try:
            q, _pdf = await _build_quote_and_pdf(
                pid, user["user_id"], body.finish, body.floor_material,
                body.include_paint, body.include_placo,
                body.include_floor, body.include_skirting,
            )
            sub_quotes.append(q)
        except HTTPException as e:
            if e.status_code == 400:
                continue
            raise
    if not sub_quotes:
        raise HTTPException(400, "Aucun plan n'a de pièce mesurée.")

    def _sum(key: str) -> float:
        return round(sum(q.get(key, 0) for q in sub_quotes), 2)

    project = {
        "project_name": (body.project_name or "Projet ConstruPilot").strip(),
        "plan_count": len(sub_quotes),
        "sub_quotes": sub_quotes,
        "totals": {
            "floor_m2": round(sum(q["totals"]["floor_m2"] for q in sub_quotes), 2),
            "ceiling_m2": round(sum(q["totals"]["ceiling_m2"] for q in sub_quotes), 2),
            "walls_net_m2": round(sum(q["totals"]["walls_net_m2"] for q in sub_quotes), 2),
            "perimeter_m": round(sum(q["totals"]["perimeter_m"] for q in sub_quotes), 2),
        },
        "subtotal_material": _sum("subtotal_material"),
        "subtotal_labour": _sum("subtotal_labour"),
        "subtotal_ht": _sum("subtotal_ht"),
        "tva_rate": 0.10,
        "tva": _sum("tva"),
        "total_ttc": _sum("total_ttc"),
        "finish": body.finish, "floor_material": body.floor_material,
    }
    branding = await _resolve_branding(user["user_id"])
    pdf_bytes = build_project_quote_pdf(project, branding=branding)
    return Response(
        content=pdf_bytes, media_type="application/pdf",
        headers={"Content-Disposition":
                 f'attachment; filename="devis-projet-{secrets.token_hex(3)}.pdf"'}
    )


# ============ TRACKING & NOTIFICATIONS ============

@quote_router.get("/track/open/{log_id}", include_in_schema=False)
async def track_quote_open(log_id: str):
    """Pixel de tracking : marque le devis comme ouvert (idempotent)."""
    from fastapi.responses import Response
    pixel = b"GIF89a\x01\x00\x01\x00\x80\x00\x00\x00\x00\x00\xff\xff\xff!\xf9\x04" \
            b"\x01\x00\x00\x00\x00,\x00\x00\x00\x00\x01\x00\x01\x00\x00\x02\x02" \
            b"D\x01\x00;"
    try:
        log = await db.quote_sends.find_one({"log_id": log_id},
                                             {"_id": 0, "opened": 1,
                                              "user_id": 1, "client_email": 1})
        if log and not log.get("opened"):
            await db.quote_sends.update_one(
                {"log_id": log_id},
                {"$set": {"opened": True, "opened_at": _now()},
                 "$inc": {"open_count": 1}}
            )
            await db.notifications.insert_one({
                "notif_id": "n_" + secrets.token_hex(5),
                "user_id": log["user_id"],
                "type": "quote_opened", "log_id": log_id,
                "message": f"Le client ({log.get('client_email')}) a ouvert votre devis.",
                "created_at": _now(), "read": False,
            })
        elif log:
            await db.quote_sends.update_one({"log_id": log_id},
                                             {"$inc": {"open_count": 1},
                                              "$set": {"last_opened_at": _now()}})
    except Exception as e:
        logger.warning(f"tracking pixel fail {log_id}: {e}")
    return Response(content=pixel, media_type="image/gif",
                    headers={"Cache-Control": "no-store, no-cache, must-revalidate"})


@quote_router.get("/notifications/me")
async def list_notifications(user=Depends(current_user)) -> dict:
    docs = await db.notifications.find(
        {"user_id": user["user_id"]}, {"_id": 0}
    ).sort("created_at", -1).limit(50).to_list(50)
    unread = sum(1 for d in docs if not d.get("read"))
    return {"unread": unread, "notifications": docs}


@quote_router.post("/notifications/{notif_id}/read")
async def mark_notification_read(notif_id: str, user=Depends(current_user)) -> dict:
    await db.notifications.update_one(
        {"notif_id": notif_id, "user_id": user["user_id"]},
        {"$set": {"read": True, "read_at": _now()}}
    )
    return {"ok": True}


@quote_router.post("/notifications/read-all")
async def mark_all_notifications_read(user=Depends(current_user)) -> dict:
    """Marque toutes les notifs comme lues."""
    res = await db.notifications.update_many(
        {"user_id": user["user_id"], "read": False},
        {"$set": {"read": True, "read_at": _now()}}
    )
    return {"ok": True, "marked": res.modified_count}


@quote_router.get("/sends/all")
async def list_all_quote_sends(user=Depends(current_user)) -> dict:
    """Suivi Devis : tous les devis envoyés par l'utilisateur, plan_ids agrégés."""
    docs = await db.quote_sends.find(
        {"user_id": user["user_id"]},
        {"_id": 0, "signature_b64": 0}
    ).sort("created_at", -1).limit(200).to_list(200)

    plan_ids = list({d.get("plan_id") for d in docs if d.get("plan_id")})
    plans = await db.plans_takeoff.find(
        {"plan_id": {"$in": plan_ids}, "user_id": user["user_id"]},
        {"_id": 0, "plan_id": 1, "filename": 1}
    ).to_list(500) if plan_ids else []
    filenames = {p["plan_id"]: p.get("filename") for p in plans}
    for d in docs:
        d["plan_filename"] = filenames.get(d.get("plan_id"))

    stats = {"sent": len(docs),
             "opened": sum(1 for d in docs if d.get("opened")),
             "signed": sum(1 for d in docs if d.get("signed")),
             "not_opened": sum(1 for d in docs
                                if not d.get("opened") and not d.get("signed"))}
    stats["total_ttc_signed"] = round(sum(d.get("total_ttc", 0) for d in docs
                                           if d.get("signed")), 2)
    stats["total_ttc_pipeline"] = round(sum(d.get("total_ttc", 0) for d in docs
                                              if not d.get("signed")), 2)
    return {"sends": docs, "stats": stats}


class SendToInvoiceIn(BaseModel):
    vat_rate: float = 10.0     # 5.5 | 10 | 20


@quote_router.post("/sends/{log_id}/to-invoice")
async def send_to_invoice(log_id: str, body: SendToInvoiceIn | None = None,
                          user=Depends(current_user)) -> dict:
    """Convertit un devis suivi (envoyé au client) en facture dans l'espace Facturation."""
    from routes.documents import _compute_totals, _next_number

    send = await db.quote_sends.find_one(
        {"log_id": log_id, "user_id": user["user_id"]}, {"_id": 0})
    if not send:
        raise HTTPException(404, "Devis introuvable")
    if send.get("invoice_doc_id"):
        return {"doc_id": send["invoice_doc_id"], "already": True}

    vat_rate = float(body.vat_rate) if body else 10.0
    if vat_rate not in (5.5, 10.0, 20.0):
        vat_rate = 10.0
    ttc = float(send.get("total_ttc") or 0)
    ht = round(ttc / (1 + vat_rate / 100), 2) if ttc else 0.0
    label = send.get("plan_filename") or send.get("work_address") \
        or send.get("client_name") or "Travaux de rénovation"
    lines = [{"designation": f"Travaux — {label}", "unit": "forfait", "qty": 1,
              "unit_price_ht": ht, "vat_rate": vat_rate, "category": ""}]
    totals = _compute_totals(lines, 0)

    name = (send.get("client_name") or "").strip()
    parts = name.split(" ", 1)
    client = {
        "name": name, "company": "", "email": send.get("client_email") or "",
        "phone": "", "address": send.get("client_address") or "",
        "postal_code": "", "city": "", "siret": "",
        "first_name": parts[0] if parts else "",
        "last_name": parts[1] if len(parts) > 1 else "",
    }
    now = _now()
    number = await _next_number(user["user_id"], "invoice")
    doc = {
        "doc_id": "doc_" + secrets.token_hex(6),
        "user_id": user["user_id"], "doc_type": "invoice", "number": number,
        "title": f"Facture — {label}", "client_id": None, "client": client,
        "status": "sent", "lines": lines, "notes": "",
        "payment_terms": "", "deposit_pct": 0, "totals": totals,
        "payments": [], "paid": 0.0, "remaining_due": totals["ttc"],
        "issue_date": now,
        "due_date": (datetime.now(timezone.utc) + timedelta(days=30)).isoformat(),
        "source": "from_quote_send", "parent_log_id": log_id,
        "created_at": now, "updated_at": now,
    }
    await db.pro_documents.insert_one(dict(doc))
    await db.quote_sends.update_one(
        {"log_id": log_id}, {"$set": {"invoice_doc_id": doc["doc_id"],
                                       "invoiced_at": now}})
    doc.pop("_id", None)
    return {"doc_id": doc["doc_id"], "number": number}



@quote_router.get("/revenue/monthly")
async def revenue_monthly(months: int = 12, user=Depends(current_user)) -> dict:
    """CA mensuel des devis SIGNÉS (N derniers mois) — pour le graphique dashboard."""
    from collections import OrderedDict
    docs = await db.quote_sends.find(
        {"user_id": user["user_id"], "signed": True},
        {"_id": 0, "total_ttc": 1, "signed_at": 1, "created_at": 1},
    ).to_list(5000)
    labels = ["janv.", "févr.", "mars", "avr.", "mai", "juin",
              "juil.", "août", "sept.", "oct.", "nov.", "déc."]
    now = datetime.now(timezone.utc)
    months = max(1, min(months, 24))
    buckets: "OrderedDict[str, dict]" = OrderedDict()
    for i in range(months - 1, -1, -1):
        mm = now.month - i
        yy = now.year
        while mm <= 0:
            mm += 12
            yy -= 1
        key = f"{yy:04d}-{mm:02d}"
        buckets[key] = {"month": key, "label": f"{labels[mm - 1]} {str(yy)[2:]}",
                        "revenue": 0.0, "count": 0}
    for d in docs:
        dt = (d.get("signed_at") or d.get("created_at") or "")[:7]
        if dt in buckets:
            buckets[dt]["revenue"] += d.get("total_ttc") or 0
            buckets[dt]["count"] += 1
    data = list(buckets.values())
    for b in data:
        b["revenue"] = round(b["revenue"], 2)
    return {"months": data,
            "total": round(sum(b["revenue"] for b in data), 2),
            "signed_count": sum(b["count"] for b in data)}


@quote_router.get("/recap-pdf")
async def recap_pdf(year: int, month: int, user=Depends(current_user)):
    """Rapport comptable PDF mensuel prêt à envoyer au comptable."""
    from fastapi.responses import Response
    from engines.pdf_takeoff import build_accounting_recap_pdf
    if not (1 <= month <= 12):
        raise HTTPException(400, "Mois invalide")
    prefix = f"{year:04d}-{month:02d}"
    sends = await db.quote_sends.find(
        {"user_id": user["user_id"]}, {"_id": 0, "signature_b64": 0},
    ).to_list(5000)
    rows = [s for s in sends
            if (s.get("created_at") or "")[:7] == prefix
            or (s.get("signed_at") or "")[:7] == prefix]
    plan_ids = list({r["plan_id"] for r in rows if r.get("plan_id")})
    plan_names = {p["plan_id"]: p.get("filename")
                  for p in await db.plans_takeoff.find(
                      {"plan_id": {"$in": plan_ids}},
                      {"_id": 0, "plan_id": 1, "filename": 1}).to_list(2000)} if plan_ids else {}
    branding = await _resolve_branding(user["user_id"]) or {}
    pdf_bytes = build_accounting_recap_pdf(prefix, rows, plan_names, branding)
    return Response(
        content=pdf_bytes, media_type="application/pdf",
        headers={"Content-Disposition":
                 f'attachment; filename="recap-comptable-{prefix}.pdf"'})


@quote_router.get("/export.csv")
async def export_clients_csv(user=Depends(current_user)):
    """Export comptabilité : devis métrés envoyés + études PC (CSV Excel FR)."""
    import csv
    from io import StringIO
    from fastapi.responses import Response
    buf = StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["Type", "Date d'envoi", "Client", "Email client", "Document",
                "Montant TTC (EUR)", "Statut", "Signé le", "Acompte payé",
                "Facture acompte"])
    sends = await db.quote_sends.find(
        {"user_id": user["user_id"]}, {"_id": 0, "signature_b64": 0}
    ).sort("created_at", -1).to_list(5000)
    plan_ids = list({d["plan_id"] for d in sends})
    plans_names = {p["plan_id"]: p.get("filename") for p in await db.plans_takeoff.find(
        {"plan_id": {"$in": plan_ids}}, {"_id": 0, "plan_id": 1, "filename": 1}
    ).to_list(2000)} if plan_ids else {}
    for s_ in sends:
        status = ("Signé" if s_.get("signed")
                  else "Ouvert" if s_.get("opened")
                  else "Envoyé")
        w.writerow(["Devis métré", (s_.get("created_at") or "")[:10],
                    s_.get("client_name") or "", s_.get("client_email") or "",
                    plans_names.get(s_.get("plan_id")) or s_.get("plan_id") or "",
                    f"{(s_.get('total_ttc') or 0):.2f}".replace(".", ","),
                    status, (s_.get("signed_at") or "")[:10], "", ""])
    pc_plans = await db.plans_takeoff.find(
        {"user_id": user["user_id"], "sent_reports.0": {"$exists": True}},
        {"_id": 0, "plan_id": 1, "filename": 1, "sent_reports": 1}
    ).to_list(2000)
    for p in pc_plans:
        for r in p.get("sent_reports") or []:
            if r.get("kind") != "pc_study":
                continue
            status = ("Signée" if r.get("accepted")
                      else "Ouverte" if r.get("opened")
                      else "Envoyée")
            w.writerow(["Étude PC", (r.get("sent_at") or "")[:10],
                        r.get("accepted_by") or "", r.get("to_email") or "",
                        p.get("filename") or p["plan_id"],
                        f"{(r.get('ttc') or 0):.2f}".replace(".", ","),
                        status, (r.get("accepted_at") or "")[:10],
                        "Oui" if r.get("deposit_paid") else "Non",
                        r.get("deposit_invoice_no") or ""])
    csv_bytes = ("\ufeff" + buf.getvalue()).encode("utf-8")
    return Response(
        content=csv_bytes, media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition":
                 'attachment; filename="construpilot-clients-devis.csv"'})


# ============ SIGNATURE CLIENT ============

class SignatureLinkIn(BaseModel):
    deposit_pct: float | None = Field(default=None, ge=0, le=100)


@router.post("/{plan_id}/quote-sends/{log_id}/signature-link")
async def create_signature_link(plan_id: str, log_id: str,
                                body: SignatureLinkIn | None = None,
                                user=Depends(current_user)) -> dict:
    """Génère un lien public unique pour que le client signe le devis."""
    log = await db.quote_sends.find_one(
        {"log_id": log_id, "plan_id": plan_id, "user_id": user["user_id"]},
        {"_id": 0}
    )
    if not log:
        raise HTTPException(404, "Envoi de devis introuvable")
    token = SIG_TOKEN_PREFIX + secrets.token_urlsafe(18)
    set_fields = {"signature_token": token, "signature_token_created_at": _now()}
    if body and body.deposit_pct is not None:
        set_fields["deposit_pct"] = float(body.deposit_pct)
    await db.quote_sends.update_one({"log_id": log_id}, {"$set": set_fields})
    frontend_url = os.getenv("FRONTEND_URL") or ""
    return {"ok": True, "token": token,
            "url": f"{frontend_url}/sign/{token}" if frontend_url else f"/sign/{token}"}


@quote_router.get("/client-portal/{token}")
async def client_portal(token: str) -> dict:
    """Portail public client — liste tous les devis reçus pour un email donné."""
    from routes.emails import _sign
    import hmac
    import base64
    if "." not in token:
        raise HTTPException(400, "Lien invalide")
    email_b64, sig = token.rsplit(".", 1)
    try:
        email = base64.urlsafe_b64decode(email_b64.encode()).decode()
    except Exception:
        raise HTTPException(400, "Lien invalide")
    if not hmac.compare_digest(sig, _sign(email.lower())):
        raise HTTPException(400, "Lien invalide ou expiré")

    cursor = db.quote_sends.find(
        {"client_email": email.lower()},
        {"_id": 0, "log_id": 1, "plan_id": 1, "user_id": 1,
         "client_name": 1, "total_ttc": 1, "sent_at": 1, "signed": 1,
         "signed_at": 1, "opened_at": 1, "signature_token": 1, "finish": 1,
         "floor_material": 1, "company_snapshot": 1}
    ).sort("sent_at", -1).limit(200)
    docs = await cursor.to_list(200)

    for d in docs:
        pro = await db.users.find_one({"user_id": d["user_id"]},
                                       {"_id": 0, "name": 1, "email": 1})
        b = await _resolve_branding(d["user_id"]) or {}
        d["pro_name"] = b.get("company_name") or (pro or {}).get("name") or "ConstruPilot Pro"
        d["pro_email"] = (pro or {}).get("email")
        d.pop("user_id", None)
        d.pop("company_snapshot", None)

    return {"email": email, "count": len(docs), "quotes": docs}


@quote_router.get("/public/sign/{token}")
async def get_signature_page_data(token: str) -> dict:
    """Endpoint public : récupère le résumé du devis à signer."""
    log = await db.quote_sends.find_one(
        {"signature_token": token},
        {"_id": 0, "plan_id": 1, "user_id": 1, "client_email": 1, "client_name": 1,
         "total_ttc": 1, "finish": 1, "floor_material": 1, "log_id": 1,
         "signed": 1, "signed_at": 1, "deposit_pct": 1, "deposit_paid": 1,
         "deposit_amount": 1, "deposit_at": 1}
    )
    if not log:
        raise HTTPException(404, "Lien invalide ou expiré")
    branding = await _resolve_branding(log["user_id"]) or {}
    ttc = float(log.get("total_ttc") or 0)
    pct = log.get("deposit_pct")
    pct = QUOTE_DEPOSIT_DEFAULT_PCT if pct is None else float(pct)
    deposit_amount = round(ttc * pct / 100.0, 2) if (ttc > 0 and pct > 0) else 0
    return {
        "log_id": log["log_id"], "plan_id": log.get("plan_id"),
        "client_email": log.get("client_email"), "client_name": log.get("client_name"),
        "total_ttc": log.get("total_ttc"),
        "finish": log.get("finish"), "floor_material": log.get("floor_material"),
        "signed": bool(log.get("signed")),
        "signed_at": log.get("signed_at"),
        "company_name": branding.get("company_name"),
        "contact_line": branding.get("contact_line"),
        "deposit": {
            "pct": pct,
            "amount": deposit_amount,
            "paid": bool(log.get("deposit_paid")),
            "paid_amount": log.get("deposit_amount"),
            "paid_at": log.get("deposit_at"),
            "enabled": deposit_amount >= 1,
        },
    }


class SignatureIn(BaseModel):
    signature_b64: str = Field(min_length=100, max_length=200_000)
    signer_name: str = Field(min_length=2, max_length=120)


def _build_pro_signature_email(pro_name: str, signer: str, ttc: float,
                                client_email: str) -> str:
    ttc_html = (f" — <strong>{int(ttc):,} € TTC</strong>".replace(",", " ")
                if ttc else "")
    return (
        f"<div style='font-family:system-ui,sans-serif;max-width:560px;margin:0 auto;padding:24px'>"
        f"<div style='background:hsl(24,96%,55%);color:white;padding:20px;border-radius:6px 6px 0 0'>"
        f"<h1 style='margin:0;font-size:22px'>✍️ Devis signé</h1></div>"
        f"<div style='border:1px solid #eee;border-top:none;padding:24px;border-radius:0 0 6px 6px'>"
        f"<p>Bonjour {pro_name or ''},</p>"
        f"<p><strong>{signer}</strong> vient de signer votre devis{ttc_html}.</p>"
        f"<p>Client : <em>{client_email}</em><br/>"
        f"Date : {_now()[:10]}</p>"
        f"<p style='margin-top:24px'>"
        f"<a href='https://construpilot.fr/devis' "
        f"style='background:hsl(24,96%,55%);color:white;padding:12px 24px;"
        f"text-decoration:none;border-radius:4px;font-weight:600;display:inline-block'>"
        f"Voir le devis signé</a></p>"
        f"<p style='margin-top:32px;color:#888;font-size:12px'>"
        f"ConstruPilot · Estimez. Comparez. Construisez.</p>"
        f"</div></div>"
    )


async def _notify_pro_on_signature(log: dict, signer: str, client_email: str) -> None:
    """Best-effort push + email notifications to the pro."""
    try:
        from routes.phase7 import notify_user
        total = log.get("total_ttc") or 0
        title = "✍️ Devis signé !"
        push_body = (f"{signer} vient de signer votre devis"
                     + (f" ({int(total):,} € TTC)".replace(",", " ") if total else ""))
        await notify_user(log["user_id"], title, push_body, "/devis")
    except Exception as e:
        logger.warning(f"Push signature notify failed: {e}")
    try:
        from engines import emailer
        pro = await db.users.find_one({"user_id": log["user_id"]},
                                       {"_id": 0, "email": 1, "name": 1})
        if pro and pro.get("email"):
            html = _build_pro_signature_email(pro.get("name") or "", signer,
                                                log.get("total_ttc", 0), client_email)
            await emailer.send_email(pro["email"], f"✍️ {signer} a signé votre devis", html)
    except Exception as e:
        logger.warning(f"Email signature notify failed: {e}")


@quote_router.post("/public/sign/{token}")
async def submit_signature(token: str, body: SignatureIn) -> dict:
    """Endpoint public : le client soumet sa signature (image base64)."""
    log = await db.quote_sends.find_one({"signature_token": token},
                                          {"_id": 0, "log_id": 1, "user_id": 1,
                                           "signed": 1, "client_email": 1,
                                           "plan_id": 1, "doc_id": 1, "client_name": 1,
                                           "total_ttc": 1})
    if not log:
        raise HTTPException(404, "Lien invalide")
    if log.get("signed"):
        raise HTTPException(400, "Ce devis a déjà été signé")
    b64_clean = body.signature_b64.split(",")[-1]
    signer = body.signer_name.strip()[:120]
    client_email = log.get("client_email") or ""
    await db.quote_sends.update_one(
        {"log_id": log["log_id"]},
        {"$set": {"signed": True, "signed_at": _now(),
                   "signer_name": signer,
                   "signature_b64": b64_clean}}
    )
    # Bascule automatique du statut du devis lié → « Signé »
    if log.get("doc_id"):
        await db.pro_documents.update_one(
            {"doc_id": log["doc_id"], "user_id": log["user_id"]},
            {"$set": {"status": "signed", "signed": True,
                      "signed_at": _now(), "updated_at": _now()}})
    await db.notifications.insert_one({
        "notif_id": "n_" + secrets.token_hex(5),
        "user_id": log["user_id"], "type": "quote_signed", "log_id": log["log_id"],
        "message": f"{signer} a signé le devis ({client_email}).",
        "created_at": _now(), "read": False,
    })
    await _notify_pro_on_signature(log, signer, client_email)
    return {"ok": True, "signed_at": _now()}


# ---- Acompte en ligne (Stripe Checkout, page publique de signature) --------

class QuoteDepositCheckoutIn(BaseModel):
    origin_url: str = Field(max_length=200)


async def _quote_deposit_flood_guard(token: str) -> None:
    since = (datetime.now(timezone.utc) - timedelta(hours=1)).isoformat()
    n = await db.payment_transactions.count_documents(
        {"sign_token": token, "type": "quote_deposit",
         "payment_status": "pending", "created_at": {"$gte": since}})
    if n >= 10:
        raise HTTPException(429, "Trop de tentatives de paiement — réessayez plus tard")


@quote_router.post("/public/sign/{token}/deposit/checkout")
async def quote_deposit_checkout(token: str, body: QuoteDepositCheckoutIn) -> dict:
    """Crée la session Stripe d'acompte pour un devis signé (public)."""
    import stripe
    log = await db.quote_sends.find_one(
        {"signature_token": token},
        {"_id": 0, "log_id": 1, "user_id": 1, "plan_id": 1, "total_ttc": 1,
         "signed": 1, "deposit_pct": 1, "deposit_paid": 1, "client_name": 1})
    if not log:
        raise HTTPException(404, "Lien invalide ou expiré")
    if not log.get("signed"):
        raise HTTPException(400, "Signez d'abord le devis avant de verser l'acompte")
    if log.get("deposit_paid"):
        raise HTTPException(400, "L'acompte a déjà été réglé — merci !")
    await _quote_deposit_flood_guard(token)
    pct = log.get("deposit_pct")
    pct = QUOTE_DEPOSIT_DEFAULT_PCT if pct is None else float(pct)
    if pct <= 0:
        raise HTTPException(400, "Le paiement d'acompte en ligne est désactivé")
    ttc = float(log.get("total_ttc") or 0)
    amount_cents = int(round(ttc * pct))  # € TTC × pct % → centimes
    if amount_cents < 100:
        raise HTTPException(400, "Montant d'acompte invalide")
    origin = body.origin_url.rstrip("/")
    branding = await _resolve_branding(log["user_id"]) or {}
    company = (branding.get("company_name") or "ConstruPilot")[:60]
    try:
        session = stripe.checkout.Session.create(
            line_items=[{"price_data": {
                "currency": "eur",
                "product_data": {"name": f"Acompte {pct:g} % — Devis {company}"},
                "unit_amount": amount_cents}, "quantity": 1}],
            mode="payment",
            success_url=f"{origin}/sign/{token}?session_id={{CHECKOUT_SESSION_ID}}",
            cancel_url=f"{origin}/sign/{token}",
            metadata={"type": "quote_deposit", "sign_token": token,
                      "log_id": log["log_id"], "plan_id": log.get("plan_id") or ""})
    except Exception:
        logger.exception("stripe quote deposit checkout")
        raise HTTPException(502, "Paiement indisponible — réessayez plus tard")
    await db.payment_transactions.insert_one({
        "session_id": session.id, "user_id": log["user_id"],
        "plan_id": log.get("plan_id"), "log_id": log["log_id"],
        "sign_token": token, "type": "quote_deposit", "amount_cents": amount_cents,
        "currency": "eur", "mode": "payment", "status": "initiated",
        "payment_status": "pending", "created_at": _now(), "updated_at": _now()})
    return {"url": session.url, "session_id": session.id,
            "amount": round(amount_cents / 100.0, 2), "pct": pct}


async def _apply_quote_deposit(rec: dict) -> None:
    """Marque l'acompte du devis payé (idempotent) + notifie le pro."""
    amount = round((rec.get("amount_cents") or 0) / 100.0, 2)
    paid_at = _now()
    r = await db.quote_sends.update_one(
        {"log_id": rec.get("log_id"), "deposit_paid": {"$ne": True}},
        {"$set": {"deposit_paid": True, "deposit_amount": amount,
                  "deposit_at": paid_at,
                  "deposit_session_id": rec.get("session_id")}})
    if not r.modified_count:
        return
    try:
        from routes.phase7 import notify_user
        await notify_user(rec["user_id"], "💶 Acompte reçu !",
                          f"Un acompte de {int(amount):,} € a été réglé sur un devis signé.".replace(",", " "),
                          "/devis")
    except Exception as e:
        logger.warning(f"notify quote deposit failed: {e}")
    try:
        from engines import emailer
        pro = await db.users.find_one({"user_id": rec["user_id"]},
                                       {"_id": 0, "email": 1, "name": 1})
        if pro and pro.get("email"):
            html = (f"<div style='font-family:system-ui,sans-serif;max-width:560px;margin:0 auto;padding:24px'>"
                    f"<div style='background:hsl(24,96%,55%);color:white;padding:20px;border-radius:6px 6px 0 0'>"
                    f"<h1 style='margin:0;font-size:22px'>💶 Acompte reçu</h1></div>"
                    f"<div style='border:1px solid #eee;border-top:none;padding:24px;border-radius:0 0 6px 6px'>"
                    f"<p>Bonjour {pro.get('name') or ''},</p>"
                    f"<p>Un acompte de <strong>{amount} €</strong> vient d'être réglé en ligne "
                    f"sur un devis signé.</p>"
                    f"<p style='margin-top:24px'><a href='https://construpilot.fr/devis' "
                    f"style='background:hsl(24,96%,55%);color:white;padding:12px 24px;"
                    f"text-decoration:none;border-radius:4px;font-weight:600;display:inline-block'>"
                    f"Voir le suivi des devis</a></p></div></div>")
            await emailer.send_email(pro["email"], "💶 Acompte réglé sur un devis signé", html)
    except Exception as e:
        logger.warning(f"email quote deposit failed: {e}")


@quote_router.get("/deposit/status/{session_id}")
async def quote_deposit_status(session_id: str) -> dict:
    """Statut public d'un paiement d'acompte devis — polling au retour Stripe."""
    q = {"session_id": session_id, "type": "quote_deposit"}
    rec = await db.payment_transactions.find_one(q, {"_id": 0})
    if not rec:
        raise HTTPException(404, "Session inconnue")
    if rec.get("payment_status") != "paid":
        import stripe
        try:
            s = stripe.checkout.Session.retrieve(session_id)
            if s.payment_status == "paid" or s.status == "complete":
                upd = await db.payment_transactions.update_one(
                    {"session_id": session_id, "payment_status": {"$ne": "paid"}},
                    {"$set": {"status": "completed", "payment_status": "paid",
                              "stripe_payment_intent_id": s.payment_intent,
                              "updated_at": _now()}})
                rec = await db.payment_transactions.find_one(q, {"_id": 0})
                if upd.modified_count:
                    await _apply_quote_deposit(rec)
            elif s.status == "expired":
                await db.payment_transactions.update_one(
                    {"session_id": session_id},
                    {"$set": {"status": "expired", "payment_status": "expired",
                              "updated_at": _now()}})
                rec["payment_status"] = "expired"
        except Exception as e:
            logger.warning(f"quote deposit status poll {session_id}: {e}")
    return {"session_id": session_id, "status": rec.get("status"),
            "payment_status": rec.get("payment_status"),
            "amount": round((rec.get("amount_cents") or 0) / 100.0, 2)}



@quote_router.get("/public/sign/{token}/pdf")
async def get_signed_quote_pdf(token: str):
    """Renvoie le PDF Devis avec le bloc signature intégré."""
    from fastapi.responses import Response
    from engines.pdf_takeoff import build_quote_pdf
    log = await db.quote_sends.find_one({"signature_token": token}, {"_id": 0})
    if not log:
        raise HTTPException(404, "Lien invalide")
    plan_id = log.get("plan_id")
    user_id = log["user_id"]
    branding = await _resolve_branding(user_id)
    signature = None
    if log.get("signed"):
        signature = {"b64": log.get("signature_b64"),
                     "name": log.get("signer_name"),
                     "at": log.get("signed_at"),
                     "client_email": log.get("client_email")}

    if not plan_id and log.get("doc_id"):
        # Envoi par document (éditeur de devis) — PDF basé sur le document
        from engines.pdf_docs import build_document_pdf
        doc = await db.pro_documents.find_one({"doc_id": log["doc_id"], "user_id": user_id}, {"_id": 0})
        if not doc:
            raise HTTPException(404, "Devis introuvable")
        try:
            pdf_bytes = build_document_pdf(doc, branding=branding, signature=signature,
                                           lang=(doc.get("lang")))
        except TypeError:
            pdf_bytes = build_document_pdf(doc, branding=branding, lang=(doc.get("lang")))
        return Response(content=pdf_bytes, media_type="application/pdf",
                        headers={"Content-Disposition": f'attachment; filename="devis-signe-{log["doc_id"]}.pdf"'})

    quote, _pdf = await _build_quote_and_pdf(
        plan_id, user_id, log.get("finish", "std"),
        log.get("floor_material", "carrelage"),
        True, True, True, True,
    )
    pdf_bytes = build_quote_pdf(quote, branding=branding, signature=signature)
    return Response(
        content=pdf_bytes, media_type="application/pdf",
        headers={"Content-Disposition":
                 f'attachment; filename="devis-signe-{plan_id}.pdf"'}
    )
