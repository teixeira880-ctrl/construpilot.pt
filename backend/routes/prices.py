from datetime import datetime, timezone
import io
import secrets

from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from core.db import db
from core.security import current_user, optional_user
from engines import price_search, affiliate
from routes.phase6 import promote as promote_sponsored

router = APIRouter(prefix="/prices", tags=["prices"])


@router.get("/suppliers")
async def suppliers():
    extra = await db.suppliers.find({}, {"_id": 0}).to_list(200)
    return price_search.SUPPLIERS + extra


class SearchIn(BaseModel):
    query: str
    limit: int = 10


@router.post("/search")
async def search(body: SearchIn, user=Depends(optional_user)):
    if len(body.query.strip()) < 2:
        raise HTTPException(status_code=400, detail="Requête trop courte")
    await db.search_stats.update_one({"query": body.query.lower()},
                                     {"$inc": {"count": 1},
                                      "$set": {"last_at": datetime.now(timezone.utc).isoformat()}},
                                     upsert=True)
    result = await price_search.search(body.query, body.limit)
    enriched = await affiliate.enrich_prices(result.get("results") or [])
    result["results"] = await promote_sponsored(enriched, body.query)
    return result


@router.get("/history")
async def history(query: str):
    return await price_search.history(query)


class CartIn(BaseModel):
    items: list[dict]
    coefficient: float = 1.0     # coefficient de prix régional (départemental)


@router.get("/departments")
async def departments():
    """Coefficients de prix régionaux (indicatifs) pour ajuster les estimations."""
    return price_search.DEPARTMENT_COEFF


@router.post("/optimize")
async def optimize(body: CartIn, user=Depends(optional_user)):
    if not body.items:
        raise HTTPException(status_code=400, detail="Panier vide")
    catalog = None
    if user:
        catalog = await db.price_catalog.find({"user_id": user["user_id"]}, {"_id": 0}).to_list(1000)
    coef = body.coefficient if 0.5 <= body.coefficient <= 2.0 else 1.0
    return price_search.optimize_cart(body.items, catalog=catalog, coefficient=coef)


@router.get("/catalog/template")
async def catalog_template():
    """Modèle Excel prêt à remplir pour l'import de tarifs."""
    import openpyxl
    from openpyxl.styles import Font, PatternFill
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Tarifs"
    headers = ["nom", "prix", "unité", "fournisseur", "catégorie"]
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="16A34A")
    examples = [
        ["Plaque BA13 250x120", 6.50, "plaque", "Point P", "Placo"],
        ["Rail R48 (barre 3m)", 3.20, "barre", "Point P", "Placo"],
        ["Carrelage grès 60x60", 24.90, "m²", "Leroy Merlin", "Carrelage"],
        ["Peinture blanche 10L", 45.90, "pot", "Leroy Merlin", "Peinture"],
    ]
    for row in examples:
        ws.append(row)
    for col, w in zip("ABCDE", (28, 10, 10, 16, 16)):
        ws.column_dimensions[col].width = w
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return StreamingResponse(
        buf, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="modele-tarifs-construpilot.xlsx"'})


# ─────────────────────────── Catalogue de prix perso ───────────────────────────

class CatalogItemIn(BaseModel):
    name: str
    price: float
    unit: str = "u"
    supplier: str = ""
    category: str = ""


@router.get("/catalog")
async def catalog_list(user=Depends(current_user)):
    rows = await db.price_catalog.find({"user_id": user["user_id"]}, {"_id": 0}) \
        .sort("name", 1).to_list(2000)
    return rows


@router.post("/catalog")
async def catalog_upsert(body: CatalogItemIn, user=Depends(current_user)):
    if not body.name.strip():
        raise HTTPException(422, "Nom requis")
    now = datetime.now(timezone.utc).isoformat()
    existing = await db.price_catalog.find_one(
        {"user_id": user["user_id"], "name_lc": body.name.lower().strip()}, {"_id": 0, "catalog_id": 1})
    doc = {"user_id": user["user_id"], "name": body.name.strip(), "name_lc": body.name.lower().strip(),
           "price": round(float(body.price), 2), "unit": body.unit, "supplier": body.supplier,
           "category": body.category, "updated_at": now}
    if existing:
        await db.price_catalog.update_one({"catalog_id": existing["catalog_id"]}, {"$set": doc})
        return {**doc, "catalog_id": existing["catalog_id"]}
    doc["catalog_id"] = "pc_" + secrets.token_hex(6)
    await db.price_catalog.insert_one(dict(doc))
    doc.pop("_id", None)
    return doc


@router.delete("/catalog/{catalog_id}")
async def catalog_delete(catalog_id: str, user=Depends(current_user)):
    await db.price_catalog.delete_one({"catalog_id": catalog_id, "user_id": user["user_id"]})
    return {"ok": True}


@router.post("/catalog/import")
async def catalog_import(file: UploadFile = File(...), user=Depends(current_user)):
    """Import CSV ou Excel. Colonnes attendues : name/désignation, price/prix, unit/unité, supplier/fournisseur, category/catégorie."""
    raw = await file.read()
    fn = (file.filename or "").lower()
    rows: list[dict] = []
    try:
        if fn.endswith(".xlsx") or fn.endswith(".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.active
            data = list(ws.iter_rows(values_only=True))
            if not data:
                raise HTTPException(400, "Fichier vide")
            headers = [str(h or "").lower().strip() for h in data[0]]
            for r in data[1:]:
                rows.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
        else:
            import csv
            text = raw.decode("utf-8-sig", errors="ignore")
            dialect = csv.Sniffer().sniff(text[:2000], delimiters=";,\t") if text.strip() else csv.excel
            reader = csv.DictReader(io.StringIO(text), dialect=dialect)
            for r in reader:
                rows.append({(k or "").lower().strip(): v for k, v in r.items()})
    except HTTPException:
        raise
    except Exception:
        raise HTTPException(400, "Fichier illisible. Utilisez un CSV (;) ou un .xlsx avec en-têtes.")

    def pick(d, *keys):
        for k in keys:
            if k in d and d[k] not in (None, ""):
                return d[k]
        return None

    def to_float(v):
        if v is None:
            return None
        s = str(v).replace("€", "").replace(",", ".").strip()
        try:
            return round(float(s), 2)
        except ValueError:
            return None

    now = datetime.now(timezone.utc).isoformat()
    ops, imported, skipped = [], 0, 0
    for d in rows:
        name = pick(d, "name", "nom", "désignation", "designation", "libellé", "libelle", "produit", "matériau", "materiau")
        price = to_float(pick(d, "price", "prix", "prix ht", "prix_ht", "tarif", "pu"))
        if not name or price is None:
            skipped += 1
            continue
        unit = pick(d, "unit", "unité", "unite", "u") or "u"
        supplier = pick(d, "supplier", "fournisseur", "négoce", "negoce") or ""
        category = pick(d, "category", "catégorie", "categorie", "corps d'état", "lot") or ""
        doc = {"user_id": user["user_id"], "name": str(name).strip(), "name_lc": str(name).lower().strip(),
               "price": price, "unit": str(unit), "supplier": str(supplier), "category": str(category),
               "updated_at": now, "catalog_id": "pc_" + secrets.token_hex(6)}
        await db.price_catalog.update_one(
            {"user_id": user["user_id"], "name_lc": doc["name_lc"]},
            {"$set": {k: v for k, v in doc.items() if k != "catalog_id"},
             "$setOnInsert": {"catalog_id": doc["catalog_id"]}}, upsert=True)
        imported += 1
    return {"imported": imported, "skipped": skipped}


@router.get("/stores")
async def stores(lat: float, lng: float, radius_km: int = 20) -> dict:
    return {"radius_km": radius_km, "stores": price_search.stores_nearby(lat, lng, radius_km)}


class AlertIn(BaseModel):
    query: str
    target_price: float
    product: str | None = None
    channels: list[str] = ["email"]  # "email" et/ou "sms"
    phone: str | None = None  # requis si "sms" dans channels (format E.164 +33…)


@router.post("/alerts")
async def create_alert(body: AlertIn, user=Depends(current_user)):
    channels = [c for c in (body.channels or ["email"]) if c in ("email", "sms")]
    if not channels:
        channels = ["email"]
    if "sms" in channels:
        if not body.phone:
            raise HTTPException(422, "Numéro requis pour recevoir les alertes SMS")
        normalized = "".join(ch for ch in body.phone if ch.isdigit() or ch == "+")
        if not normalized.startswith("+") or len(normalized) < 8:
            raise HTTPException(422, "Numéro invalide : utilisez le format +33…")
    doc = {"alert_id": f"al_{datetime.now(timezone.utc).timestamp():.0f}",
           "user_id": user["user_id"], "query": body.query, "target_price": body.target_price,
           "product": body.product or body.query, "active": True,
           "channels": channels,
           "phone": ("".join(ch for ch in body.phone if ch.isdigit() or ch == "+")
                     if "sms" in channels and body.phone else None),
           "created_at": datetime.now(timezone.utc).isoformat()}
    await db.price_alerts.insert_one(doc)
    doc.pop("_id", None)
    return doc


@router.get("/alerts")
async def list_alerts(user=Depends(current_user)):
    alerts = await db.price_alerts.find({"user_id": user["user_id"]}, {"_id": 0}).to_list(200)
    for a in alerts:
        offers = price_search.demo_search(a["query"], 10)
        best = min((o["price"] for o in offers), default=None)
        a["best_price"] = best
        a["triggered"] = bool(best and best <= a["target_price"])
    return alerts


@router.post("/alerts/check-now")
async def check_alerts_now(user=Depends(current_user)) -> dict:
    """Force la vérification immédiate des alertes de l'utilisateur + envoi email des déclenchées."""
    from engines import alerting
    sent = await alerting.check_user_alerts(user["user_id"])
    return {"ok": True, "notified": sent}


@router.delete("/alerts/{alert_id}")
async def delete_alert(alert_id: str, user=Depends(current_user)) -> dict:
    await db.price_alerts.delete_one({"alert_id": alert_id, "user_id": user["user_id"]})
    return {"ok": True}
